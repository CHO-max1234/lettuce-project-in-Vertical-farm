"""Scale validation calculations used by the final cycle pipeline."""
from pathlib import Path as _Path
import sys as _sys
_ROOT = _Path(__file__).resolve().parents[1]
for _folder in (_ROOT, _ROOT / 'pipeline', _ROOT / 'preprocessing'):
    _sys.path.insert(0, str(_folder))
import settings as CYCLE
from runtime import resolved_spec, run_source, moved_tree, moved_path, yolo_device
from pathlib import Path
CODE = Path(__file__).with_name('scale_estimation.py')
ROOT = CYCLE.PREPROCESS_OUTPUT / '3. 밝기보정과 Scale'
RUN = ROOT / '_실행기록' / '261009_SCALE'
WORKBOOK = ROOT / '4작기_scale_map.xlsx'
import argparse
import csv
import hashlib
import importlib.util
import json
import math
import statistics
import sys
from collections import Counter
from datetime import datetime
sys.dont_write_bytecode = True
import cv2
import numpy as np
from PIL import Image, ImageDraw, ImageFont

def module():
    spec = resolved_spec('season4_scale', CODE)
    m = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(m)
    return m

def journal():
    for path in sorted(RUN.glob('detection_records*.jsonl')):
        with path.open(encoding='utf-8') as f:
            for line in f:
                try:
                    yield json.loads(line)
                except json.JSONDecodeError:
                    break

def boards(ongoing=False):
    m = module()
    ref = m.load_reference()
    groups = {c: [] for c in ('cam0', 'cam1')}
    for x in journal():
        r = x['row']
        c = r['camera']
        if c in groups and r['window_mode'] == 'X_EXPANDED_ON_MISS_TARGET_MASK' and (len(groups[c]) < 3):
            groups[c].append(x)
        if all((len(v) == 3 for v in groups.values())):
            break
    outputs = []
    font = ImageFont.truetype('C:/Windows/Fonts/malgun.ttf', 26)
    for (c, records) in groups.items():
        views = []
        for x in records:
            r = x['row']
            a = x['audit']
            im = ref.read_image(Path(r['source_path']))
            if hashlib.sha256(Path(r['source_path']).read_bytes()).hexdigest() != r['source_sha256']:
                raise ValueError('QC sample hash changed')
            view = ref.overlay(im, np.asarray(a['roi_polygon'], np.int32), a['accepted_candidates'], 'X WINDOW RECOVERY; SAME HOLE CONDITIONS')
            p = Image.fromarray(cv2.cvtColor(view, cv2.COLOR_BGR2RGB))
            p.thumbnail((2000, 1000))
            views.append((r, p))
        canvas = Image.new('RGB', (2000, sum((v.height + 80 for (_, v) in views)) + 80), '#f4f5f6')
        draw = ImageDraw.Draw(canvas)
        draw.text((15, 15), f'{c.upper()} | 미검출 범위 보완 후 추가 검출 3장', font=font, fill='#17222a')
        y = 80
        for (r, v) in views:
            draw.text((15, y), f"{r['image_name']} | {r['selected_diameter_px']:.2f}px", font=font, fill='#17222a')
            canvas.paste(v, (0, y + 45))
            y += v.height + 80
        p = RUN / f'261009_{c.upper()}_범위보완_결과확인.png'
        canvas.save(p)
        outputs.append(str(p))
    print(json.dumps(dict(review_boards=outputs, ongoing=ongoing), ensure_ascii=False))

def verify():
    m = module()
    inputs = m.load_inputs()
    ref = m.load_reference()
    with (RUN / 'scale_map.csv').open(encoding='utf-8-sig', newline='') as f:
        rows = list(csv.DictReader(f))
    if len(rows) != 39401 or [r['image_key'] for r in rows] != [r['image_key'] for r in inputs]:
        raise ValueError('final row count/order mismatch')
    by_key = {r['image_key']: r for r in rows}
    if len(by_key) != 39401:
        raise ValueError('duplicate scale key')
    counters = Counter()
    for (r, upstream) in zip(rows, inputs):
        assert r['label_accepted'] == 'True' and r['scale_verified'] == 'False'
        assert r['source_path'] == upstream['output_path'] and r['source_sha256'] == upstream['output_sha256']
        assert r['ocr_qc_status'] == upstream['qc_status'] and r['crop_bbox'] == upstream['crop_bbox']
        assert int(r['native_width']) == upstream['metadata']['native_width']
        assert int(r['native_height']) == upstream['metadata']['native_height']
        assert r['source_camera'] == upstream['source_camera'] and r['date'] == upstream['date']
        assert r['input_hash_verified'] == 'True'
        ready = r['scale_ready'] == 'True'
        source = r['scale_source']
        counters[source] += 1
        if ready:
            assert math.isclose(float(r['mm_per_px']) * float(r['diameter_px']), 50, abs_tol=1e-10)
        else:
            assert not r['diameter_px'] and (not r['mm_per_px'])
        if source == 'DIRECT':
            assert float(r['diameter_px']) == float(r['selected_diameter_px']) and int(r['reference_count']) == 0
        elif source == 'MEDIAN_SAME_GROUP':
            n = int(r['reference_count'])
            assert 1 <= n <= 3
            refs = [by_key[r[f'reference_{i}']] for i in range(1, n + 1)]
            assert all((x['scale_source'] == 'DIRECT' and x['reference_seed_eligible'] == 'True' and (x['geometry_group'] == r['geometry_group']) for x in refs))
            assert all(((x['camera'], x['native_width'], x['native_height'], x['source_camera']) == (r['camera'], r['native_width'], r['native_height'], r['source_camera']) for x in refs))
            assert math.isclose(statistics.median((float(x['selected_diameter_px']) for x in refs)), float(r['diameter_px']), abs_tol=1e-10)
    selected = {}
    candidate_count = 0
    with (RUN / 'candidates.csv').open(encoding='utf-8-sig', newline='') as f:
        for c in csv.DictReader(f):
            candidate_count += 1
            assert c['image_key'] in by_key
            assert math.isclose(float(c['mm_per_px']) * float(c['diameter_px']), 50, abs_tol=1e-10)
            if c['selected'] == 'True':
                assert c['image_key'] not in selected
                selected[c['image_key']] = c
    assert len(selected) == counters['DIRECT']
    for (key, c) in selected.items():
        assert c['overlay_id'] == by_key[key]['selected_overlay_id']
        assert float(c['diameter_px']) == float(by_key[key]['selected_diameter_px'])
    contract = json.loads((RUN / 'execution_contract.json').read_text(encoding='utf-8'))
    assert hashlib.sha256(CODE.read_bytes()).hexdigest() == contract['script_sha256']
    assert m.digest(m.REFERENCE_CODE) == contract['reference_sha256']
    assert m.digest(m.MANIFEST) == contract['brightness_manifest_sha256']
    assert [m.digest(p) for p in m.ROI_MANIFESTS] == contract['roi_manifest_sha256']
    from openpyxl import load_workbook
    book = load_workbook(WORKBOOK, read_only=True, data_only=True)
    assert book.sheetnames == ['scale_map', 'candidate', '산정기준']
    ws = book['scale_map']
    iterator = ws.iter_rows(values_only=True)
    headers = list(next(iterator))
    assert headers == ref.SCALE_FIELDS
    n = 0
    for values in iterator:
        r = rows[n]
        n += 1
        for (k, v) in zip(headers, values):
            expected = r[k]
            if v is None:
                assert expected == ''
            elif isinstance(v, bool):
                assert str(v) == expected
            elif isinstance(v, (int, float)):
                assert math.isclose(float(v), float(expected), rel_tol=1e-12, abs_tol=1e-10)
            else:
                assert v == expected
    assert n == 39401
    ci = book['candidate'].iter_rows(values_only=True)
    assert list(next(ci)) == ref.CANDIDATE_FIELDS
    with (RUN / 'candidates.csv').open(encoding='utf-8-sig', newline='') as f:
        cn = 0
        for (values, c) in zip(ci, csv.DictReader(f)):
            cn += 1
            for (k, v) in zip(ref.CANDIDATE_FIELDS, values):
                if isinstance(v, bool):
                    assert str(v) == c[k]
                elif isinstance(v, (int, float)):
                    assert math.isclose(v, float(c[k]), rel_tol=1e-12, abs_tol=1e-10)
                else:
                    assert v == c[k]
        assert cn == candidate_count and next(ci, None) is None
    book.close()
    check = dict(verified_at=datetime.now().isoformat(timespec='seconds'), rows=39401, status_counts=dict(counters), row_order_and_accepted_set_preserved=True, sha_matches_brightness_manifest=39401, csv_xlsx_all_scale_rows_equal=True, csv_xlsx_candidate_rows_equal=candidate_count, direct_candidate_linkage_valid=True, reference_groups_and_medians_valid=True, real_diameter_identity_mm=50, native_resolutions_and_source_camera_reference_groups_separated=True, upstream_code_manifests_unchanged=True, physical_scale_verified=False, workbook_sha256=m.digest(WORKBOOK))
    m.write_json(RUN / 'final_verification.json', check)
    summary = json.loads((RUN / 'summary.json').read_text(encoding='utf-8'))
    status = 'COMPLETE_WITH_UNAVAILABLE' if counters['ERROR'] or counters['UNRESOLVED'] else 'COMPLETE'
    m.write_json(RUN / 'completion.json', dict(status=status, **summary, workbook=str(WORKBOOK), verification=check))
    print(json.dumps(check, ensure_ascii=False, indent=2))
if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--ongoing', action='store_true')
    args = parser.parse_args()
    if not args.ongoing:
        verify()
    boards(args.ongoing)
