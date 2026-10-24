"""Scale writer calculations used by the final cycle pipeline."""
from pathlib import Path as _Path
import sys as _sys
_ROOT = _Path(__file__).resolve().parents[1]
for _folder in (_ROOT, _ROOT / 'pipeline', _ROOT / 'preprocessing'):
    _sys.path.insert(0, str(_folder))
import settings as CYCLE
from runtime import resolved_spec, run_source, moved_tree, moved_path, yolo_device
from pathlib import Path
CODE_DIR = Path(__file__).resolve().parent
ROOT = CYCLE.PREPROCESS_OUTPUT / '3. 밝기보정과 Scale'
RUN = ROOT / '_실행기록' / '261009_SCALE'
TEMP = ROOT / '4작기_scale_map.writing.xlsx'
import csv, json, importlib.util, sys
sys.dont_write_bytecode = True
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
BOOLEAN = {'detected', 'scale_ready', 'scale_verified'}
INTEGER = {'n_candidates', 'reference_count', 'image_width', 'image_height'}
NUMBER = {'selected_x_px', 'selected_y_px', 'selected_diameter_px', 'selected_minor_px', 'selected_mm_per_px', 'bottom_angle_deg', 'fallback_median_diameter_px', 'fallback_median_mm_per_px', 'diameter_px', 'mm_per_px', 'reference_max_angle_diff_deg', 'reference_diameter_range_ratio', 'bottom_left_x', 'bottom_left_y', 'bottom_right_x', 'bottom_right_y'}

def csv_records(path, fields, candidates=False):
    with path.open(encoding='utf-8-sig', newline='') as f:
        for r in csv.DictReader(f):
            values = []
            for h in fields:
                v = r[h]
                if v == '':
                    v = None
                elif candidates and h == 'selected' or h in BOOLEAN:
                    v = v == 'True'
                elif h in INTEGER:
                    v = int(v)
                elif h in NUMBER or (candidates and h in fields[6:]):
                    v = float(v)
                values.append(v)
            yield values

def main():
    spec = resolved_spec('season4_export', CODE_DIR / 'scale_estimation.py')
    m = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(m)
    ref = m.load_reference()
    summary = json.loads((RUN / 'summary.json').read_text(encoding='utf-8'))
    book = Workbook(write_only=True)

    def sheet(name, fields, records):
        ws = book.create_sheet(name)
        ws.sheet_view.showGridLines = False
        if name != '산정기준':
            ws.freeze_panes = 'D2'
        for (i, h) in enumerate(fields, 1):
            width = 20
            if h in ('image_key', 'image_name'):
                width = 53
            if h in ('qc', 'scale_qc'):
                width = 74
            if h == 'geometry_group' or h == 'scale_version':
                width = 62
            if h in ('source_path', 'source_sha256', 'reference_1', 'reference_2', 'reference_3'):
                width = 70
            if name == '산정기준':
                width = [29, 64, 98][i - 1]
            ws.column_dimensions[get_column_letter(i)].width = width
        ws.row_dimensions[1].height = 46
        header = []
        for h in fields:
            c = WriteOnlyCell(ws, h)
            c.font = Font(name='맑은 고딕', size=10, bold=True, color='FFFFFF')
            c.fill = PatternFill('solid', fgColor='24425E')
            c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            header.append(c)
        ws.append(header)
        count = 0
        for values in records:
            count += 1
            cells = []
            if name == '산정기준':
                ws.row_dimensions[count + 1].height = 50
            for (h, v) in zip(fields, values):
                c = WriteOnlyCell(ws, v)
                c.font = Font(name='맑은 고딕', size=10, color='17222A')
                if isinstance(v, str):
                    c.data_type = 's'
                c.alignment = Alignment(vertical='center', wrap_text=name == '산정기준')
                if isinstance(v, float):
                    c.number_format = '0.0000' if 'mm_per_px' in h else '0.0%' if h == 'reference_diameter_range_ratio' else '0.00'
                if h == 'scale_source' and v != 'DIRECT' or (h == 'detected' and v is False):
                    c.fill = PatternFill('solid', fgColor='FFF0CC')
                if h == 'selected' and v is True:
                    c.fill = PatternFill('solid', fgColor='FFD6D6')
                cells.append(c)
            ws.append(cells)
        if name != '산정기준':
            ws.auto_filter.ref = f'A1:{get_column_letter(len(fields))}{count + 1}'
        print(f'{name}: {count} rows', flush=True)
        return count
    assert sheet('scale_map', ref.SCALE_FIELDS, csv_records(RUN / 'scale_map.csv', ref.SCALE_FIELDS)) == 39401
    assert sheet('candidate', ref.CANDIDATE_FIELDS, csv_records(RUN / 'candidates.csv', ref.CANDIDATE_FIELDS, True)) == 86958
    notes = [['대상 장수', 39401, '직접 검출 38,321장. 같은 그룹 참조 보충 1,080장. 미해결/처리 오류 0장.'], ['홀 기준 지름(mm)', 50, '3작기 설정값 유지. 4작기 판의 실측 검증은 하지 않음.'], ['픽셀 지름', '타원 장축', '단축과 후보별 점수는 candidate 시트에 보존.'], ['mm_per_px', '50 / diameter_px', 'Python 계산값. 이미지 크기는 변경하지 않음.'], ['입력', '4작기 밝기보정 완료 PNG', 'label_accepted=True인 39,401장. OCR REVIEW 상태도 CSV에 보존.'], ['검출 조건', '3작기 함수 재사용', '밝기 임계값, 장축/단축, 둘레 완전성, 후보 점수 유지.'], ['가로 탐색 보완', 'CAM0/1 미검출 시에만 확장', '상하 경계 연장 후 upstream 대상 bed polygon 안으로 제한. 기존 성공 선택값 유지.'], ['직접 검출', 'DIRECT', '최고 점수 후보 장축 사용. 그룹 이상치여도 계산값 유지 + QC 표시.'], ['미검출 보충', 'MEDIAN_SAME_GROUP', '같은 카메라·원본 해상도·source camera·기존 구도에서 하단 기울기가 가까운 최대 3장 장축 중앙값.'], ['참조 제외', '그룹 중앙값 ±30% 밖 직접값 / 대상 polygon 밖', '직접값은 보존. 미검출 보충의 참조에서만 제외. 다른 QC 참조 사용은 별도 표시.'], ['보충 불가', 'UNRESOLVED', '같은 그룹에 직접 참조가 없으면 값 없음. 이번 결과는 0장.'], ['참조 QC', '각도 차 2.5도 / 지름범위비 25%', '초과 시 계산 유지 + QC 기록. 참고 3장 미만과 각도 범위 밖도 표시.'], ['scale_verified', False, '실측으로 정확성을 확인한 상태가 아님. scale_ready는 계산값 확보 여부.'], ['후기 솟은 잎', '높이 보정 미적용', '홀은 판 평면의 환산 기준. 잎 높이와 위치에 따른 원근 차이는 남음.'], ['조인키', 'image_key + camera', 'image_key에 cam/날짜/파일명 보존. 추가 OCR·원본 해상도 기록은 CSV에 보존.'], ['계산 재현', 'execution_contract.json', '검출 코드·입력 SHA와 적용 정책 보존. 기준/계산 변경 시 코드에서 다시 계산.'], ['대용량 저장', '3작기 write_only 저장 방식', 'artifact-tool 전체 표 렌더가 8GB heap에서 실패하여 행 단위 저장으로 전환. 계산값 변경 없음.']]
    sheet('산정기준', ['항목', '값', '설명'], notes)
    TEMP.parent.mkdir(parents=True, exist_ok=True)
    book.save(TEMP)
    print('Exported:', TEMP, flush=True)
    m.write_json(RUN / 'workbook_export_backend.json', dict(backend='openpyxl_write_only_legacy_style', reason='artifact-tool full workbook render failed with 8GB JavaScript heap OOM', cells=39401 * len(ref.SCALE_FIELDS) + 86958 * len(ref.CANDIDATE_FIELDS), scale_values_changed=False, candidates=86958))
if __name__ == '__main__':
    main()
