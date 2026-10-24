"""Preserve raw sensor observations and compute causal partial environment features."""
from pathlib import Path as _Path
import sys as _sys
_ROOT = _Path(__file__).resolve().parents[1]
for _folder in (_ROOT, _ROOT / 'pipeline', _ROOT / 'preprocessing'):
    _sys.path.insert(0, str(_folder))
import settings as CYCLE
from runtime import resolved_spec, run_source, moved_tree, moved_path, yolo_device
from pathlib import Path
SOURCE = CYCLE.ENVIRONMENT_FILE
TARGETS = CYCLE.TARGET_FILE
C_CODE = CYCLE.PROJECT_ROOT / 'model/environment_contract.py'
REFERENCE2 = CYCLE.DATA_ROOT / 'optional_development_environment.xlsx'
MODEL_ROOT = CYCLE.C_MODELS
OUTPUT_ROOT = CYCLE.ENVIRONMENT_ROOT
(START, END) = ('2026-07-23', '2026-08-25')
MIN_STRESS_MINUTES = 48
RAW_EC_DIVISOR = 1000.0
import argparse
from collections import Counter
from contextlib import closing
from datetime import datetime, time, timezone
import csv
import hashlib
import importlib.util
import json
import math
import sys
import time as timer
from zoneinfo import ZoneInfo
sys.dont_write_bytecode = True
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from tqdm import tqdm
ALLOWED = ['heat_hours_28_ratio', 'heat_degree_hours_28', 'vpd_hours_ratio_1_8', 'vpd_degree_hours_1_8', 'hot_dry_load_28_1_8', 'past_7d_hot_dry_load_sum']
RAW_OUTPUT = ['source_excel_row', 'date', 'time', 'timestamp_local', 'epoch_raw', 'temp_raw', 'rh_raw', 'ec_raw', 'EC_source_formula_value', 'ph_raw', 'water_temp_raw', 'ec_observed_mS_cm', 'temp_rh_model_valid', 'ec_conversion_zero_from_missing', 'ph_outside_review_0_14', 'qc_reasons']

def sha(path):
    h = hashlib.sha256()
    with Path(path).open('rb') as f:
        for block in iter(lambda : f.read(1024 * 1024), b''):
            h.update(block)
    return h.hexdigest()

def clean(x):
    if isinstance(x, dict):
        return {str(k): clean(v) for (k, v) in x.items()}
    if isinstance(x, (tuple, list)):
        return [clean(v) for v in x]
    if isinstance(x, (Path, datetime, pd.Timestamp)):
        return str(x)
    if isinstance(x, np.generic):
        return clean(x.item())
    if isinstance(x, float) and (not math.isfinite(x)):
        return None
    return x

def save_json(path, value):
    Path(path).write_text(json.dumps(clean(value), ensure_ascii=False, indent=2, allow_nan=False), encoding='utf-8')

def number(x):
    try:
        x = float(x)
        return x if math.isfinite(x) else np.nan
    except (ValueError, TypeError):
        return np.nan

def csv_text(x):
    return '' if x is None else str(x)

def raw_digest_update(digest, row):
    fields = [csv_text(row[k]) for k in ['source_excel_row', 'ec_raw', 'EC_source_formula_value', 'ph_raw']]
    digest.update((json.dumps(fields, ensure_ascii=False) + '\n').encode('utf-8'))

def load_c_contract():
    spec = resolved_spec('frozen_c_environment', C_CODE)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    assert len(module.ENV_COLS) == 17 and len(module.RAW) == 12
    return module

def source_minutes(out):
    (all_dates, offsets, flags) = (Counter(), Counter(), Counter())
    source_digest = hashlib.sha256()
    (in_scope, seen) = ([], {})
    (duplicates, conflicts, units_checked) = (0, 0, 0)
    first_last = []
    with closing(load_workbook(SOURCE, read_only=True, data_only=True)) as wb, (out / '01_environment_minute_preserved.csv').open('w', encoding='utf-8-sig', newline='') as f:
        sheet = wb.worksheets[0]
        rows = sheet.iter_rows(values_only=True)
        headers = list(next(rows))
        names = ['날짜', '시간', 'epoch', 'temp', 'rh', 'ec', 'EC', 'ph', 'water_temp']
        ix = {k: headers.index(k) for k in names}
        assert len(set(ix.values())) == len(ix)
        writer = csv.DictWriter(f, RAW_OUTPUT)
        writer.writeheader()
        for (excel_row, values) in enumerate(tqdm(rows, total=(sheet.max_row or 1) - 1, desc='Processing', unit='row', mininterval=5), 2):
            (day_raw, tm_raw) = (values[ix['날짜']], values[ix['시간']])
            if day_raw is None:
                flags['blank_date_rows'] += 1
                continue
            day = pd.Timestamp(day_raw).date().isoformat()
            tm = tm_raw.strftime('%H:%M:%S') if isinstance(tm_raw, (datetime, time)) else str(tm_raw)
            ts = datetime.fromisoformat(day + ' ' + tm)
            all_dates[day] += 1
            ep = number(values[ix['epoch']])
            if math.isfinite(ep):
                utc = datetime.fromtimestamp(ep, timezone.utc).replace(tzinfo=None)
                offsets[str(round((ts - utc).total_seconds() / 3600, 6))] += 1
            if not START <= day <= END:
                continue
            (t, rh, ec, converted, ph, water) = [number(values[ix[k]]) for k in ['temp', 'rh', 'ec', 'EC', 'ph', 'water_temp']]
            valid_core = math.isfinite(t) and -20 <= t <= 60 and math.isfinite(rh) and (0 <= rh <= 100)
            ec_valid = math.isfinite(ec) and ec >= 0
            zero_from_missing = not math.isfinite(ec) and converted == 0
            ph_review = math.isfinite(ph) and (not 0 <= ph <= 14)
            reasons = []
            for (flag, condition) in [('TEMP_RH_UNAVAILABLE', not valid_core), ('EC_RAW_MISSING', not math.isfinite(ec)), ('EC_NEGATIVE_REVIEW', math.isfinite(ec) and ec < 0), ('EC_FORMULA_ZERO_FROM_MISSING', zero_from_missing), ('PH_MISSING', not math.isfinite(ph)), ('PH_OUTSIDE_REVIEW_0_14', ph_review), ('WATER_MISSING', not math.isfinite(water))]:
                if condition:
                    reasons.append(flag)
                    flags[flag] += 1
            if math.isfinite(ec) and math.isfinite(converted):
                assert math.isclose(converted, ec / RAW_EC_DIVISOR, rel_tol=1e-10, abs_tol=1e-10), 'Input contract failed: math.isclose(converted, ec / RAW_EC_DIVISOR, rel_tol=1e-10, abs_tol=1e-10)'
                units_checked += 1
            row = dict(source_excel_row=excel_row, date=day, time=tm, timestamp_local=ts.isoformat(sep=' '), epoch_raw=values[ix['epoch']], temp_raw=values[ix['temp']], rh_raw=values[ix['rh']], ec_raw=values[ix['ec']], EC_source_formula_value=values[ix['EC']], ph_raw=values[ix['ph']], water_temp_raw=values[ix['water_temp']], ec_observed_mS_cm=ec / RAW_EC_DIVISOR if ec_valid else '', temp_rh_model_valid=valid_core, ec_conversion_zero_from_missing=zero_from_missing, ph_outside_review_0_14=ph_review, qc_reasons=';'.join(reasons))
            raw_digest_update(source_digest, row)
            writer.writerow(row)
            key = ts.isoformat()
            fingerprint = tuple((csv_text(row[k]) for k in RAW_OUTPUT[4:10]))
            if key in seen:
                duplicates += 1
                if seen[key] != fingerprint:
                    conflicts += 1
            else:
                seen[key] = fingerprint
                in_scope.append(dict(row, _t=t, _rh=rh, _ec=ec, _ph=ph, _water=water, _ts=ts))
            first_last.append(key)
        metadata = dict(source_sheet=sheet.title, exact_source_columns=ix, all_source_rows=sum(all_dates.values()), all_source_date_min=min(all_dates), all_source_date_max=max(all_dates), all_source_days=len(all_dates), scope_rows=len(first_last), scope_unique_timestamps=len(seen), scope_first=min(first_last), scope_last=max(first_last), duplicate_timestamps=duplicates, conflicting_duplicates=conflicts, raw_ec_divisor=RAW_EC_DIVISOR, ec_unit_relation_verified_rows=units_checked, epoch_utc_to_local_hour_offsets=offsets, raw_flags=flags, raw_ec_ph_fingerprint=source_digest.hexdigest())
    assert conflicts == 0, 'Input contract failed: conflicts == 0'
    return (pd.DataFrame(in_scope), metadata)

def aggregate_minutes(frame, out, contract):
    f = frame.copy()
    f['date'] = pd.to_datetime(f.date)
    f['hour'] = pd.to_datetime(f.timestamp_local).dt.floor('h')
    f['ec_scaled'] = f._ec.where(f._ec >= 0) / RAW_EC_DIVISOR
    f['water_valid'] = f._water.where(f._water.between(-20, 60))
    valid = f[f.temp_rh_model_valid].copy()
    valid['vpd'] = [0.6108 * math.exp(17.27 * t / (t + 237.3)) * (1 - r / 100) for (t, r) in zip(valid._t, valid._rh)]
    hourly = valid.groupby('hour', sort=True).agg(minute_count=('_t', 'size'), temp_mean=('_t', 'mean'), rh_mean=('_rh', 'mean'), vpd_mean=('vpd', 'mean'), ec_mean=('ec_scaled', 'mean'), ec_count=('ec_scaled', 'count'), ph_mean_raw=('_ph', 'mean'), ph_count=('_ph', 'count'), water_temp_mean=('water_valid', 'mean'), water_count=('water_valid', 'count')).reset_index()
    hourly['stress_hour_valid'] = hourly.minute_count >= MIN_STRESS_MINUTES
    hourly['date'] = hourly.hour.dt.normalize()
    hourly.to_csv(out / '02_environment_hourly_audit.csv', index=False, encoding='utf-8-sig')
    (base_rows, audit_rows) = ([], [])
    for day in pd.date_range(START, END):
        a = f[f.date.eq(day)]
        v = valid[valid.date.eq(day)]
        h = hourly[hourly.date.eq(day)]
        s = h[h.stress_hour_valid]
        n = len(s)
        base = dict(date=day, temp_mean=v._t.mean(), heat_hours_28_ratio=(s.temp_mean > 28).mean() if n else np.nan, heat_degree_hours_28=(s.temp_mean - 28).clip(lower=0).sum() if n else np.nan, vpd_hours_ratio_1_8=(s.vpd_mean > 1.8).mean() if n else np.nan, vpd_degree_hours_1_8=(s.vpd_mean - 1.8).clip(lower=0).sum() if n else np.nan, hot_dry_load_28_1_8=((s.temp_mean - 28).clip(lower=0) * (s.vpd_mean - 1.8).clip(lower=0)).sum() if n else np.nan, water_air_gap_day=np.nan, water_air_gap_day_abs=np.nan, ec_mean=v.ec_scaled.mean(), minute_count=len(v), valid_stress_hour_count=n)
        base_rows.append(base)
        gap_minutes = pd.to_datetime(a.timestamp_local).sort_values().diff().dt.total_seconds().div(60)
        row = dict(date=day, observed_minutes=len(a), temp_rh_valid_minutes=len(v), coverage_ratio=len(a) / 1440, hour_count=len(h), valid_stress_hour_count=n, incomplete_stress_hours=int((~h.stress_hour_valid).sum()), max_gap_between_observed_minutes=gap_minutes.max(), temp_mean_raw=a._t.mean(), temp_min_raw=a._t.min(), temp_max_raw=a._t.max(), rh_mean_raw=a._rh.mean(), rh_min_raw=a._rh.min(), rh_max_raw=a._rh.max(), ec_raw_count=a._ec.count(), ec_raw_mean=a._ec.mean(), ec_raw_min=a._ec.min(), ec_raw_max=a._ec.max(), ec_observed_mS_cm_mean=a.ec_scaled.mean(), EC_formula_mean_including_missing_as_zero=pd.to_numeric(a.EC_source_formula_value).mean(), EC_formula_zero_from_missing_count=int(a.ec_conversion_zero_from_missing.sum()), ph_raw_count=a._ph.count(), ph_raw_mean=a._ph.mean(), ph_raw_min=a._ph.min(), ph_raw_max=a._ph.max(), ph_outside_review_0_14_count=int(a.ph_outside_review_0_14.sum()), water_raw_count=a._water.count(), water_raw_mean=a._water.mean(), water_raw_min=a._water.min(), water_raw_max=a._water.max(), daynight_basis='UNAVAILABLE_NO_CONFIRMED_SCHEDULE', light_status='EXCLUDED_USER_CONFIRMED_SENSOR_FAILURE')
        audit_rows.append(row)
    base = pd.DataFrame(base_rows)[contract.RAW]
    audit = pd.DataFrame(audit_rows)
    base.to_csv(out / '03_environment_daily_base_Cstep.csv', index=False, encoding='utf-8-sig', date_format='%Y-%m-%d')
    audit.to_csv(out / '07_environment_daily_observation_audit.csv', index=False, encoding='utf-8-sig', date_format='%Y-%m-%d')
    return (base, audit, hourly)

def derive_available(base, contract):
    blocked = base.copy()
    blocked['ec_mean'] = np.nan
    derived = contract.derive(blocked)
    derived['ec_mean_observed_preserved'] = base.ec_mean.to_numpy()
    assert derived.ec_penalty.isna().all() and derived.water_air_gap_day.isna().all()
    return derived

def asof_frames(derived, contract):
    dates = pd.DataFrame({'manifest_row_id': range(len(derived)), 'prediction_cutoff': derived.date})
    by_day = contract.join_current(dates, derived)
    by_day['cultivation_day'] = (pd.to_datetime(by_day.prediction_cutoff) - pd.Timestamp(START)).dt.days
    with closing(load_workbook(TARGETS, read_only=True, data_only=True)) as wb:
        it = wb['생육매칭'].iter_rows(values_only=True)
        headers = next(it)
        targets = [dict(zip(headers, r)) for r in it if any((x is not None for x in r))]
    assert len(targets) == 89
    meta = pd.DataFrame([dict(manifest_row_id=i, item_id=r['item_id'], record_id=r['record_id'], prediction_cutoff=r['capture_timestamp'], survey_date=str(r['survey_date'])[:10]) for (i, r) in enumerate(targets)])
    current = meta.merge(contract.join_current(meta, derived), on=['manifest_row_id', 'prediction_cutoff'], validate='one_to_one')
    for df in [by_day, current]:
        df['available_model_environment_count'] = df[ALLOWED].notna().sum(axis=1)
        df['blocked_environment_reason'] = 'EC_TARGET_UNDOCUMENTED;DAYNIGHT_UNDOCUMENTED;IS_EC_SHOCK_FROZEN_NAN'
        source = pd.to_datetime(df.env_daily_source_date)
        cutoff = pd.to_datetime(df.prediction_cutoff).dt.normalize()
        assert (source.isna() | (source == cutoff - pd.Timedelta(days=1))).all()
    return (by_day, current)

def column_audit(contract, by_day, current, out):
    selected = {}
    for mode in ['direct_stacking', 'residual']:
        p = MODEL_ROOT / mode / 'ensemble.json'
        selected[mode] = json.loads(p.read_text(encoding='utf-8'))['features']
    rows = []
    for name in contract.ENV_COLS:
        reason = 'OBSERVED_TEMP_RH_HOURLY_48MIN_CALENDAR_PAST' if name in ALLOWED else 'FROZEN_NAN_NOT_SELECTED' if name == 'is_ec_shock' else 'DAYNIGHT_UNDOCUMENTED' if name.startswith('water_air_gap') else 'EC_TARGET_UNDOCUMENTED_OR_DEPENDENT'
        rows.append(dict(column=name, model_usable=name in ALLOWED, reason=reason, selected_direct_stacking=name in selected['direct_stacking'], selected_residual=name in selected['residual'], finite_calendar_days=int(by_day[name].notna().sum()), total_calendar_days=len(by_day), finite_current89=int(current[name].notna().sum()), current89_min=current[name].min(), current89_max=current[name].max()))
    result = pd.DataFrame(rows)
    result.to_csv(out / '06_Cstep_column_availability.csv', index=False, encoding='utf-8-sig')
    return (result, selected)

def verify_saved(out, source_metadata, base, derived, by_day, current, contract):
    digest = hashlib.sha256()
    with (out / '01_environment_minute_preserved.csv').open(encoding='utf-8-sig', newline='') as f:
        for r in csv.DictReader(f):
            raw_digest_update(digest, r)
    assert digest.hexdigest() == source_metadata['raw_ec_ph_fingerprint']
    saved = pd.read_csv(out / '01_environment_minute_preserved.csv')
    assert len(saved) == source_metadata['scope_rows']
    saved['date'] = pd.to_datetime(saved.date)
    saved['hour'] = pd.to_datetime(saved.timestamp_local).dt.floor('h')
    valid = saved[saved.temp_rh_model_valid].copy()
    temperature = valid.temp_raw.to_numpy(float)
    valid['independent_vpd'] = 0.6108 * np.exp(17.27 * temperature / (temperature + 237.3)) * (1 - valid.rh_raw.to_numpy(float) / 100)
    groups = valid.groupby('hour').agg(t=('temp_raw', 'mean'), v=('independent_vpd', 'mean'), n=('temp_raw', 'size'))
    errors = []
    for r in base.itertuples():
        a = valid[valid.date.eq(r.date)]
        h = groups[(groups.index.normalize() == r.date) & (groups.n >= MIN_STRESS_MINUTES)]
        expected = [a.temp_raw.mean(), (h.t > 28).mean(), (h.t - 28).clip(lower=0).sum(), (h.v > 1.8).mean(), (h.v - 1.8).clip(lower=0).sum(), ((h.t - 28).clip(lower=0) * (h.v - 1.8).clip(lower=0)).sum(), pd.to_numeric(a.ec_raw).where(pd.to_numeric(a.ec_raw) >= 0).mean() / RAW_EC_DIVISOR]
        actual = [r.temp_mean, r.heat_hours_28_ratio, r.heat_degree_hours_28, r.vpd_hours_ratio_1_8, r.vpd_degree_hours_1_8, r.hot_dry_load_28_1_8, r.ec_mean]
        assert np.allclose(actual, expected, atol=1e-10, rtol=1e-10, equal_nan=True)
        errors.extend((abs(x - y) for (x, y) in zip(actual, expected) if math.isfinite(x) and math.isfinite(y)))
    names = contract.ENV_COLS
    blocked = [c for c in names if c not in ALLOWED]
    assert by_day[blocked].isna().all().all() and current[blocked].isna().all().all()
    for cutoff in pd.to_datetime(by_day.prediction_cutoff):
        prefix = derive_available(base[base.date < cutoff].copy(), contract)
        q = pd.DataFrame({'manifest_row_id': [0], 'prediction_cutoff': [cutoff]})
        expected = contract.join_current(q, prefix)[names].to_numpy(float)
        actual = by_day.loc[pd.to_datetime(by_day.prediction_cutoff).eq(cutoff), names].to_numpy(float)
        assert np.allclose(actual, expected, equal_nan=True, atol=1e-10)
    extra = derived.iloc[[-1]].copy()
    extra['date'] = pd.Timestamp(END) + pd.Timedelta(days=10)
    for c in ALLOWED:
        if c in extra:
            extra[c] = 9999.0
    contaminated = pd.concat([derived, extra], ignore_index=True)
    q = pd.DataFrame({'manifest_row_id': range(len(by_day)), 'prediction_cutoff': by_day.prediction_cutoff})
    future = contract.join_current(q, contaminated)[names].to_numpy(float)
    assert np.allclose(future, by_day[names].to_numpy(float), equal_nan=True)
    saved_asof = pd.read_csv(out / '05_Cstep_environment_asof.csv')
    saved_current = pd.read_csv(out / '08_Cstep_environment_current89.csv')
    assert np.allclose(saved_asof[names], by_day[names], equal_nan=True)
    assert np.allclose(saved_current[names], current[names], equal_nan=True)
    assert saved_current.record_id.tolist() == current.record_id.tolist()
    assert not any(('light' in c.lower() or 'ph_raw' == c for c in names))
    return dict(status='PASS', raw_ec_ph_saved_fingerprint_equal=True, independent_daily_checks=len(base) * 7, max_numeric_difference=max(errors, default=0), prefix_cutoff_checks=len(by_day), future_insertion_invariance=True, blocked_features_all_nan=True, source_previous_day_exact=True, saved_current89_rows=len(saved_current), saved_model_values_equal=True, no_interpolation=True, no_ec_target_fabrication=True, light_excluded=True)

def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output-root', type=Path, default=OUTPUT_ROOT)
    args = parser.parse_args()
    begun = timer.monotonic()
    stamp = datetime.now(ZoneInfo('Asia/Seoul')).strftime('%Y%m%d_%H%M%S_%f')
    out = args.output_root / f'run_{stamp}'
    out.mkdir(parents=True, exist_ok=False)
    contract = load_c_contract()
    hashes = {str(p): sha(p) for p in [SOURCE, TARGETS, C_CODE, Path(__file__)]}
    (frame, metadata) = source_minutes(out)
    (base, daily, hourly) = aggregate_minutes(frame, out, contract)
    derived = derive_available(base, contract)
    derived.to_csv(out / '04_environment_daily_components.csv', index=False, encoding='utf-8-sig', date_format='%Y-%m-%d')
    (by_day, current) = asof_frames(derived, contract)
    by_day.to_csv(out / '05_Cstep_environment_asof.csv', index=False, encoding='utf-8-sig')
    current.to_csv(out / '08_Cstep_environment_current89.csv', index=False, encoding='utf-8-sig')
    (columns, selected) = column_audit(contract, by_day, current, out)
    validation = verify_saved(out, metadata, base, derived, by_day, current, contract)
    assert all((sha(p) == digest for (p, digest) in hashes.items())), 'Input contract failed: all((sha(p) == digest for (p, digest) in hashes.items()))'
    save_json(out / 'verification.json', validation)
    reference = dict(path=str(REFERENCE2), exists=REFERENCE2.is_file(), comparison_status='VISUALIZATION_NOT_RUN', caveats=['2작기 원본에는 주야간 열 존재; 4작기는 확인되지 않음', '원 frozen 일별표의 정규화/누적값과 4작기 부분 환경을 그대로 같은 완전 입력으로 비교하지 않음', '다음 시각화는 온습도·EC·pH·수온과 공통 C-step 스트레스 열부터 비교'])
    if REFERENCE2.is_file():
        reference['sha256'] = sha(REFERENCE2)
        with closing(load_workbook(REFERENCE2, read_only=True, data_only=True)) as wb:
            it = wb.worksheets[0].iter_rows(values_only=True)
            hd = list(next(it))
            rs = [r for r in it if r[0] is not None]
            reference.update(rows=len(rs), date_min=str(min((r[0] for r in rs))), date_max=str(max((r[0] for r in rs))), cstep_columns_present=[c for c in contract.ENV_COLS if c in hd])
    save_json(out / '09_2nd_crop_comparison_reference.json', reference)
    report = dict(status='ENVIRONMENT_PROCESSING_COMPLETE_WITH_LIMITATIONS', started_local=stamp, finished_local=str(datetime.now(ZoneInfo('Asia/Seoul'))), elapsed_seconds=timer.monotonic() - begun, run=str(out), source_hashes=hashes, period=[START, END], metadata=metadata, days=len(base), hourly_rows=len(hourly), cstep_environment_columns=contract.ENV_COLS, selected_environment_count={k: len(set(v) & set(contract.ENV_COLS)) for (k, v) in selected.items()}, usable_columns=ALLOWED, blocked_columns=[c for c in contract.ENV_COLS if c not in ALLOWED], current89_usable_count_distribution=current.available_model_environment_count.value_counts().to_dict(), ready_for_unmodified_cstep=False, season4_adapter_must_load_explicit_17column_asof=True, no_model_inference=True, validation=validation, policies=dict(light='USER_CONFIRMED_SENSOR_FAILURE; raw original only; excluded from input and interpretation', ec='raw ec retained; fixed /1000 conversion checked against source EC; no missing-to-zero', ph='raw retained including outside 0..14 review values; not a frozen-model feature', ec_target='undocumented; penalty and dependent variables remain NaN', daynight='undocumented; model daytime water-air difference remains NaN', daily_partial='observed-minute estimates and >=48-minute hourly rule; coverage explicit; no filling'))
    report['outputs'] = {p.name: dict(bytes=p.stat().st_size, sha256=sha(p)) for p in sorted(out.iterdir()) if p.is_file()}
    save_json(out / 'completion.json', report)
    save_json(args.output_root / 'latest_completed.json', dict(run=str(out)))
    print(json.dumps(clean({k: report[k] for k in ['status', 'run', 'elapsed_seconds', 'days', 'hourly_rows', 'metadata', 'current89_usable_count_distribution', 'usable_columns', 'blocked_columns', 'validation']}), ensure_ascii=False, indent=2), flush=True)
if __name__ == '__main__':
    main()
