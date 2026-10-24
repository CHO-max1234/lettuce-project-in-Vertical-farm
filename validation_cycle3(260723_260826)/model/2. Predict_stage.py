"""Verify target crops, masks and identities; separate labels; run the frozen CNN."""
from pathlib import Path as _Path
import sys as _sys
_ROOT = _Path(__file__).resolve().parents[1]
for _folder in (_ROOT, _ROOT / 'pipeline', _ROOT / 'preprocessing'):
    _sys.path.insert(0, str(_folder))
import settings as CYCLE
from runtime import resolved_spec, run_source, moved_tree, moved_path, yolo_device
from pathlib import Path
PROJECT = CYCLE.PROJECT_ROOT
REFERENCE = CYCLE.MAIN_ROOT
MATCHING = CYCLE.MATCHING
TARGET_FILE = CYCLE.TARGET_FILE
SOURCE115 = CYCLE.SOURCE115
CROP_ROOT = CYCLE.CROP_ROOT
CROP_DB = CYCLE.CROP_DB
GROWTH_FILE = CYCLE.GROWTH_FILE
MODEL_BASE = CYCLE.ARTIFACT_ROOT
CNN_MODEL = CYCLE.STAGE_MODEL
A_MODEL = CYCLE.A_MODEL
INPUT_ROOT = CYCLE.INPUT_ROOT
CNN_ROOT = CYCLE.CNN_ROOT
EXPECTED_CNN_SHA = 'dc2b180b471ba47f67e617d4d6aa3f43c08016460dee6cf9cfb8ead226587eb2'
(BATCH_SIZE, ROW_CHUNK, IMAGE_WORKERS) = (16, 256, 4)
import argparse
from collections import Counter
from concurrent.futures import ThreadPoolExecutor
from contextlib import closing
from datetime import datetime
import csv
import hashlib
import importlib.util
import io
import json
import math
import re
import shutil
import sqlite3
import sys
import time
import zlib
from zoneinfo import ZoneInfo
sys.dont_write_bytecode = True
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from PIL import Image
from tqdm import tqdm
IMAGE_NAME = re.compile('^(bed\\d+)_(20\\d{6})_(\\d{6})_(cam[012])(?:_r(?:f)?)?_+([bt]\\d+)(?:_p\\d+)?\\.(?:png|jpg|jpeg)$', re.I)

def sha(path):
    from runtime import digest, resolve_code
    return digest(resolve_code(path))

def stamp():
    return datetime.now(ZoneInfo('Asia/Seoul')).strftime('%Y%m%d_%H%M%S_%f')

def clean(x):
    if isinstance(x, dict):
        return {str(k): clean(v) for (k, v) in x.items()}
    if isinstance(x, (list, tuple)):
        return [clean(v) for v in x]
    if isinstance(x, (datetime, pd.Timestamp, Path)):
        return str(x)
    if isinstance(x, np.generic):
        return clean(x.item())
    if isinstance(x, float) and (not math.isfinite(x)):
        return None
    return x

def js(path, value):
    Path(path).write_text(json.dumps(clean(value), ensure_ascii=False, indent=2, allow_nan=False), encoding='utf-8')

def read_json(path):
    return json.loads(Path(path).read_text(encoding='utf-8-sig'))

def load(name, path):
    sys.path.insert(0, str(REFERENCE))
    spec = resolved_spec(name, path)
    m = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(m)
    return m

def read_sheet(path, sheet):
    with closing(load_workbook(path, read_only=True, data_only=True)) as wb:
        it = wb[sheet].iter_rows(values_only=True)
        h = list(next(it))
        assert len(h) == len(set(h))
        return [dict(zip(h, r)) for r in it if any((x is not None for x in r))]

def write_csv(path, rows, columns=None):
    columns = columns or list(dict.fromkeys((k for r in rows for k in r)))
    with Path(path).open('w', encoding='utf-8-sig', newline='') as f:
        writer = csv.DictWriter(f, fieldnames=columns)
        writer.writeheader()
        writer.writerows(({k: clean(r.get(k)) for k in columns} for r in rows))

def read_csv(path):
    with Path(path).open(encoding='utf-8-sig', newline='') as f:
        return list(csv.DictReader(f))

def protected_sources():
    files = [TARGET_FILE, SOURCE115, GROWTH_FILE, CNN_MODEL, A_MODEL / 'selected_model_bundle.json', A_MODEL / 'MODEL_SELECTION_FROZEN.json', REFERENCE / 'model/1. Predict_stage.py', REFERENCE / 'model/2. Predict_weight.py', REFERENCE / 'pipeline/crop_integrity.py', REFERENCE / 'preprocessing/mask_record.py', Path(__file__)]
    bundle = read_json(A_MODEL / 'selected_model_bundle.json')
    for fold in bundle['folds']:
        for rec in fold['files']:
            p = A_MODEL / rec['path']
            assert sha(p) == rec['sha256'], 'A model changed: ' + str(p)
            files.append(p)
    result = {str(p): sha(p) for p in files}
    assert result[str(CNN_MODEL)] == EXPECTED_CNN_SHA
    return result

def check_destructive_identity():
    with closing(load_workbook(GROWTH_FILE, read_only=True, data_only=True)) as wb:
        data = list(wb['26.3회차 파괴'].iter_rows(values_only=True))[2:]
    raw = [r for r in data if r[0] and isinstance(r[1], (int, float)) and isinstance(r[2], (int, float)) and isinstance(r[7], (int, float))]
    keys = [(int(r[1]), int(r[2])) for r in raw]
    assert len(raw) == 128 and len(set(keys)) == 128, 'Input contract failed: len(raw) == 128 and len(set(keys)) == 128'
    return len(raw)

def prepare():
    begun = time.monotonic()
    protected = protected_sources()
    targets = read_sheet(TARGET_FILE, '생육매칭')
    hist = read_sheet(TARGET_FILE, '과거Feature')
    assert len(targets) == 89 and len(hist) == 1313
    assert len({r['record_id'] for r in hist}) == 1313
    target_by_item = {r['item_id']: r for r in targets}
    history_by_id = {r['record_id']: r for r in hist}
    assert len(target_by_item) == 89
    for t in targets:
        assert t['record_id'] in history_by_id and t['view'] == 'top' and (t['A_C_main_view_eligible'] is True)
        assert t['crop_use_decision'] == 'USER_APPROVED_TOP_89'
        assert t['physical_slot'] == ('t' + str(int(t['sample_number'])) if t['sample_number'] <= 12 else 'b' + str(int(t['sample_number']) - 12))
        assert int(t['Stage_human']) in [0, 1, 2]
    assert check_destructive_identity() == 128
    a = load('s4_a_input_contract', REFERENCE / 'model/2. Predict_weight.py')
    (bundle, core) = a.checked_bundle(A_MODEL)
    gate = load('s4_mask_integrity', REFERENCE / 'pipeline/crop_integrity.py')
    packer = load('s4_explicit_mask', REFERENCE / 'preprocessing/mask_record.py')
    assert not Path(str(CROP_DB) + '-wal').exists() or Path(str(CROP_DB) + '-wal').stat().st_size == 0
    db_before = (CROP_DB.stat().st_size, CROP_DB.stat().st_mtime_ns)
    ordered = [history_by_id[t['record_id']] for t in targets] + [r for r in hist if r['role'] == 'HISTORY_ONLY']
    assert len(ordered) == 1313 and len({r['record_id'] for r in ordered}) == 1313
    model_rows = []
    audits = []
    gates = []
    with closing(sqlite3.connect(CROP_DB.as_uri() + '?mode=ro&immutable=1', uri=True)) as db:
        for (index, h) in enumerate(tqdm(ordered, desc='Processing', unit='image', mininterval=5)):
            t = target_by_item[h['item_id']]
            assert pd.Timestamp(h['capture_timestamp']) <= pd.Timestamp(h['target_cutoff']) == pd.Timestamp(t['capture_timestamp'])
            blob = db.execute('SELECT payload,mask FROM slots WHERE record_id=?', (h['record_id'],)).fetchone()
            assert blob and blob[1] is not None, 'Input contract failed: blob and blob[1] is not None'
            r = moved_tree(json.loads(zlib.decompress(blob[0])))
            for k in ['record_id', 'camera', 'capture_date', 'capture_time', 'crop_sha256']:
                assert str(h[k]) == str(r[k]), f"{h['record_id']} metadata mismatch {k}"
            assert r['bed_id'] == t['bed_id'] and r['physical_slot'] == t['physical_slot']
            assert r['target_available'] is True and r['selection_status'] == 'SELECTED'
            mask = np.unpackbits(np.frombuffer(zlib.decompress(blob[1]), np.uint8), bitorder='little', count=math.prod(r['target_mask_shape'])).reshape(r['target_mask_shape']).astype(bool)
            assert hashlib.sha256(mask.tobytes()).hexdigest() == r['target_mask_pixels_sha256']
            manifest_mask = packer.unpack(r) > 0
            assert np.array_equal(mask, manifest_mask) and int(mask.sum()) == int(r['mask_pixels'])
            path = Path(moved_path(h['crop_path'])).resolve()
            assert path.is_relative_to(CROP_ROOT.resolve())
            raw = path.read_bytes()
            assert hashlib.sha256(raw).hexdigest() == h['crop_sha256'] == r['crop_sha256']
            m = IMAGE_NAME.fullmatch(path.name)
            assert m and m.groups() == tuple((str(r[k]) for k in ['bed_id', 'capture_date', 'capture_time', 'camera', 'physical_slot']))
            with Image.open(io.BytesIO(raw)) as im:
                rgb = np.asarray(im.convert('RGB'))
                (width, height) = im.size
            assert mask.shape == (height, width) and (not np.any(rgb[~mask]))
            assert hashlib.sha256(np.ascontiguousarray(rgb[:, :, ::-1]).tobytes()).hexdigest() == r['crop_pixels_sha256']
            box = r['crop_source_box']
            assert [box[2] - box[0], box[3] - box[1]] == [width, height]
            assert all((isinstance(h[c], (int, float)) and math.isfinite(h[c]) for c in core))
            if h['role'] == 'CURRENT':
                assert all((math.isclose(float(t[c]), float(h[c]), rel_tol=1e-12, abs_tol=1e-12) for c in core))
                assert int(t['mask_pixels']) == int(r['mask_pixels'])
            packed = dict(r, has_crop='True', crop_path=str(path))
            (allowed, status, reason) = gate.judge(packed, straight_ratio=0.8, minimum=20, main_share=0.9)
            row = {c: h[c] for c in core}
            row.update(record_id=r['record_id'], season_id='4작기', bed_id=r['bed_id'], physical_slot=r['physical_slot'], camera=r['camera'], capture_date=r['capture_date'], capture_time=r['capture_time'], capture_timestamp=h['capture_timestamp'], crop_path=str(path), crop_sha256=h['crop_sha256'], mask_pixels=r['mask_pixels'], crop_source_box=json.dumps(box), selected_camera=r['camera'], selected_record_id=r['record_id'], slot_namespace_version='physical_12col_v1', source_slot_namespace_version=r['slot_namespace_version'], slot_mapping_basis=t['slot_mapping_basis'], feature_status='OK', 분석가능=allowed, front_height_cm=None, volume_proxy=None, scale_verified=False, identity_confirmed=False, source_camera=h['source_camera'], native_width=h['native_width'], native_height=h['native_height'], item_id=t['item_id'], row_role=h['role'], target_record_id=t['record_id'], target_cutoff=h['target_cutoff'], source_history_row=hist.index(h) + 2, crop_integrity_status=status, crop_integrity_reason=reason, crop_integrity_version=gate.VERSION, scale_qc_v2_2=h['scale_qc_v2_2'], body_qc_v2_2=h['body_qc_v2_2'])
            assert not set(row) & {'Stage_human', 'fresh_weight_g', 'human_label', 'survey_weight_g'}
            model_rows.append(row)
            audits.append(dict(embedding_row_index=index, record_id=r['record_id'], item_id=t['item_id'], role=h['role'], hash_verified=True, mask_equal=True, background_zero=True, core25_missing=0, source_namespace=r['slot_namespace_version'], model_namespace='physical_12col_v1', model_eligible=allowed, integrity_status=status, integrity_reason=reason))
            gates.append(dict(record_id=r['record_id'], model_eligible=allowed, integrity_status=status, integrity_reason=reason))
    assert (CROP_DB.stat().st_size, CROP_DB.stat().st_mtime_ns) == db_before
    assert all((sha(p) == v for (p, v) in protected.items())), 'Input contract failed: all((sha(p) == v for (p, v) in protected.items()))'
    out = INPUT_ROOT / ('run_' + stamp())
    out.mkdir(parents=True, exist_ok=False)
    shutil.copy2(__file__, out / '실행코드_보존.py')
    for (name, rows) in [('model_input_1313.csv', model_rows), ('evaluation_targets89.csv', targets), ('history_source_mapping1313.csv', hist), ('crop_integrity.csv', gates), ('input_verification.csv', audits)]:
        write_csv(out / name, rows)
    report = dict(status='SEASON4_INPUT_PREPARED', run=str(out), source_hashes=protected, source_db_snapshot=dict(path=str(CROP_DB), bytes=db_before[0], mtime_ns=db_before[1]), total_rows=1313, current_rows=89, history_only_rows=1224, core_numeric_features=core, crop_gate_thresholds=dict(straight_ratio=0.8, minimum=20, main_share=0.9), gate_all=dict(Counter((r['crop_integrity_status'] for r in model_rows))), gate_current89=dict(Counter((r['crop_integrity_status'] for r in model_rows[:89]))), gate_current_reasons=dict(Counter((r['crop_integrity_reason'] for r in model_rows[:89] if r['crop_integrity_reason']))), original_identity_review_preserved=True, labels_in_model_input=False, history_has_no_duplicated_truth=True, current_rows_first_in_source_order=True, all_core25_current_history_agree=True, model_training=False, elapsed_seconds=time.monotonic() - begun)
    report['outputs'] = {p.name: sha(p) for p in out.iterdir() if p.is_file()}
    js(out / 'completion.json', report)
    js(INPUT_ROOT / 'latest_completed.json', dict(run=str(out)))
    print(json.dumps(clean({k: report[k] for k in ['status', 'run', 'total_rows', 'current_rows', 'gate_all', 'gate_current89', 'gate_current_reasons', 'elapsed_seconds']}), ensure_ascii=False, indent=2), flush=True)
    return out

def validate_prepared(prepared):
    report = read_json(prepared / 'completion.json')
    assert report['status'] == 'SEASON4_INPUT_PREPARED'
    for (n, v) in report['outputs'].items():
        assert sha(prepared / n) == v
    for (p, v) in report['source_hashes'].items():
        assert sha(p) == v
    snap = report['source_db_snapshot']
    p = Path(snap['path'])
    assert (p.stat().st_size, p.stat().st_mtime_ns) == (snap['bytes'], snap['mtime_ns'])
    return report

def run_cnn(prepared):
    report = validate_prepared(prepared)
    rows = read_csv(prepared / 'model_input_1313.csv')
    mod = load('s4_frozen_single_cnn', REFERENCE / 'model/1. Predict_stage.py')
    mod.IMAGE_NAME = IMAGE_NAME
    mod.CROP_ROOT = CROP_ROOT
    mod.MODEL_PT = CNN_MODEL
    mod.DEVICE = __import__('runtime').torch_device()
    mod.BATCH_SIZE = BATCH_SIZE
    mod.IMAGE_WORKERS = IMAGE_WORKERS
    (model, transform, device, size, model_hash, signature) = mod.load_model()
    out = CNN_ROOT / ('run_' + stamp())
    out.mkdir(parents=True, exist_ok=False)
    js(out / 'started.json', dict(status='RUNNING', prepared=str(prepared), input_sha256=sha(prepared / 'model_input_1313.csv')))
    shutil.copy2(__file__, out / '실행코드_보존.py')
    headers = list(rows[0])
    counters = Counter()
    begun = time.monotonic()
    cols = headers + mod.CNN_COLUMNS
    with closing(sqlite3.connect(out / 'cnn_predictions_cache.sqlite3')) as db, (out / 'crop_features_CNN.csv').open('w', encoding='utf-8-sig', newline='') as f, ThreadPoolExecutor(max_workers=IMAGE_WORKERS) as pool:
        db.execute('CREATE TABLE predictions(signature TEXT,image_sha256 TEXT,probabilities TEXT,PRIMARY KEY(signature,image_sha256))')
        writer = csv.DictWriter(f, fieldnames=cols)
        writer.writeheader()
        for start in tqdm(range(0, len(rows), ROW_CHUNK), desc='Processing', unit='chunk', mininterval=5):
            part = rows[start:start + ROW_CHUNK]
            records = [(start + j + 2, [r.get(c) for c in headers]) for (j, r) in enumerate(part)]
            result = mod.infer_chunk(records, headers, model, transform, device, size, model_hash, signature, db, pool, counters)
            for (r, p) in zip(part, result):
                assert p['cnn_status'] != 'PENDING'
                writer.writerow(dict(r, **p))
                counters['status_' + p['cnn_status']] += 1
    del model
    import torch
    torch.cuda.empty_cache()
    saved = read_csv(out / 'crop_features_CNN.csv')
    assert len(saved) == 1313
    for (r, s) in zip(rows, saved):
        assert all((str(r[k]) == str(s[k]) for k in headers))
    assert all((sha(p) == v for (p, v) in report['source_hashes'].items()))
    done = dict(status='COMPLETE' if counters['status_OK'] == 1313 else 'COMPLETE_WITH_ROW_ERRORS', prepared=str(prepared), rows=1313, current_rows=89, counts=dict(counters), current_status_counts=dict(Counter((r['cnn_status'] for r in saved[:89]))), model_sha256=model_hash, reference_code_sha256=sha(REFERENCE / 'model/1. Predict_stage.py'), input_sha256=sha(prepared / 'model_input_1313.csv'), output_sha256=sha(out / 'crop_features_CNN.csv'), frozen_preprocessing=dict(image_size=384, resize='PIL_BILINEAR', mean=[0.485, 0.456, 0.406], std=[0.229, 0.224, 0.225], amp=True), filename_adapter=IMAGE_NAME.pattern, model_training=False, labels_used=False, elapsed_seconds=time.monotonic() - begun)
    js(out / 'completion.json', done)
    js(CNN_ROOT / 'latest_completed.json', dict(run=str(out), prepared=str(prepared)))
    print(json.dumps(clean(dict(run=str(out), **done)), ensure_ascii=False, indent=2), flush=True)
    return out
if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--prepare-only', action='store_true')
    args = parser.parse_args()
    result = prepare()
    if not args.prepare_only:
        run_cnn(result)
