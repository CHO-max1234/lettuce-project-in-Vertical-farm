"""Portable paths, shared frozen-model loading and value-only workbook output."""
from pathlib import Path, PureWindowsPath
import sys
import importlib.util
import importlib.abc
import json
import hashlib
import runpy
import settings as CYCLE
LOCAL_CODES = {'261005_10_윗면_CAM01_전체_ROI_BOX.py': 'preprocessing/1. Segment_beds_top.py', '261005_24_CAM2_비우고_새모델재생성.py': 'preprocessing/1. Segment_beds_front.py', '261005_2_81장_기존bedseg_비교테스트.py': 'preprocessing/bed_detection.py', '261005_7_윗면_추론설정_회복진단.py': 'preprocessing/top_bed_policy.py', '261005_15_정면모델_샘플검증과공통함수.py': 'preprocessing/front_bed_policy.py', '261005_16_정면CAM2_전체ROI_재개가능.py': 'preprocessing/front_bed_storage.py', '261009_01_밝기보정_3작기방식.py': 'preprocessing/3. Normalize_brightness.py', '261009_00_SCALE_한번에.py': 'preprocessing/4. Estimate_scale.py', '261009_04_SCALE_4작기.py': 'preprocessing/scale_estimation.py', '261009_05_SCALE_스트리밍엑셀.py': 'preprocessing/scale_writer.py', '261009_06_SCALE_결과검증.py': 'preprocessing/scale_validation.py', '261010_22_CAM2개선_원스톱.py': 'preprocessing/5. Refine_slot_geometry.py', '261009_10_manifest_병렬출력.py': 'preprocessing/layout_manifest.py', '261009_cam2_geometry_ver2.py': 'preprocessing/front_geometry_base.py', '261010_16_CAM2전수_병합과검증.py': 'preprocessing/layout_builder.py', '261010_17_bbox_최신자리배치입력.py': 'preprocessing/layout_reader.py', '261010_cam2_geometry_ver3.py': 'preprocessing/front_geometry.py', '261010_cam2_ver3_구도기준.json': 'preprocessing/front_geometry_calibration.json', '261011_자리배치_기하계약과표시.py': 'preprocessing/layout_geometry.py', '261011_자리배치_입출력공통.py': 'preprocessing/layout_records.py', '261009_cam01_geometry_ver2.py': 'preprocessing/top_geometry.py', '261011_01_bbox와_crop_표본.py': 'preprocessing/plant_selection.py', '261011_1_bbox_to_masked_crop_전체.py': 'preprocessing/crop_runtime.py', '261011_2_프로세스_병렬실행.py': 'preprocessing/6. Extract_plant_crops.py', '261011_01_전체_Feature추출.py': 'preprocessing/7. Extract_features.py', '261013_16_v2_2_전체적용과표본비교.py': 'preprocessing/8. Apply_feature_calibration.py', '261014_22_115개체_crop_overlay와_과거필요량.py': 'preprocessing/survey_matching.py', '261014_25_최종89_115_Feature준비.py': 'preprocessing/9. Build_survey_inputs.py', '261014_28_115개체_생육QC.py': 'preprocessing/10. Assess_growth.py', 'cam01_selection_policy.py': 'preprocessing/cam01_selection_policy.py', 'camera_selection.py': 'preprocessing/camera_selection.py', 'feature_calculation.py': 'preprocessing/feature_calculation.py', 'workbook_stream.py': 'preprocessing/workbook_stream.py', '261015_01_환경확인과_가공.py': 'model/1. Prepare_environment.py', '261015_01_입력연결과_CNN.py': 'model/2. Predict_stage.py', '261015_04_Astep_실행과_평가.py': 'model/3. Predict_weight.py', '261015_05_Cstep_실행과_A비교.py': 'model/4. Fuse_growth_history.py', '261015_02_정면26_보조CNN.py': 'model/5. Evaluate_front_stage.py'}
SHARED_CODES = {'0. CNN Softlabel and merge.py': ['main', 'model/1. Predict_stage.py'], '1. A_step embedding and regression.py': ['main', 'model/2. Predict_weight.py'], '2. C-step regression.py': ['main', 'model/3. Fuse_growth_history.py'], 'crop_integrity.py': ['main', 'pipeline/crop_integrity.py'], 'mask_record.py': ['main', 'preprocessing/mask_record.py'], '3. 밝기보정_3작기.py': ['third', 'preprocessing/3. Normalize_brightness.py'], '2. scale_map_3작기.py': ['third', 'preprocessing/4. Estimate_scale.py'], '261001_4. 자리배치_bbox_final_whole.py': ['third', 'preprocessing/5. Build_plant_bboxes.py'], '261011_bbox_후보와본체기록.py': ['third', 'preprocessing/bbox_records.py'], '261001_bbox_seg_selector_ver2.2.py': ['third', 'preprocessing/mask_selection.py'], '261003_Cstep_실증.py': ['third', 'model/3. Fuse_growth_history.py'], '261002_2_feature_to_growth_qc.py': ['third', 'preprocessing/9. Assess_growth.py']}
LEGACY_ENVIRONMENT_COLUMNS = ['effective_gdd_cum', 'effective_gdd_day', 'gdd_loss_ratio', 'daily_env_penalty', 'past_7d_env_penalty_mean', 'heat_hours_28_ratio', 'heat_degree_hours_28', 'vpd_hours_ratio_1_8', 'vpd_degree_hours_1_8', 'hot_dry_load_28_1_8', 'past_7d_hot_dry_load_sum', 'ec_penalty', 'water_air_gap_day', 'water_air_gap_day_abs', 'is_ec_shock', 'past_7d_effective_gdd_sum', 'past_7d_gdd_loss_ratio_mean']

def moved_path(value):
    if not isinstance(value, (str, Path, PureWindowsPath)) or not str(value):
        return value
    path = PureWindowsPath(str(value))
    aliases = [(CYCLE.LEGACY_SOURCE_ROOT, CYCLE.SOURCE_ROOT), (CYCLE.LEGACY_RAW_ROOT + '\\RGB_윗면\\0. 원본\\4작기', CYCLE.RAW_TOP_ROOT), (CYCLE.LEGACY_RAW_ROOT + '\\RGB_정면\\0. 원본\\4작기', CYCLE.RAW_FRONT_ROOT)]
    for (old, new) in aliases:
        parts = PureWindowsPath(old).parts
        if tuple((p.casefold() for p in path.parts[:len(parts)])) == tuple((p.casefold() for p in parts)):
            return str(new.joinpath(*path.parts[len(parts):]))
    return value

def moved_tree(value):
    if isinstance(value, dict):
        return {moved_path(k): moved_tree(v) for (k, v) in value.items()}
    if isinstance(value, list):
        return [moved_tree(v) for v in value]
    if isinstance(value, tuple):
        return tuple((moved_tree(v) for v in value))
    return moved_path(value)

def legacy_identity_path(value):
    path = Path(value).resolve()
    for base in (CYCLE.PREPROCESS_OUTPUT, CYCLE.SOURCE_ROOT):
        if path.is_relative_to(base):
            return str(PureWindowsPath(CYCLE.LEGACY_SOURCE_ROOT).joinpath(*path.relative_to(base).parts))
    return str(value)

def digest(path):
    h = hashlib.sha256()
    with Path(path).open('rb') as f:
        for block in iter(lambda : f.read(4 * 1024 * 1024), b''):
            h.update(block)
    return h.hexdigest()

def yolo_device():
    import torch
    return '0' if CYCLE.DEVICE == 'cuda' or (CYCLE.DEVICE == 'auto' and torch.cuda.is_available()) else 'cpu'

def torch_device():
    import torch
    return 'cuda' if CYCLE.DEVICE == 'auto' and torch.cuda.is_available() else 'cpu' if CYCLE.DEVICE == 'auto' else CYCLE.DEVICE

def resolve_code(path):
    path = Path(path)
    name = path.name
    if name in LOCAL_CODES:
        return CYCLE.PROJECT_ROOT / LOCAL_CODES[name]
    if name in LOCAL_CODES.values():
        return CYCLE.PROJECT_ROOT / name
    if name in SHARED_CODES:
        (group, relative) = SHARED_CODES[name]
        return (CYCLE.MAIN_ROOT if group == 'main' else CYCLE.THIRD_ROOT) / relative
    return path

class SharedLoader(importlib.abc.Loader):

    def __init__(self, loader, path):
        (self.loader, self.path) = (loader, path)

    def create_module(self, spec):
        return None

    def exec_module(self, module):
        sys.modules[module.__name__] = module
        self.loader.exec_module(module)
        if self.path == CYCLE.MAIN_ROOT / 'model/3. Fuse_growth_history.py':
            module.ENV_COLS = list(LEGACY_ENVIRONMENT_COLUMNS)
            module.load_models = load_legacy_components

def resolved_spec(name, path):
    path = resolve_code(path)
    for folder in (CYCLE.PROJECT_ROOT / 'preprocessing', CYCLE.MAIN_ROOT, CYCLE.MAIN_PIPELINE, CYCLE.MAIN_PREPROCESS, CYCLE.MAIN_PREPROCESS / '_support/cam01', CYCLE.THIRD_ROOT / 'preprocessing'):
        if str(folder) not in sys.path:
            sys.path.append(str(folder))
    spec = importlib.util.spec_from_file_location(name, path)
    if spec is None or spec.loader is None:
        raise ImportError(str(path))
    spec.loader = SharedLoader(spec.loader, path)
    return spec

def load_module(name, path):
    spec = resolved_spec(name, path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module

def load_shared(name, relative):
    return load_module(name, CYCLE.MAIN_ROOT / relative)

def run_source(path, **kwargs):
    return runpy.run_path(str(resolve_code(path)), **kwargs)

def load_legacy_components():
    """Read the unchanged final C models for this missing-environment protocol."""
    from catboost import CatBoostRegressor
    (result, hashes) = ({}, {})
    for mode in ('direct_stacking', 'residual'):
        folder = CYCLE.C_MODELS / mode
        path = folder / 'ensemble.json'
        metadata = json.loads(path.read_text(encoding='utf-8-sig'))
        if metadata['mode'] != mode or not metadata['seeds']:
            raise ValueError('Invalid frozen ensemble: ' + str(path))
        estimators = []
        for item in metadata['seeds']:
            model_path = folder / f"seed_{item['seed']}.cbm"
            actual = digest(model_path)
            if actual != item['model_sha256']:
                raise ValueError('Frozen C model changed: ' + str(model_path))
            model = CatBoostRegressor()
            model.load_model(str(model_path))
            if model.feature_names_ != metadata['features']:
                raise ValueError('Frozen feature order differs: ' + str(model_path))
            estimators.append(model)
            hashes[str(model_path)] = actual
        result[mode] = (metadata['features'], estimators)
        hashes[str(path)] = digest(path)
    if result['direct_stacking'][0] != result['residual'][0] + ['pred_weight_cnn']:
        raise ValueError('Stacking/residual feature order differs')
    for path in (CYCLE.PCA_FILE, CYCLE.BACKBONE_FILE, CYCLE.A_MODEL / 'selected_model_bundle.json'):
        hashes[str(path)] = digest(path)
    return (result, hashes)

def write_workbooks(specifications):
    """Write value-only XLSX without a Codex or Node runtime."""
    from datetime import datetime
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    date_columns = set()
    for specification in specifications:
        path = Path(specification['path'])
        path.parent.mkdir(parents=True, exist_ok=True)
        workbook = Workbook(write_only=True)
        for sheet in specification['sheets']:
            worksheet = workbook.create_sheet(sheet['name'])
            worksheet.freeze_panes = 'C2'
            worksheet.sheet_view.showGridLines = False
            headers = []
            for (column, value) in enumerate(sheet['columns'], 1):
                cell = WriteOnlyCell(worksheet, value)
                cell.font = Font(bold=True, color='FFFFFF')
                cell.fill = PatternFill('solid', fgColor='274D67')
                headers.append(cell)
                worksheet.column_dimensions[get_column_letter(column)].width = 24
            worksheet.append(headers)
            for row in sheet['rows']:
                cells = []
                for (column, value) in zip(sheet['columns'], row):
                    if column in date_columns | set(sheet.get('date_columns', [])) and value:
                        if isinstance(value, (int, float)):
                            from openpyxl.utils.datetime import from_excel
                            value = from_excel(value)
                        elif isinstance(value, str):
                            value = datetime.strptime(value, '%Y%m%d') if len(value) == 8 and value.isdigit() else datetime.fromisoformat(value)
                    cell = WriteOnlyCell(worksheet, value)
                    if isinstance(value, str):
                        cell.data_type = 's'
                    elif isinstance(value, datetime):
                        cell.number_format = 'yyyy-mm-dd'
                    cells.append(cell)
                worksheet.append(cells)
            worksheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(sheet['rows']) + 1}"
        temporary = path.with_name(path.stem + '.writing.xlsx')
        workbook.save(temporary)
        temporary.replace(path)
