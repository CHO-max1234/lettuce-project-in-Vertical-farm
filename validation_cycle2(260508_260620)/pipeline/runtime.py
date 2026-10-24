"""Path resolution, workbook output and frozen external-model adapters."""
from pathlib import Path, PureWindowsPath
import sys
PROJECT_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT_ROOT))
import settings as CYCLE
import importlib.util
import json
import hashlib
STORAGE_ROOT = CYCLE.DATA_ROOT
ROI_ROOT = CYCLE.PREPROCESS_OUTPUT
CROP_FOLDER = CYCLE.CROP_FOLDER
OLD_CROP_FOLDER = '5. BBOX to crop'
OLD_STORAGE_ROOT = PureWindowsPath(CYCLE.LEGACY_STORAGE_ROOT)

def moved_path(value):
    """Resolve saved paths in memory without rewriting source records."""
    if not isinstance(value, (str, Path, PureWindowsPath)) or not str(value):
        return value
    path = PureWindowsPath(str(value))
    for root in (CYCLE.LEGACY_STORAGE_ROOT, CYCLE.RELOCATED_STORAGE_ROOT):
        parts = PureWindowsPath(root).parts
        if tuple((p.casefold() for p in path.parts[:len(parts)])) != tuple((p.casefold() for p in parts)):
            continue
        suffix = path.parts[len(parts):]
        if len(suffix) >= 2 and suffix[0] == '3작기 실증 적용' and (suffix[1] in {'2. ROI_BOX', '2. 전처리'}):
            rest = suffix[2:]
            if rest and rest[0] == OLD_CROP_FOLDER:
                rest = (CROP_FOLDER,) + rest[1:]
            return str(CYCLE.SOURCE_PREPROCESS_ROOT.joinpath(*rest))
    return value

def moved_tree(value):
    if isinstance(value, dict):
        return {moved_path(k): moved_tree(v) for (k, v) in value.items()}
    if isinstance(value, list):
        return [moved_tree(v) for v in value]
    if isinstance(value, tuple):
        return tuple((moved_tree(v) for v in value))
    return moved_path(value)

def legacy_identity_path(value):
    """Preserve historical path-derived frame IDs when a dataset moves."""
    path = PureWindowsPath(str(value))
    for candidate in (CYCLE.PREPROCESS_OUTPUT, CYCLE.SOURCE_PREPROCESS_ROOT):
        parts = PureWindowsPath(candidate).parts
        if tuple((p.casefold() for p in path.parts[:len(parts)])) == tuple((p.casefold() for p in parts)):
            suffix = path.parts[len(parts):]
            if suffix and suffix[0] == CROP_FOLDER:
                suffix = (OLD_CROP_FOLDER,) + suffix[1:]
            return str(OLD_STORAGE_ROOT.joinpath('3작기 실증 적용', '2. ROI_BOX', *suffix))
    return str(value)

def digest(path):
    result = hashlib.sha256()
    with Path(path).open('rb') as stream:
        for block in iter(lambda : stream.read(4 * 1024 * 1024), b''):
            result.update(block)
    return result.hexdigest()

def yolo_device():
    """Resolve the configured torch device to the YOLO device spelling."""
    if CYCLE.DEVICE == 'auto':
        import torch
        return '0' if torch.cuda.is_available() else 'cpu'
    return '0' if CYCLE.DEVICE == 'cuda' else CYCLE.DEVICE

def load_module(name, path):
    spec = importlib.util.spec_from_file_location(name, Path(path))
    if spec is None or spec.loader is None:
        raise ImportError(str(path))
    module = importlib.util.module_from_spec(spec)
    sys.modules[name] = module
    spec.loader.exec_module(module)
    return module
LEGACY_ENVIRONMENT_COLUMNS = ['effective_gdd_cum', 'effective_gdd_day', 'gdd_loss_ratio', 'daily_env_penalty', 'past_7d_env_penalty_mean', 'heat_hours_28_ratio', 'heat_degree_hours_28', 'vpd_hours_ratio_1_8', 'vpd_degree_hours_1_8', 'hot_dry_load_28_1_8', 'past_7d_hot_dry_load_sum', 'ec_penalty', 'water_air_gap_day', 'water_air_gap_day_abs', 'is_ec_shock', 'past_7d_effective_gdd_sum', 'past_7d_gdd_loss_ratio_mean']

def load_legacy_components():
    """Read the unchanged final C models for this missing-environment protocol."""
    from catboost import CatBoostRegressor
    (result, hashes) = ({}, {})
    for mode in ('direct_stacking', 'residual'):
        folder = CYCLE.C_MODELS / mode
        path = folder / 'ensemble.json'
        metadata = json.loads(path.read_text(encoding='utf-8-sig'))
        if metadata['mode'] != mode or not metadata['seeds']:
            raise ValueError('Invalid frozen ensemble: ' + str(path))
        estimators = []
        for item in metadata['seeds']:
            model_path = folder / f"seed_{item['seed']}.cbm"
            actual = digest(model_path)
            if actual != item['model_sha256']:
                raise ValueError('Frozen C model changed: ' + str(model_path))
            model = CatBoostRegressor()
            model.load_model(str(model_path))
            if model.feature_names_ != metadata['features']:
                raise ValueError('Frozen feature order differs: ' + str(model_path))
            estimators.append(model)
            hashes[str(model_path)] = actual
        result[mode] = (metadata['features'], estimators)
        hashes[str(path)] = digest(path)
    if result['direct_stacking'][0] != result['residual'][0] + ['pred_weight_cnn']:
        raise ValueError('Stacking/residual feature order differs')
    for path in (CYCLE.PCA_FILE, CYCLE.BACKBONE_FILE, CYCLE.A_MODEL / 'selected_model_bundle.json'):
        hashes[str(path)] = digest(path)
    return (result, hashes)

def missing_environment(frame, daily):
    """Keep every legacy environment field unavailable for this external cycle."""
    import numpy as np
    import pandas as pd
    if len(daily):
        raise ValueError('This protocol requires excluded environmental observations')
    rows = []
    for row in frame.itertuples():
        values = {name: np.nan for name in LEGACY_ENVIRONMENT_COLUMNS}
        values.update(manifest_row_id=row.manifest_row_id, prediction_cutoff=row.prediction_cutoff, env_daily_source_date=pd.NaT, env_past_7d_observed_days=0, env_missing_reason='EXTERNAL_ENVIRONMENT_EXCLUDED', env_daily_join_source='missing', env_lag_join_source='missing')
        rows.append(values)
    return pd.DataFrame(rows)

def load_shared(name, relative):
    """Use shared code with this cycle's paths and explicit frozen contract."""
    for path in (CYCLE.MAIN_ROOT, CYCLE.MAIN_PIPELINE, CYCLE.MAIN_PREPROCESS, CYCLE.MAIN_PREPROCESS / '_support' / 'cam01'):
        sys.path.insert(0, str(path))
    module = load_module(name, CYCLE.MAIN_ROOT / relative)
    if relative == 'model/3. Fuse_growth_history.py':
        module.ENV_COLS = list(LEGACY_ENVIRONMENT_COLUMNS)
        module.load_models = load_legacy_components
        module.join_current = missing_environment
    return module

def write_workbooks(specifications):
    """Write value-only XLSX without a Codex or Node runtime."""
    from datetime import datetime
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    date_columns = {'survey_date', 'image_date', 'qc_date', 'qc_window_start_date', 'qc_window_end_date'}
    for specification in specifications:
        path = Path(specification['path'])
        path.parent.mkdir(parents=True, exist_ok=True)
        workbook = Workbook(write_only=True)
        for sheet in specification['sheets']:
            worksheet = workbook.create_sheet(sheet['name'])
            worksheet.freeze_panes = 'C2'
            worksheet.sheet_view.showGridLines = False
            headers = []
            for (column, value) in enumerate(sheet['columns'], 1):
                cell = WriteOnlyCell(worksheet, value)
                cell.font = Font(bold=True, color='FFFFFF')
                cell.fill = PatternFill('solid', fgColor='274D67')
                headers.append(cell)
                worksheet.column_dimensions[get_column_letter(column)].width = 24
            worksheet.append(headers)
            for row in sheet['rows']:
                cells = []
                for (column, value) in zip(sheet['columns'], row):
                    if column in date_columns and value:
                        value = datetime.fromisoformat(str(value))
                    cell = WriteOnlyCell(worksheet, value)
                    if isinstance(value, str):
                        cell.data_type = 's'
                    elif isinstance(value, datetime):
                        cell.number_format = 'yyyy-mm-dd'
                    cells.append(cell)
                worksheet.append(cells)
            worksheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(sheet['rows']) + 1}"
        temporary = path.with_name(path.stem + '.writing.xlsx')
        workbook.save(temporary)
        temporary.replace(path)
