"""Prepare reviewed cycle-3 targets and apply the frozen main CNN/A models. Human labels and measured weights are used only for evaluation."""
import sys as _sys
from pathlib import Path as _Path
_CYCLE_ROOT = _Path(__file__).resolve().parents[1]
if str(_CYCLE_ROOT) not in _sys.path:
    _sys.path.insert(0, str(_CYCLE_ROOT))
_sys.path.insert(0, str(_CYCLE_ROOT / 'pipeline'))
import settings as CYCLE
ROI_ROOT = CYCLE.PREPROCESSED_INPUT
from pathlib import Path
import sys
sys.dont_write_bytecode = True
PROJECT = CYCLE.PROJECT_ROOT
sys.path.insert(0, str(PROJECT))
from runtime import ROI_ROOT, CROP_FOLDER, moved_path
ROI_ROOT = CYCLE.PREPROCESSED_INPUT
PYTHON = CYCLE.PYTHON
REFERENCE = CYCLE.MAIN_ROOT
MODEL_BASE = CYCLE.ARTIFACT_ROOT
CNN_MODEL = CYCLE.STAGE_MODEL
A_MODEL = CYCLE.A_MODEL
LABEL_ROOT = CYCLE.LABEL_FILE.parent
LABEL_FILE = CYCLE.LABEL_FILE
CNN_ROOT = CYCLE.MODEL_OUTPUT / '01_CNN'
INPUT_ROOT = CYCLE.MODEL_OUTPUT / '00_input'
A_ROOT = CYCLE.MODEL_OUTPUT / '02_A-step'
FEATURE_DB = CYCLE.INPUT_FEATURE_DB
CROP_DB = CYCLE.INPUT_CROP_DB
MASK_DB = CYCLE.INPUT_MASK_DB
CROP_ROOT = CYCLE.INPUT_CROP_ROOT
WORK = CYCLE.OUTPUT_ROOT / 'work' / 'model'
BATCH_SIZE = CYCLE.BATCH_SIZE
DEVICE = CYCLE.DEVICE
EXPECTED_CNN_SHA = 'dc2b180b471ba47f67e617d4d6aa3f43c08016460dee6cf9cfb8ead226587eb2'
import argparse
from collections import Counter
from contextlib import closing
from datetime import datetime
from zoneinfo import ZoneInfo
import csv
import hashlib
import json
import math
import re
import sqlite3
import numpy as np
from openpyxl import load_workbook
from PIL import Image
from tqdm import tqdm
IMAGE_NAME = re.compile('^(bed\\d+)_(20\\d{6})_(\\d{6})_(cam[012])(?:_r(?:f)?)?_+([bt]\\d+)(?:_p\\d+)?\\.(?:png|jpg|jpeg)$', re.I)

def require(ok, message):
    if not ok:
        raise ValueError(message)

def sha(path):
    h = hashlib.sha256()
    with Path(path).open('rb') as f:
        for block in iter(lambda : f.read(4 * 1024 * 1024), b''):
            h.update(block)
    return h.hexdigest()

def stamp():
    return datetime.now(ZoneInfo('Asia/Seoul')).strftime('%Y%m%d_%H%M%S_%f')

def clean(value):
    if isinstance(value, dict):
        return {str(k): clean(v) for (k, v) in value.items()}
    if isinstance(value, (list, tuple)):
        return [clean(v) for v in value]
    if isinstance(value, (datetime, Path)):
        return str(value)
    if isinstance(value, np.generic):
        return clean(value.item())
    if isinstance(value, float) and (not math.isfinite(value)):
        return None
    return value

def write_json(path, value):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_suffix(path.suffix + '.writing')
    tmp.write_text(json.dumps(clean(value), ensure_ascii=False, indent=2, allow_nan=False), encoding='utf-8')
    tmp.replace(path)

def read_json(path):
    return json.loads(Path(path).read_text(encoding='utf-8-sig'))

def rows_xlsx(path, sheet):
    with closing(load_workbook(path, read_only=True, data_only=True)) as wb:
        it = wb[sheet].iter_rows(values_only=True)
        headers = next(it)
        require(len(headers) == len(set(headers)), 'Duplicate workbook columns')
        return [dict(zip(headers, r)) for r in it if any((v is not None for v in r))]

def load_reference(name, filename):
    from runtime import load_shared
    return load_shared(name, filename)

def readonly_db(path):
    require(path.is_file(), 'Missing source SQLite: ' + str(path))
    require(not Path(str(path) + '-wal').exists() or Path(str(path) + '-wal').stat().st_size == 0, 'Source SQLite has an active WAL; stop the writer before immutable read')
    return sqlite3.connect(path.as_uri() + '?mode=ro&immutable=1', uri=True)

def signature_sources():
    files = [LABEL_FILE, CNN_MODEL, A_MODEL / 'selected_model_bundle.json', A_MODEL / 'MODEL_SELECTION_FROZEN.json', REFERENCE / 'model/1. Predict_stage.py', REFERENCE / 'model/2. Predict_weight.py', REFERENCE / 'pipeline/crop_integrity.py', REFERENCE / 'preprocessing/mask_record.py', Path(__file__), CYCLE.PROJECT_ROOT / 'settings.py', CYCLE.PROJECT_ROOT / 'pipeline/runtime.py']
    files += [ROI_ROOT / CROP_FOLDER / name for name in ['completion.json', 'feature_completion.json']]
    hashes = {str(p): sha(p) for p in files}
    for fold in read_json(A_MODEL / 'selected_model_bundle.json')['folds']:
        for artifact in fold['files']:
            path = A_MODEL / artifact['path']
            actual = sha(path)
            require(actual == artifact['sha256'], 'Frozen A artifact changed: ' + str(path))
            hashes[str(path)] = actual
    require(hashes[str(CNN_MODEL)] == EXPECTED_CNN_SHA, 'Frozen CNN hash mismatch')
    snapshots = {str(p): [p.stat().st_size, p.stat().st_mtime_ns] for p in [FEATURE_DB, CROP_DB, MASK_DB]}
    return (hashes, snapshots)

def prepare_data():
    (source_hashes, snapshots) = signature_sources()
    candidates = rows_xlsx(LABEL_FILE, 'feature')
    unmatched = rows_xlsx(LABEL_FILE, '미매칭_조사')
    require(len(candidates) == 36 and len(unmatched) == 36, 'Unexpected survey cohort size; review the contract')
    require(len({r['sample_key'] for r in candidates + unmatched}) == 72, 'Survey sample keys are not unique')
    selected = [r for r in candidates if r['match_eligible'] is True]
    require(len(selected) == 32, 'Confirmed matched cohort is not 32')
    bundle = read_json(A_MODEL / 'selected_model_bundle.json')
    require(bundle['selected_model'] == 'cnn_only' and len(bundle['folds']) == 5, 'Wrong frozen A model')
    core = bundle['core_numeric_features']
    require(len(core) == 25, 'A numeric feature count mismatch')
    gate_mod = load_reference('s3_gate_reference', 'pipeline/crop_integrity.py')
    (input_rows, gates, lineage) = ([], [], [])
    with closing(readonly_db(FEATURE_DB)) as features, closing(readonly_db(CROP_DB)) as crops, closing(readonly_db(MASK_DB)) as slots:
        for r in tqdm(selected, desc='32 targets: label / feature / PNG / mask'):
            require(type(r['human_label']) is int and 0 <= r['human_label'] <= 4, 'Invalid human label: ' + r['sample_key'])
            require(isinstance(r['fresh_weight_g'], (int, float)) and math.isfinite(r['fresh_weight_g']) and (r['fresh_weight_g'] > 0), 'Invalid target weight: ' + r['sample_key'])
            record = features.execute('SELECT payload FROM features WHERE record_id=?', (r['record_id'],)).fetchone()
            crop_record = crops.execute('SELECT payload FROM crops WHERE record_id=?', (r['record_id'],)).fetchone()
            slot_record = slots.execute('SELECT bbox_json,length(mask) FROM slots WHERE record_id=?', (r['record_id'],)).fetchone()
            require(record and crop_record and slot_record and (slot_record[1] > 0), 'Missing saved source record')
            (f, c, b) = (json.loads(record[0]), json.loads(crop_record[0]), json.loads(slot_record[0]))
            for key in ['record_id', 'season_id', 'bed_id', 'physical_slot', 'camera', 'capture_date', 'capture_time', 'crop_sha256']:
                require(str(f[key]) == str(r[key]), 'Workbook/source identity mismatch: ' + key)
            for key in core:
                require(math.isclose(float(f[key]), float(r[key]), rel_tol=1e-12, abs_tol=1e-12), 'Numeric feature changed: ' + key)
            require(b['slot_namespace_version'] in {'physical_12col_v1', 'image_left_to_right_v1'} and b['target_available'], 'Unverified physical slot')
            sample = int(r['샘플'])
            survey_slot = f't{sample}' if sample <= 12 else f'b{sample - 12}'
            require(r['physical_slot'] == survey_slot and r['source_slot_verified'] is True, 'Survey sample/physical slot mismatch')
            for key in ['bed_id', 'physical_slot', 'camera', 'capture_date', 'capture_time']:
                require(str(b[key]) == str(r[key]), 'Stored mask identity mismatch: ' + key)
            require(int(f['mask_pixels']) == int(b['mask_pixels']), 'Feature/mask pixel count mismatch')
            src = Path(moved_path(f['crop_path'])).resolve()
            require(src.is_relative_to(CROP_ROOT.resolve()), 'Source crop outside season-3 crop root')
            require(sha(src) == f['crop_sha256'] == c['crop_sha256'] == sha(r['label_crop_path']), 'Crop hash mismatch')
            match = IMAGE_NAME.fullmatch(src.name)
            expected = tuple((str(r[k]) for k in ['bed_id', 'capture_date', 'capture_time', 'camera', 'physical_slot']))
            require(match and match.groups() == expected, 'Crop filename identity mismatch')
            require((r['survey_date'] - r['image_date']).days == r['lead_days'] and 0 <= r['lead_days'] <= 3, 'Survey/image day offset changed')
            mask = gate_mod.unpack(c)
            require(int(np.count_nonzero(mask)) == int(f['mask_pixels']), 'Packed mask pixel count mismatch')
            with Image.open(src) as im:
                rgb = np.asarray(im.convert('RGB'))
            require(rgb.shape[:2] == mask.shape and (not np.any(rgb[mask == 0])), 'PNG/mask dimensions or zero background mismatch')
            c.update(crop_path=str(src), has_crop='True')
            (allow, status, reason) = gate_mod.judge(c, straight_ratio=0.8, minimum=20, main_share=0.9)
            gates.append(dict(record_id=r['record_id'], model_eligible=allow, integrity_status=status, integrity_reason=reason))
            row = {k: f[k] for k in ['record_id', 'season_id', 'bed_id', 'physical_slot', 'camera', 'capture_date', 'capture_time', 'crop_sha256', 'feature_status', 'mask_pixels', 'scale_verified', 'shape_qc', 'feature_qc'] + core}
            row.update(crop_path=str(src), source_crop_path=f['crop_path'], selected_record_id=r['record_id'], selected_camera=r['camera'], slot_namespace_version='physical_12col_v1', source_slot_namespace_version=b['slot_namespace_version'], slot_mapping_basis='USER_CONFIRMED_t1_1_t12_12_b1_13_b12_24; original slot unchanged', front_height_cm=None, volume_proxy=None, crop_source_box=json.dumps(c['crop_source_box']))
            row['분석가능'] = True
            input_rows.append(row)
            lineage.append(dict(sample_key=r['sample_key'], record_id=r['record_id'], crop_sha256=f['crop_sha256'], resolved_crop_path=str(src), model_eligible=allow, integrity_status=status, integrity_reason=reason, lead_days=r['lead_days'], human_label=r['human_label'], fresh_weight_g=r['fresh_weight_g']))
    for (p, before) in snapshots.items():
        stat = Path(p).stat()
        require([stat.st_size, stat.st_mtime_ns] == before, 'Source SQLite changed during preparation')
    require(sha(LABEL_FILE) == source_hashes[str(LABEL_FILE)], 'Human workbook changed during preparation')
    signature = hashlib.sha256(json.dumps(dict(hashes=source_hashes, snapshots=snapshots), sort_keys=True).encode()).hexdigest()
    report = dict(status='READ_ONLY_INPUT_CHECK_PASS', signature=signature, source_hashes=source_hashes, source_db_snapshots=snapshots, source_survey_rows=72, targets=32, identity_review=4, unmatched=36, inference_allowed=sum((g['model_eligible'] for g in gates)), mask_review=sum((not g['model_eligible'] for g in gates)), gate_thresholds=dict(straight_ratio=0.8, minimum=20, main_share=0.9), stage_counts=dict(Counter((r['human_label'] for r in selected))), offset_counts=dict(Counter((r['lead_days'] for r in selected))), adaptation='season3_storage.moved_path and S3 filename parser only; no image/feature/model alteration', inference_labels_present=False, front_used_in_prediction=False, scale_verified=False, current_targets_only=True, future_qc_used=False, lineage=lineage)
    return (input_rows, gates, candidates, unmatched, report)

def sheet_spec(name, rows, columns=None):
    columns = columns or list(dict.fromkeys((k for r in rows for k in r)))
    return dict(name=name, columns=columns, rows=[[clean(r.get(k)) for k in columns] for r in rows])

def save_workbooks(specs, name):
    from runtime import write_workbooks
    write_workbooks(specs)
    for spec in specs:
        with closing(load_workbook(spec['path'], read_only=True, data_only=True)) as wb:
            require(wb.sheetnames == [s['name'] for s in spec['sheets']], 'Saved sheet order mismatch')
            for sheet in spec['sheets']:
                rows = wb[sheet['name']].iter_rows(values_only=True)
                require(list(next(rows)) == sheet['columns'], 'Saved column order mismatch')
                for (got, wanted) in zip(rows, sheet['rows']):
                    for (actual, expected) in zip(got, wanted):
                        if isinstance(actual, datetime) and isinstance(expected, str):
                            actual = actual.strftime('%Y-%m-%d') if len(expected) == 10 else str(actual)
                        if isinstance(expected, (int, float)) and (not isinstance(expected, bool)):
                            require(isinstance(actual, (int, float)) and math.isclose(actual, expected, rel_tol=1e-12, abs_tol=1e-12), 'Saved numeric value changed')
                        else:
                            require(actual == (None if expected == '' else expected), 'Saved workbook value changed')

def save_prepared(data):
    global INPUT_ROOT
    (inputs, gates, candidates, unmatched, report) = data
    done = INPUT_ROOT / 'completion.json'
    if done.exists():
        old = read_json(done)
        if old['signature'] == report['signature']:
            for (p, expected) in old['output_hashes'].items():
                require(sha(p) == expected, 'Prepared output changed: ' + p)
            return old
        INPUT_ROOT = INPUT_ROOT / ('run_' + stamp())
    INPUT_ROOT.mkdir(parents=True, exist_ok=True)
    cohorts = [dict(r, model_eligible=g['model_eligible'], integrity_status=g['integrity_status'], integrity_reason=g['integrity_reason']) for (r, g) in zip([r for r in candidates if r['match_eligible']], gates)]
    specs = [dict(path=str(INPUT_ROOT / 'model_input.xlsx'), sheets=[sheet_spec('전체', inputs)]), dict(path=str(INPUT_ROOT / 'target_cohort.xlsx'), sheets=[sheet_spec('Targets', cohorts), sheet_spec('Identity_review', [r for r in candidates if not r['match_eligible']]), sheet_spec('미매칭_조사', unmatched)])]
    save_workbooks(specs, 'prepared_workbooks')
    gate_file = INPUT_ROOT / 'crop_integrity.csv'
    with gate_file.open('w', encoding='utf-8-sig', newline='') as f:
        writer = csv.DictWriter(f, fieldnames=list(gates[0]))
        writer.writeheader()
        writer.writerows(gates)
    report.update(status='CNN_A_INPUT_PREPARED_WITH_MASK_REVIEW', prepared_at=stamp(), output_hashes={str(p): sha(p) for p in [gate_file] + [Path(s['path']) for s in specs]})
    write_json(done, report)
    return report

def run_cnn(report):
    out = CNN_ROOT / ('run_' + stamp())
    out.mkdir(parents=True, exist_ok=False)
    write_json(out / 'started.json', dict(status='RUNNING', input_signature=report['signature']))
    mod = load_reference('s3_cnn_reference', 'model/1. Predict_stage.py')
    mod.IMAGE_NAME = IMAGE_NAME
    mod.FEATURE_XLSX = INPUT_ROOT / 'model_input.xlsx'
    (mod.CROP_ROOT, mod.MODEL_PT) = (CROP_ROOT, CNN_MODEL)
    mod.OUTPUT_XLSX = out / 'crop_features_CNN.xlsx'
    mod.CACHE_SQLITE = out / 'cnn_predictions_cache.sqlite3'
    (mod.GATE_CSV, mod.LABEL_CSV, mod.DEVICE) = (INPUT_ROOT / 'crop_integrity.csv', None, DEVICE)
    mod.main()
    rows = rows_xlsx(mod.OUTPUT_XLSX, '전체')
    require(len(rows) == 32, 'CNN output lost target rows')
    successful = [r for r in rows if r['cnn_status'] == 'OK']
    for r in successful:
        p = np.array([r['cnn_p' + str(i)] for i in range(5)])
        require(np.isfinite(p).all() and (p >= 0).all() and (abs(p.sum() - 1) < 1e-05), 'Invalid CNN probability')
    completion = dict(status='COMPLETE_WITH_ROW_ERRORS' if len(successful) < len(rows) else 'COMPLETE', input_signature=report['signature'], rows=len(rows), predicted=len(successful), status_counts=dict(Counter((r['cnn_status'] for r in rows))), model_sha256=sha(CNN_MODEL), output=str(mod.OUTPUT_XLSX), output_sha256=sha(mod.OUTPUT_XLSX), filename_adapter=IMAGE_NAME.pattern, model_training=False, classes=[0, 1, 2, 3, 4])
    write_json(out / 'completion.json', completion)
    write_json(CNN_ROOT / 'latest_completed.json', dict(run=str(out), input_signature=report['signature']))
    return (out, rows)

def run_a(report, cnn_run):
    mod = load_reference('s3_a_reference', 'model/2. Predict_weight.py')
    mod.IMAGE_NAME = IMAGE_NAME
    before = set(A_ROOT.glob('run_*'))
    args = ['a_step', '--input', str(cnn_run / 'crop_features_CNN.xlsx'), '--sheet', '전체', '--output-dir', str(A_ROOT), '--run', str(A_MODEL), '--crop-root', str(CROP_ROOT), '--device', DEVICE, '--batch-size', str(BATCH_SIZE)]
    old = sys.argv
    try:
        sys.argv = args
        mod.main()
    finally:
        sys.argv = old
    completed = [p for p in set(A_ROOT.glob('run_*')) - before if (p / 'completion.json').is_file()]
    require(len(completed) == 1, 'A completed run resolution failed')
    out = completed[0]
    with (out / '01_A_step_predictions.csv').open(encoding='utf-8-sig', newline='') as f:
        rows = list(csv.DictReader(f))
    require(len(rows) == 32, 'A lost target rows')
    write_json(A_ROOT / 'latest_completed.json', dict(run=str(out), input_signature=report['signature'], cnn_run=str(cnn_run)))
    return (out, rows)

def evaluate(candidates, cnn_rows, a_rows):
    cnn = {r['record_id']: r for r in cnn_rows}
    a = {r['record_id']: r for r in a_rows}
    results = []
    for r in candidates:
        if not r['match_eligible']:
            continue
        (c, p) = (cnn[r['record_id']], a[r['record_id']])
        results.append(dict(sample_key=r['sample_key'], record_id=r['record_id'], survey_date=r['survey_date'].strftime('%Y-%m-%d'), image_date=r['image_date'].strftime('%Y-%m-%d'), lead_days=r['lead_days'], camera=r['camera'], human_label=r['human_label'], fresh_weight_g=r['fresh_weight_g'], cnn_status=c['cnn_status'], cnn_pred_stage=c['cnn_pred_stage'], cnn_expected_stage=c['cnn_expected_stage'], cnn_max_prob=c['cnn_max_prob'], **{f'cnn_p{i}': c[f'cnn_p{i}'] for i in range(5)}, a_step_status=p['a_step_status'], a_step_error=p['a_step_error'], pred_weight_g=float(p['pred_weight_g']) if p['pred_weight_g'] else None, area_cm2=r['area_cm2'], label_crop_path=r['label_crop_path'], capture_before_survey_verified=r['capture_before_survey_verified'], integrity_reason=c['crop_integrity_reason'], scale_verified=False))
    summary = []
    groups = [('ALL_0_TO_3_DAY_OFFSET', results), ('SAME_CALENDAR_DAY_SURVEY_TIME_UNKNOWN', [r for r in results if r['lead_days'] == 0]), ('PRIOR_1_TO_3_DAYS', [r for r in results if r['lead_days'] > 0])]
    groups += [(d, [r for r in results if r['survey_date'] == d]) for d in sorted({r['survey_date'] for r in results})]
    for (name, group) in groups:
        stage = [r for r in group if r['cnn_status'] == 'OK']
        weight = [r for r in group if r['a_step_status'] == 'OK' and r['pred_weight_g'] is not None]
        y = np.array([r['fresh_weight_g'] for r in weight])
        pred = np.array([r['pred_weight_g'] for r in weight])
        err = pred - y
        f1 = []
        for k in range(5):
            tp = sum((r['human_label'] == k and r['cnn_pred_stage'] == k for r in stage))
            fp = sum((r['human_label'] != k and r['cnn_pred_stage'] == k for r in stage))
            fn = sum((r['human_label'] == k and r['cnn_pred_stage'] != k for r in stage))
            if tp + fp + fn:
                f1.append(2 * tp / (2 * tp + fp + fn))
        summary.append(dict(group=name, target_n=len(group), cnn_n=len(stage), a_n=len(weight), unavailable_n=len(group) - len(weight), accuracy=sum((r['human_label'] == r['cnn_pred_stage'] for r in stage)) / len(stage) if stage else None, macro_f1_present_union=float(np.mean(f1)) if f1 else None, expected_stage_mae=float(np.mean([abs(r['cnn_expected_stage'] - r['human_label']) for r in stage])) if stage else None, hard_stage_mae=float(np.mean([abs(r['cnn_pred_stage'] - r['human_label']) for r in stage])) if stage else None, weight_mae_g=float(abs(err).mean()) if len(err) else None, weight_rmse_g=float(np.sqrt((err ** 2).mean())) if len(err) else None, weight_bias_g=float(err.mean()) if len(err) else None, weight_r2=float(1 - (err ** 2).sum() / ((y - y.mean()) ** 2).sum()) if len(y) > 1 and np.var(y) > 0 else None))
    confusion = np.zeros((5, 5), int)
    for r in results:
        if r['cnn_status'] == 'OK':
            confusion[r['human_label'], r['cnn_pred_stage']] += 1
    class_metrics = []
    for k in range(5):
        n = int(confusion[k].sum())
        pred_n = int(confusion[:, k].sum())
        tp = int(confusion[k, k])
        class_metrics.append(dict(human_stage=k, support=n, predicted_n=pred_n, correct_n=tp, recall=tp / n if n else None, precision=tp / pred_n if pred_n else None, f1=2 * tp / (n + pred_n) if n + pred_n else None))
    return (results, summary, confusion, class_metrics)

def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--stage', choices=['prepare', 'cnn-a'], default='prepare')
    parser.add_argument('--check', action='store_true')
    args = parser.parse_args()
    data = prepare_data()
    (inputs, gates, candidates, unmatched, report) = data
    print(json.dumps({k: v for (k, v) in report.items() if k not in ['source_hashes', 'source_db_snapshots', 'lineage']}, ensure_ascii=False, indent=2), flush=True)
    if args.check:
        return
    prepared = save_prepared(data)
    print('PREPARATION COMPLETE:', INPUT_ROOT, flush=True)
    if args.stage == 'prepare':
        return
    (cnn_run, cnn_rows) = run_cnn(prepared)
    (a_run, a_rows) = run_a(prepared, cnn_run)
    (results, summary, confusion, class_metrics) = evaluate(candidates, cnn_rows, a_rows)
    out = a_run / 'evaluation'
    out.mkdir()
    save_workbooks([dict(path=str(out / 'CNN_Astep_evaluation.xlsx'), sheets=[sheet_spec('Metrics', summary), sheet_spec('Target_predictions', results), sheet_spec('Stage_metrics', class_metrics), sheet_spec('Identity_review', [r for r in candidates if not r['match_eligible']]), sheet_spec('미매칭_조사', unmatched)])], 'evaluation_workbook')
    protected_ok = all((sha(p) == h for (p, h) in prepared['source_hashes'].items()))
    require(protected_ok, 'Source/model/code changed during inference')
    for (p, before) in prepared['source_db_snapshots'].items():
        s = Path(p).stat()
        require([s.st_size, s.st_mtime_ns] == before, 'Source database modified')
    report = dict(status='CNN_A_EXTERNAL_EVALUATION_COMPLETE_WITH_ROW_ERRORS', model_training=False, cnn_run=str(cnn_run), a_run=str(a_run), input_signature=prepared['signature'], counts=dict(survey=72, targets=32, cnn_predicted=int(confusion.sum()), a_predicted=sum((r['a_step_status'] == 'OK' for r in results)), identity_review=4, unmatched=36), metrics=summary, confusion=confusion.tolist(), classes=class_metrics, source_models_unchanged=protected_ok, limitations=['Images precede survey by 0-3 days; survey time unknown', 'Only three stage-4 targets', 'Date/camera/stage confounded', 'No disease or marketability ground truth', 'Scale remains unverified', 'A input has current targets only; prepare causal history before C'], output_hashes={str(p): sha(p) for p in out.iterdir() if p.is_file()})
    write_json(out / 'completion.json', report)
    write_json(A_ROOT / 'evaluation_latest_completed.json', dict(evaluation=str(out), cnn_run=str(cnn_run), a_run=str(a_run), input_signature=prepared['signature']))
    print(json.dumps(clean(report), ensure_ascii=False, indent=2), flush=True)
if __name__ == '__main__':
    if hasattr(sys.stdout, 'reconfigure'):
        sys.stdout.reconfigure(encoding='utf-8')
    main()
