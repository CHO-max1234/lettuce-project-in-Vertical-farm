"""Match destructive survey records to confirmed physical plants and recent past images. Generate a new blank label workbook without overwriting reviewed labels."""
import sys as _sys
from pathlib import Path as _Path
_CYCLE_ROOT = _Path(__file__).resolve().parents[1]
if str(_CYCLE_ROOT) not in _sys.path:
    _sys.path.insert(0, str(_CYCLE_ROOT))
_sys.path.insert(0, str(_CYCLE_ROOT / 'pipeline'))
import settings as CYCLE
ROI_ROOT = CYCLE.PREPROCESSED_INPUT
from pathlib import Path
import sys
PROJECT = CYCLE.PROJECT_ROOT
sys.path.insert(0, str(PROJECT))
from runtime import moved_path
ROI_ROOT = CYCLE.PREPROCESSED_INPUT
SURVEY = CYCLE.SURVEY_FILE
FEATURE_ROOT = CYCLE.PREPROCESSED_INPUT / CYCLE.CROP_FOLDER
FEATURE_CSV = FEATURE_ROOT / 'features_all.csv'
QC_ROOT = CYCLE.QC_OUTPUT
QC_CSV = QC_ROOT / 'growth_qc.csv'
MASK_DB = CYCLE.INPUT_MASK_DB
OUTPUT_ROOT = CYCLE.MATCH_OUTPUT
WORK_ROOT = CYCLE.OUTPUT_ROOT / 'work' / 'survey_matching'
CHUNK_ROWS = 40000
MAX_LEAD_DAYS = 3
MIN_MEDIAN_RATIO = 2 / 3
COPY_WORKERS = 4
YEAR = 2026
MANUAL_REVIEW = {'bed07_260514_t3': 'BED_NUMBER_NOT_VISIBLE: bed07 number is not visible in the source; slot/mask evidence exists', 'bed07_260514_b2': 'BED_NUMBER_NOT_VISIBLE: bed07 number is not visible in the source; slot/mask evidence exists', 'bed07_260514_b4': 'BED_NUMBER_NOT_VISIBLE: bed07 number is not visible in the source; slot/mask evidence exists', 'bed04_260618_t9': 'BED_IDENTITY_CONFLICT_SUSPECTED: source frames 042052/042102 show label 1; bed04 identity remains unconfirmed'}
import argparse
from collections import Counter
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime
from zoneinfo import ZoneInfo
import hashlib
import json
import math
import re
import shutil
import sqlite3
import time
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from PIL import Image
try:
    from tqdm import tqdm
except ImportError:

    def tqdm(iterable, total=None, desc=''):
        start = time.monotonic()
        last = start
        for (i, item) in enumerate(iterable, 1):
            yield item
            now = time.monotonic()
            if now - last >= 1:
                print(f"{desc}: {i}/{total or '?'}, {now - start:.1f}s", file=sys.stderr)
                last = now
        print(f'tqdm: {desc} {time.monotonic() - start:.1f} ', file=sys.stderr)

def require(ok, message):
    if not ok:
        raise ValueError(message)

def sha(path):
    h = hashlib.sha256()
    with Path(path).open('rb') as f:
        for b in iter(lambda : f.read(4 * 1024 * 1024), b''):
            h.update(b)
    return h.hexdigest()

def scalar(v):
    if v is None or (not isinstance(v, (str, list, dict)) and pd.isna(v)):
        return None
    if isinstance(v, np.generic):
        return v.item()
    return v

def typed(v, col):
    if v == '' or v is None:
        return None
    if col in {'capture_date', 'capture_time'} or col.endswith('_date') or col.endswith('_time'):
        return str(v)
    if isinstance(v, str) and v in {'True', 'False'}:
        return v == 'True'
    if isinstance(v, str) and re.fullmatch('[-+]?(?:\\d+\\.?\\d*|\\.\\d+)(?:[eE][-+]?\\d+)?', v):
        n = float(v)
        return n if math.isfinite(n) else None
    return scalar(v)

def read_survey():
    wb = load_workbook(SURVEY, read_only=True, data_only=True)
    ws = wb['파괴']
    headers = list(next(ws.iter_rows(min_row=2, max_row=2, max_col=11, values_only=True)))
    rows = []
    for (source_row, values) in enumerate(ws.iter_rows(min_row=3, max_col=11, values_only=True), 3):
        if values[0] is None:
            continue
        if isinstance(values[0], datetime):
            day = pd.Timestamp(values[0]).normalize()
        else:
            (month, day_number) = map(int, re.findall('\\d+', str(values[0]))[:2])
            day = pd.Timestamp(YEAR, month, day_number)
        (bed, sample) = (int(values[1]), int(values[2]))
        require(1 <= sample <= 24, f'Validation condition failed in read_survey: 1 <= sample <= 24: {source_row} ')
        slot = f't{sample}' if sample <= 12 else f'b{sample - 12}'
        bed_id = f'bed{bed:02d}'
        key = f'{bed_id}_{day:%y%m%d}_{slot}'
        rows.append(dict(sample_key=key, lettuce_key=f'3작기|{bed_id}|{slot}', survey_date=f'{day:%Y-%m-%d}', survey_source_row=source_row, survey_source_sheet='파괴', survey_source_path=str(SURVEY), fresh_weight_g=values[7], expected_camera='cam0' if int(slot[1:]) <= 6 else 'cam1', **dict(zip(headers, values))))
    transplant_note = ws['A1'].value
    wb.close()
    require(len(rows) == 72, 'Validation condition failed in read_survey: len(rows) == 72')
    require(len({r['sample_key'] for r in rows}) == len(rows), "Validation condition failed in read_survey: len({r['sample_key'] for r in rows}) == len(rows)")
    require(all((isinstance(r['fresh_weight_g'], (int, float)) for r in rows)), "Validation condition failed in read_survey: all((isinstance(r['fresh_weight_g'], (int, float)) for r in rows))")
    return (rows, headers, transplant_note)

def read_features(survey_rows):
    keys = {r['lettuce_key'] for r in survey_rows}
    date_pairs = {(r['lettuce_key'], (pd.Timestamp(r['survey_date']) - pd.Timedelta(days=d)).strftime('%Y%m%d')) for r in survey_rows for d in range(MAX_LEAD_DAYS + 1)}
    (pieces, n) = ([], 0)
    for c in tqdm(pd.read_csv(FEATURE_CSV, dtype=str, keep_default_na=False, chunksize=CHUNK_ROWS), total=math.ceil(717121 / CHUNK_ROWS), desc='read features'):
        n += len(c)
        c = c[c.lettuce_key.isin(keys) & c.camera.isin(['cam0', 'cam1'])].copy()
        c = c[[p in date_pairs for p in zip(c.lettuce_key, c.capture_date)]]
        if len(c):
            pieces.append(c)
    require(n == 717121, 'Validation condition failed in read_features: n == 717121')
    f = pd.concat(pieces, ignore_index=True)
    columns = list(f.columns)
    require(not f.record_id.duplicated().any(), 'Validation condition failed in read_features: not f.record_id.duplicated().any()')
    area = pd.to_numeric(f.area_cm2, errors='coerce')
    scale = pd.to_numeric(f.mm_per_px, errors='coerce')
    f['_area'] = area
    f['_numeric_ok'] = f.feature_status.eq('OK') & f.crop_path.ne('') & np.isfinite(area) & area.gt(0) & f.feature_qc.eq('QC_PASS') & f.shape_qc.eq('PASS') & np.isfinite(scale) & scale.gt(0)
    f['_identity_ok'] = ~(f.ocr_status.eq('NO_VALID_OCR') | f.upstream_qc.str.contains('OCR_IDENTITY_REVIEW|IDENTITY_CONFLICT', regex=True))
    f['_strict_ok'] = f._numeric_ok & f._identity_ok
    return (f, columns)

def choose_on_day(group, tier):
    g = group[group['_strict_ok' if tier == 'STRICT' else '_numeric_ok']].sort_values(['capture_time', 'record_id'])
    if g.empty:
        return None
    median = float(g._area.median())
    for (_, r) in g[g._area.gt(median * MIN_MEDIAN_RATIO)].iterrows():
        if Path(moved_path(r.crop_path)).is_file():
            return (r, dict(day_valid_crop_n=len(g), day_area_median_cm2=median, area_to_day_median_ratio=float(r._area / median), earliest_record_id=g.iloc[0].record_id, replaced_earliest=bool(r.record_id != g.iloc[0].record_id)))
    return None

def select(survey_rows, features, feature_columns):
    (selected, unmatched) = ([], [])
    for s in survey_rows:
        chosen = None
        for tier in ['STRICT', 'REVIEW']:
            for lead in range(MAX_LEAD_DAYS + 1):
                day = pd.Timestamp(s['survey_date']) - pd.Timedelta(days=lead)
                group = features[(features.lettuce_key == s['lettuce_key']) & (features.camera == s['expected_camera']) & (features.capture_date == day.strftime('%Y%m%d'))]
                found = choose_on_day(group, tier)
                if found:
                    (r, calculation) = found
                    usable = tier == 'STRICT'
                    review_note = '' if usable else MANUAL_REVIEW.get(s['sample_key'], 'OCR_IDENTITY_REVIEW: inspect the source numberplate')
                    status = ('MATCHED_SAME_DAY' if lead == 0 else 'MATCHED_PREVIOUS_DAY') if usable else 'REVIEW_IDENTITY_CONFLICT' if 'CONFLICT' in review_note else 'REVIEW_BED_IDENTITY'
                    relative = f"crops/{('' if usable else '_review/')}{s['sample_key']}.png"
                    chosen = {**s, 'image_date': f'{day:%Y-%m-%d}', 'lead_days': lead, 'match_status': status, 'match_eligible': usable, 'review_note': review_note, 'survey_time_known': False, 'capture_before_survey_verified': lead > 0, 'feature_basis': 'ONE_SELECTED_CAPTURE', 'same_day_median_basis': f'{tier}_QUALITY_PASS_SAME_KEY_CAMERA_DATE', **calculation, 'label_crop_relative_path': relative, 'label_crop_path': str(OUTPUT_ROOT / relative), 'source_crop_path': r.crop_path, 'resolved_source_crop_path': moved_path(r.crop_path), **{k: typed(r[k], k) for k in feature_columns}}
                    require(chosen['lettuce_key'] == s['lettuce_key'], 'feature key conflict')
                    selected.append(chosen)
                    break
            if chosen is not None:
                break
        if chosen is None:
            unmatched.append({**s, 'match_status': 'NO_USABLE_CROP_WITHIN_3_DAYS', 'match_eligible': False, 'review_note': 'No same-bed/physical-slot/assigned-camera candidate in the preceding 0-3 days', 'survey_time_known': False, 'capture_before_survey_verified': False})
    for rows in [selected, unmatched]:
        rows.sort(key=lambda r: (r['lettuce_key'], r['survey_date']))
    return (selected, unmatched)

def join_qc(selected):
    qc = pd.read_csv(QC_CSV, dtype=str, keep_default_na=False)
    require(len(qc) == 60282 and (not qc.duplicated(['lettuce_key', 'date']).any()), "Validation condition failed in join_qc: len(qc) == 60282 and (not qc.duplicated(['lettuce_key', 'date']).any())")
    indexed = qc.set_index(['lettuce_key', 'date'], drop=False)
    out = []
    for r in selected:
        key = (r['lettuce_key'], r['image_date'])
        require(key in indexed.index, f'Validation condition failed in join_qc: key in indexed.index: {key} ')
        q = indexed.loc[key]
        entry = {k: r.get(k) for k in ['sample_key', 'lettuce_key', 'survey_date', 'image_date', 'lead_days', 'fresh_weight_g', 'match_status', 'match_eligible', 'review_note', 'record_id', 'camera', 'capture_time', 'survey_source_row', 'label_crop_relative_path', 'label_crop_path', 'capture_before_survey_verified']}
        entry.update(qc_basis='EXISTING_IMAGE_DATE_DAILY_AGGREGATE', qc_date=r['image_date'], qc_camera_matches_capture=q.top_camera == r['camera'], qc_cutoff_verified=False, qc_usage='DIAGNOSTIC_ONLY_DAILY_AGGREGATE; C_STEP_RECOMPUTE_TO_SELECTED_CAPTURE_CUTOFF')
        entry.update({f'qc_{k}': typed(q[k], k) for k in qc.columns})
        out.append(entry)
    return out

def verify_sources(selected):
    with sqlite3.connect(MASK_DB.as_uri() + '?mode=ro', uri=True) as conn:
        for r in tqdm(selected, desc='verify sources'):
            p = Path(r['resolved_source_crop_path'])
            require(sha(p) == r['crop_sha256'], f"PNG hash {r['sample_key']}")
            with Image.open(p) as im:
                im.verify()
            rec = conn.execute('select frame_id,bbox_json,length(mask) from slots where record_id=?', (r['record_id'],)).fetchone()
            require(rec is not None and rec[2] > 0, 'Validation condition failed in verify_sources: rec is not None and rec[2] > 0')
            b = json.loads(rec[1])
            require(b['target_available'] and b['physical_slot'] == r['physical_slot'], 'target/slot mismatch')
            require(b['bed_id'] == r['bed_id'] and b['capture_date'] == r['capture_date'], 'bed/date mismatch')
            require(int(b['mask_pixels']) == int(r['mask_pixels']), 'mask pixels mismatch')
            r['source_slot_verified'] = True
            r['identity_verification_basis'] = 'UPSTREAM_OCR_PASS' if r['match_eligible'] else 'MANUAL_IDENTITY_NOT_CONFIRMED'

def make_payload(selected, unmatched, qc, headers, transplant_note):
    ordered = ['sample_key', 'lettuce_key', 'survey_date', 'image_date', 'lead_days', 'match_status', 'match_eligible', 'fresh_weight_g', 'label_crop_relative_path', 'review_note'] + headers
    rest = [k for r in selected for k in r if k not in ordered]
    columns = list(dict.fromkeys(ordered + rest)) + ['human_label']
    qc_columns = list(qc[0])

    def pack(rows, cols):
        return [[scalar(r.get(c)) for c in cols] for r in rows]
    return dict(workbooks=[dict(filename='feature.xlsx', sheets=[dict(name='feature', columns=columns, rows=pack(selected, columns)), dict(name='미매칭_조사', columns=columns, rows=pack(unmatched, columns))]), dict(filename='qc.xlsx', sheets=[dict(name='qc', columns=qc_columns, rows=pack(qc, qc_columns)), dict(name='미매칭_조사', columns=list(unmatched[0]), rows=pack(unmatched, list(unmatched[0])))])], transplant_note=transplant_note)

def verify_export(payload):
    checked_cells = 0
    for book in payload['workbooks']:
        wb = load_workbook(OUTPUT_ROOT / book['filename'], read_only=True, data_only=True)
        require(wb.sheetnames == [s['name'] for s in book['sheets']], 'XLSX sheets')
        for sheet in book['sheets']:
            actual = list(wb[sheet['name']].iter_rows(values_only=True))
            require(list(actual[0]) == sheet['columns'], 'XLSX columns')
            require(len(actual) == len(sheet['rows']) + 1, 'XLSX rows')
            for (a, expected) in zip(actual[1:], sheet['rows']):
                for (v, e, col) in zip(a, expected, sheet['columns']):
                    if isinstance(v, datetime):
                        v = v.strftime('%Y-%m-%d')
                    if e == '':
                        e = None
                    if isinstance(e, (int, float)) and (not isinstance(e, bool)):
                        require(isinstance(v, (int, float)) and math.isclose(v, e, rel_tol=1e-12, abs_tol=1e-12), f'XLSX numeric {col}')
                    else:
                        require(v == e, f'XLSX value {col}: {v!r} != {e!r}')
                    checked_cells += 1
        wb.close()
    return checked_cells

def copy_crops(selected):

    def one(r):
        dst = Path(r['label_crop_path'])
        dst.parent.mkdir(parents=True, exist_ok=True)
        src = Path(r['resolved_source_crop_path'])
        if dst.exists():
            require(sha(dst) == sha(src), f'Validation condition failed in one: sha(dst) == sha(src): {dst} ')
        else:
            shutil.copy2(src, dst)
        require(sha(dst) == r['crop_sha256'], 'crop copy hash')
    with ThreadPoolExecutor(max_workers=COPY_WORKERS) as pool:
        list(tqdm(pool.map(one, selected), total=len(selected), desc='copy crops'))

def summary(selected, unmatched):
    return dict(total_survey_rows=len(selected) + len(unmatched), crop_candidates=len(selected), match_eligible=sum((r['match_eligible'] for r in selected)), identity_review=sum((not r['match_eligible'] for r in selected)), unmatched=len(unmatched), selected_lead_days=dict(Counter((str(r['lead_days']) for r in selected))), earliest_replacements=sum((r['replaced_earliest'] for r in selected)), review_keys=[r['sample_key'] for r in selected if not r['match_eligible']], by_date={day: dict(selected=sum((r['survey_date'] == day for r in selected)), eligible=sum((r['survey_date'] == day and r['match_eligible'] for r in selected)), review=sum((r['survey_date'] == day and (not r['match_eligible']) for r in selected)), unmatched=sum((r['survey_date'] == day for r in unmatched))) for day in sorted({r['survey_date'] for r in selected + unmatched})})

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--check', action='store_true')
    ap.add_argument('--prepare', action='store_true')
    args = ap.parse_args()
    protected_paths = [SURVEY, FEATURE_ROOT / 'completion.json', FEATURE_ROOT / 'feature_completion.json', QC_ROOT / 'growth_qc_completion.json']
    protected = {str(p): sha(p) for p in protected_paths}
    db_stat = (MASK_DB.stat().st_size, MASK_DB.stat().st_mtime_ns)
    cc = json.loads((FEATURE_ROOT / 'completion.json').read_text(encoding='utf-8'))
    require(cc['held_frames_preserved'] == 134 and cc['inference_calls'] == 0, "Validation condition failed in main: cc['held_frames_preserved'] == 134 and cc['inference_calls'] == 0")
    (rows, headers, note) = read_survey()
    (f, cols) = read_features(rows)
    (selected, unmatched) = select(rows, f, cols)
    require(len(selected) + len(unmatched) == 72, 'Validation condition failed in main: len(selected) + len(unmatched) == 72')
    require(len({r['record_id'] for r in selected}) == len(selected), "Validation condition failed in main: len({r['record_id'] for r in selected}) == len(selected)")
    verify_sources(selected)
    qc = join_qc(selected)
    payload = make_payload(selected, unmatched, qc, headers, note)
    result = summary(selected, unmatched)
    require(all((0 <= r['lead_days'] <= 3 for r in selected)), "Validation condition failed in main: all((0 <= r['lead_days'] <= 3 for r in selected))")
    require(all((r['area_to_day_median_ratio'] > MIN_MEDIAN_RATIO for r in selected)), "Validation condition failed in main: all((r['area_to_day_median_ratio'] > MIN_MEDIAN_RATIO for r in selected))")
    print(json.dumps(result, ensure_ascii=False, indent=2))
    if args.check:
        return
    WORK_ROOT.mkdir(parents=True, exist_ok=True)
    payload_path = WORK_ROOT / 'workbook_payload.json'
    payload_path.write_text(json.dumps(payload, ensure_ascii=False, allow_nan=False), encoding='utf-8')
    if args.prepare:
        return
    OUTPUT_ROOT.mkdir(parents=True, exist_ok=True)
    for book in payload['workbooks']:
        dst = OUTPUT_ROOT / book['filename']
        if dst.exists():
            require(False, f'Validation condition failed in main: False: {dst} ')
    from runtime import write_workbooks
    write_workbooks([dict(path=str(OUTPUT_ROOT / book['filename']), sheets=book['sheets']) for book in payload['workbooks']])
    copy_crops(selected)
    for book in payload['workbooks']:
        sheet = book['sheets'][0]
        pd.DataFrame(sheet['rows'], columns=sheet['columns']).to_csv(OUTPUT_ROOT / book['filename'].replace('.xlsx', '.csv'), index=False, encoding='utf-8-sig')
    checked_cells = verify_export(payload)
    require(len(list((OUTPUT_ROOT / 'crops').rglob('*.png'))) == len(selected), "Validation condition failed in main: len(list((OUTPUT_ROOT / 'crops').rglob('*.png'))) == len(selected)")
    require(protected == {str(p): sha(p) for p in protected_paths}, 'Validation condition failed in main: protected == {str(p): sha(p) for p in protected_paths}')
    require(db_stat == (MASK_DB.stat().st_size, MASK_DB.stat().st_mtime_ns), 'Validation condition failed in main: db_stat == (MASK_DB.stat().st_size, MASK_DB.stat().st_mtime_ns)')
    completion = dict(status='SURVEY_MATCH_COMPLETE_WITH_IDENTITY_REVIEW', finished_at=datetime.now(ZoneInfo('Asia/Seoul')).isoformat(), **result, xlsx_checked_cells=checked_cells, source_rows_preserved=72, human_label_filled=0, human_identity_review_complete=False, front_required=False, inference_calls=0, held_frames_preserved=134, scale_verified=False, qc_usage='DIAGNOSTIC_ONLY_DAILY_AGGREGATE', qc_cutoff_verified=False, history_policy='C_STEP: HISTORY_CUTOFF_IS_SELECTED_IMAGE_TIMESTAMP', inputs_sha256={str(p): sha(p) for p in [SURVEY, FEATURE_CSV, QC_CSV] + protected_paths[1:]}, upstream_unchanged=True, source_sqlite_read_only=True, code_sha256=sha(Path(__file__)), artifacts={p.name: sha(p) for p in OUTPUT_ROOT.iterdir() if p.is_file()})
    (OUTPUT_ROOT / 'completion.json').write_text(json.dumps(completion, ensure_ascii=False, indent=2), encoding='utf-8')
    print('SAVED_AND_VERIFIED', str(OUTPUT_ROOT))
if __name__ == '__main__':
    main()
