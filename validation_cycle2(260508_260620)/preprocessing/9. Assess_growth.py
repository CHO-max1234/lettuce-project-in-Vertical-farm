"""Calculate provisional top-area growth QC using actual observation dates. Break time-series windows across gaps or camera/scale changes."""
import sys as _sys
from pathlib import Path as _Path
_CYCLE_ROOT = _Path(__file__).resolve().parents[1]
if str(_CYCLE_ROOT) not in _sys.path:
    _sys.path.insert(0, str(_CYCLE_ROOT))
_sys.path.insert(0, str(_CYCLE_ROOT / 'pipeline'))
import settings as CYCLE
from pathlib import Path
import sys
PROJECT = CYCLE.PROJECT_ROOT
sys.path.insert(0, str(PROJECT))
from runtime import STORAGE_ROOT
INPUT_ROOT = CYCLE.PREPROCESS_OUTPUT / CYCLE.CROP_FOLDER
INPUT_CSV = INPUT_ROOT / 'features_all.csv'
FEATURE_VIS_ROOT = STORAGE_ROOT / '3작기 실증 적용' / '4. 시각화와 QC' / '1. Feature 시각화'
OUTPUT_ROOT = CYCLE.QC_OUTPUT
CHUNK_ROWS = 20000
WINDOW_DAYS = 7
MIN_VALID_DAYS = 3
MIN_SPAN_DAYS = 2
GOOD_VALID_DAYS = 5
MAX_GAP_DAYS = 3
AREA_THRESHOLD_PCT_DAY = 1.0
SLOWING_RATIO = 0.5
SCALE_PROTOCOL_ID = 'SEASON3_EXISTING_BRIGHTNESS_SCALE_ESTIMATED'
WORKERS = 3
import argparse
from collections import Counter
from datetime import datetime
import hashlib
import json
import math
import time
import numpy as np
import pandas as pd
from tqdm import tqdm
import matplotlib
matplotlib.use('Agg')
METRICS = ['area_cm2', 'circularity', 'solidity', 'angular_cv', 'inner_ratio_50', 'brightness_mean', 'dark_ratio']
SOURCE_COLS = ['lettuce_key', 'capture_key', 'record_id', 'camera', 'bed_id', 'physical_slot', 'capture_date', 'capture_time', 'crop_path', 'feature_status', 'feature_qc', 'shape_qc', 'status', 'selection_status', 'ocr_status', 'upstream_qc', 'mm_per_px', 'scale_status', 'scale_verified', 'scale_source', 'scale_qc', 'is_crop_clipped', 'source_boundary_contact'] + METRICS
OP_COLUMNS = ['lettuce_key', 'date', 'bed_name', 'slot_name', 'top_camera', 'growth_basis', 'area_cm2_top', 'circularity_top', 'solidity_top', 'angular_cv_top', 'inner_ratio_50_top', 'brightness_mean_top', 'dark_ratio_top', 'area_growth_pct_day', 'previous_area_growth_pct_day', 'area_change_pct', 'growth_window_days', 'valid_obs_n', 'n_source_records', 'n_valid_crops', 'growth_qc', 'timeseries_qc', 'source_qc', 'scale_qc', 'final_message']
GROWTH_LABELS = {'GROWTH_OBSERVED': 'Area increased', 'GROWTH_SLOWING': 'Area increase slowed', 'PLATEAU_CANDIDATE': 'Possible area plateau', 'DECREASE_OBSERVED': 'Area decreased', 'INSUFFICIENT_DATA': 'Insufficient valid observations', 'SOURCE_REVIEW': 'Review source quality or identity', 'NO_MEASUREMENT': 'No measurement'}

def require(condition, message):
    if not condition:
        raise ValueError(message)

def sha(path):
    h = hashlib.sha256()
    with Path(path).open('rb') as f:
        for b in iter(lambda : f.read(4 * 1024 * 1024), b''):
            h.update(b)
    return h.hexdigest()

def json_safe(value):
    if isinstance(value, dict):
        return {str(k): json_safe(v) for (k, v) in value.items()}
    if isinstance(value, (list, tuple)):
        return [json_safe(v) for v in value]
    if value is None or value is pd.NaT:
        return None
    if isinstance(value, (pd.Timestamp, datetime)):
        return value.isoformat()[:10]
    if isinstance(value, np.generic):
        value = value.item()
    if value is None or (not isinstance(value, str) and pd.isna(value)):
        return None
    return value

def read_source():
    fc = json.loads((INPUT_ROOT / 'feature_completion.json').read_text(encoding='utf-8'))
    cc = json.loads((INPUT_ROOT / 'completion.json').read_text(encoding='utf-8'))
    require(fc['feature_errors'] == 0 and fc['status'] == 'FULL_FEATURE_COMPLETE_WITH_REVIEW', "Validation condition failed in read_source: fc['feature_errors'] == 0 and fc['status'] == 'FULL_FEATURE_COMPLETE_WITH_REVIEW'")
    require(cc['held_frames_preserved'] == 134, "Validation condition failed in read_source: cc['held_frames_preserved'] == 134")
    protected = {str(p): sha(p) for p in [INPUT_ROOT / 'feature_completion.json', INPUT_ROOT / 'completion.json']}
    parts = []
    count = Counter()
    top_status = Counter()
    n = 0
    seen = set()
    for c in tqdm(pd.read_csv(INPUT_CSV, usecols=SOURCE_COLS, dtype=str, keep_default_na=False, chunksize=CHUNK_ROWS), total=math.ceil(fc['input_rows'] / CHUNK_ROWS), desc='Read source chunks'):
        require(not c.record_id.duplicated().any(), 'Validation condition failed in read_source: not c.record_id.duplicated().any()')
        ids = set(c.record_id)
        require(not ids.intersection(seen), 'Validation condition failed in read_source: not ids.intersection(seen)')
        seen.update(ids)
        n += len(c)
        count.update(c.feature_status)
        c = c.loc[c.camera.isin(['cam0', 'cam1'])].copy()
        top_status.update(c.feature_status)
        for col in METRICS + ['mm_per_px']:
            c[col] = pd.to_numeric(c[col], errors='coerce')
        c['date'] = pd.to_datetime(c.capture_date, format='%Y%m%d', errors='raise')
        c['has_crop'] = c.feature_status.eq('OK') & c.crop_path.ne('') & c.area_cm2.gt(0) & np.isfinite(c.area_cm2)
        c['identity_ok'] = ~(c.ocr_status.eq('NO_VALID_OCR') | c.upstream_qc.str.contains('OCR_IDENTITY_REVIEW|IDENTITY_CONFLICT', regex=True))
        c['valid'] = c.has_crop & c.identity_ok & c.feature_qc.eq('QC_PASS') & c.shape_qc.eq('PASS') & c.mm_per_px.gt(0)
        for (src, out) in [('is_crop_clipped', 'clip_warning'), ('source_boundary_contact', 'boundary_warning')]:
            c[out] = c[src].str.lower().isin(['true', '1']) & c.has_crop
        c['suspect'] = c.has_crop & c.shape_qc.eq('SUSPECT')
        parts.append(c)
    require(n == fc['input_rows'] and count['OK'] == fc['crops'], "Validation condition failed in read_source: n == fc['input_rows'] and count['OK'] == fc['crops']")
    return (pd.concat(parts, ignore_index=True), fc, protected, dict(count))

def build_daily(raw):
    keys = ['lettuce_key', 'camera', 'bed_id', 'physical_slot', 'date']
    primary = pd.read_csv(FEATURE_VIS_ROOT / '_근거' / 'track_selection.csv', dtype={'lettuce_key': str, 'camera': str})
    require(not primary.lettuce_key.duplicated().any(), 'Validation condition failed in build_daily: not primary.lettuce_key.duplicated().any()')
    fallback = raw.groupby(['lettuce_key', 'camera'], as_index=False).agg(valid_n=('valid', 'sum'), crop_n=('has_crop', 'sum'), source_n=('record_id', 'size'))
    fallback = fallback.sort_values(['lettuce_key', 'valid_n', 'crop_n', 'source_n', 'camera'], ascending=[True, False, False, False, True]).drop_duplicates('lettuce_key')
    fallback = fallback.loc[~fallback.lettuce_key.isin(primary.lettuce_key)]
    chosen = pd.concat([primary[['lettuce_key', 'camera']], fallback[['lettuce_key', 'camera']]], ignore_index=True)
    all_counts = raw.groupby(['lettuce_key', 'date'], as_index=False).agg(n_all_top_records=('record_id', 'size'), n_all_top_crops=('has_crop', 'sum'), all_top_record_ids=('record_id', lambda s: json.dumps(list(s), ensure_ascii=False)))
    r = raw.merge(chosen, on=['lettuce_key', 'camera'], validate='many_to_one')
    d = r.groupby(keys, as_index=False).agg(n_source_records=('record_id', 'size'), n_crop_records=('has_crop', 'sum'), n_valid_crops=('valid', 'sum'), n_suspect_crops=('suspect', 'sum'), clip_n=('clip_warning', 'sum'), boundary_n=('boundary_warning', 'sum'), source_record_ids=('record_id', lambda s: json.dumps(list(s), ensure_ascii=False)), source_status=('status', lambda s: ';'.join(sorted(set(s)))), selection_status=('selection_status', lambda s: ';'.join(sorted(set(s)))), source_upstream_qc=('upstream_qc', lambda s: ';'.join(sorted(set(s)))), source_scale_qc=('scale_qc', lambda s: ';'.join(sorted(set(s)))))
    valid = r.loc[r.valid]
    vals = valid.groupby(keys, as_index=False).agg(**{col: (col, 'mean') for col in METRICS}, mm_per_px=('mm_per_px', 'median'), scale_source_methods=('scale_source', lambda s: ';'.join(sorted(set(s)))))
    numeric = r.loc[r.has_crop].groupby(keys, as_index=False)[METRICS].mean().rename(columns={c: 'raw_' + c for c in METRICS})
    rep = valid.sort_values(['lettuce_key', 'date', 'clip_warning', 'boundary_warning', 'capture_time', 'record_id'], ascending=[True, True, True, True, True, True]).drop_duplicates(keys)
    rep = rep[keys + ['record_id', 'crop_path']].rename(columns={'record_id': 'representative_record_id', 'crop_path': 'representative_crop'})
    d = d.merge(vals, on=keys, how='left', validate='one_to_one').merge(numeric, on=keys, how='left', validate='one_to_one')
    d = d.merge(rep, on=keys, how='left', validate='one_to_one').merge(all_counts, on=['lettuce_key', 'date'], how='right', validate='one_to_one')
    d['primary_camera_observed'] = d.camera.notna()
    d['camera'] = d.camera.fillna(d.lettuce_key.map(chosen.set_index('lettuce_key').camera))
    d['bed_id'] = d.bed_id.fillna(d.lettuce_key.str.split('|').str[1])
    d['physical_slot'] = d.physical_slot.fillna(d.lettuce_key.str.split('|').str[2])
    for col in ['n_source_records', 'n_crop_records', 'n_valid_crops', 'n_suspect_crops', 'clip_n', 'boundary_n']:
        d[col] = d[col].fillna(0).astype(int)
    for col in ['source_status', 'selection_status']:
        d[col] = d[col].fillna('PRIMARY_CAMERA_NOT_OBSERVED')
    for col in ['source_upstream_qc', 'source_scale_qc']:
        d[col] = d[col].fillna('')
    d['source_record_ids'] = d.source_record_ids.fillna('[]')
    require(len(d) == len(all_counts) and d.n_all_top_records.sum() == len(raw), 'Validation condition failed in build_daily: len(d) == len(all_counts) and d.n_all_top_records.sum() == len(raw)')
    d = d.rename(columns={'bed_id': 'bed_name', 'physical_slot': 'slot_name', 'camera': 'top_camera', **{c: c + '_top' for c in METRICS}})
    d['growth_basis'] = 'TOP_AREA_ONLY'
    d['scale_qc'] = 'ESTIMATED_NOT_VERIFIED'
    d['scale_verified'] = False
    d['scale_signature'] = SCALE_PROTOCOL_ID
    d['source_qc'] = np.select([d.n_valid_crops.gt(0), d.n_crop_records.gt(0)], ['USABLE', 'REVIEW'], default='NO_MEASUREMENT')
    d['primary_camera_basis'] = np.where(d.lettuce_key.isin(primary.lettuce_key), 'FEATURE_VIS_FIXED_PRIMARY', 'NO_VALID_TRACK_FALLBACK')
    d['side'] = d.slot_name.str[0]
    d['slot_number'] = d.slot_name.str[1:].astype(int)
    d['bed_number'] = d.bed_name.str.extract('(\\d+)')[0].astype(int)
    d = d.sort_values(['lettuce_key', 'date']).reset_index(drop=True)
    require(not d.duplicated(['lettuce_key', 'date']).any(), "Validation condition failed in build_daily: not d.duplicated(['lettuce_key', 'date']).any()")
    return d

def rate_window(w):
    if len(w) < MIN_VALID_DAYS:
        return None
    span = (w[-1]['date'] - w[0]['date']).days
    if span < MIN_SPAN_DAYS:
        return None
    (first, last) = (w[0]['area_cm2_top'], w[-1]['area_cm2_top'])
    return {'area_growth_pct_day': 100 * math.expm1(math.log(last / first) / span), 'area_change_pct': 100 * (last / first - 1), 'growth_window_days': span, 'window_start_date': w[0]['date'], 'window_end_date': w[-1]['date'], 'window_start_area_cm2': first, 'window_end_area_cm2': last, 'valid_obs_n': len(w), 'window_record_ids': json.dumps([x['representative_record_id'] for x in w], ensure_ascii=False)}

def temporal(d, progress=True):
    rows = []
    groups = d.groupby('lettuce_key', sort=False)
    for (key, g) in tqdm(groups, total=d.lettuce_key.nunique(), desc='Past-only top growth', disable=not progress):
        history = []
        last = None
        signature = None
        seg = 0
        for r in g.sort_values('date').to_dict('records'):
            r.update(area_growth_pct_day=np.nan, previous_area_growth_pct_day=np.nan, area_change_pct=np.nan, growth_window_days=np.nan, valid_obs_n=0, window_start_date=pd.NaT, window_end_date=pd.NaT, window_start_area_cm2=np.nan, window_end_area_cm2=np.nan, window_record_ids='[]', timeseries_qc='INSUFFICIENT', growth_qc='INSUFFICIENT_DATA', temporal_segment_id=f'{key}|{seg}')
            if r['source_qc'] != 'USABLE':
                r['growth_qc'] = 'SOURCE_REVIEW' if r['source_qc'] == 'REVIEW' else 'NO_MEASUREMENT'
                rows.append(r)
                continue
            current = (r['top_camera'], r['scale_signature'])
            if last is not None and ((r['date'] - last).days > MAX_GAP_DAYS or signature != current):
                seg += 1
                history = []
            r['temporal_segment_id'] = f'{key}|{seg}'
            last = r['date']
            signature = current
            history.append(r)
            history = [x for x in history if (r['date'] - x['date']).days < 2 * WINDOW_DAYS]
            w = [x for x in history if (r['date'] - x['date']).days < WINDOW_DAYS]
            r['valid_obs_n'] = len(w)
            result = rate_window(w)
            if result:
                r.update(result)
                r['timeseries_qc'] = 'GOOD' if len(w) >= GOOD_VALID_DAYS else 'USABLE'
                previous = rate_window([x for x in history if WINDOW_DAYS <= (r['date'] - x['date']).days < 2 * WINDOW_DAYS])
                prev = previous['area_growth_pct_day'] if previous else np.nan
                r['previous_area_growth_pct_day'] = prev
                v = r['area_growth_pct_day']
                if v < -AREA_THRESHOLD_PCT_DAY:
                    r['growth_qc'] = 'DECREASE_OBSERVED'
                elif abs(v) <= AREA_THRESHOLD_PCT_DAY:
                    if len(w) >= GOOD_VALID_DAYS:
                        r['growth_qc'] = 'PLATEAU_CANDIDATE'
                elif math.isfinite(prev) and prev > AREA_THRESHOLD_PCT_DAY and (v < prev * SLOWING_RATIO):
                    r['growth_qc'] = 'GROWTH_SLOWING'
                else:
                    r['growth_qc'] = 'GROWTH_OBSERVED'
            rows.append(r)
    out = pd.DataFrame(rows)

    def message(r):
        text = GROWTH_LABELS[r.growth_qc]
        if not getattr(r, 'primary_camera_observed', True):
            text += '; Primary camera unavailable'
        if math.isfinite(r.area_growth_pct_day):
            text += f'; Last {int(r.growth_window_days)} days: {r.area_growth_pct_day:+.2f}% per day'
        return text + '; Top-area basis'
    out['final_message'] = out.apply(message, axis=1)
    out['threshold_status'] = 'PROVISIONAL_TOP_AREA_RULES_NOT_BIOLOGICALLY_VALIDATED'
    out['area_threshold_pct_day'] = AREA_THRESHOLD_PCT_DAY
    out['slowing_ratio'] = SLOWING_RATIO
    return out.sort_values(['lettuce_key', 'date']).reset_index(drop=True)

def verify(d):
    require(d.loc[d.source_qc != 'USABLE', 'area_growth_pct_day'].isna().all(), "Validation condition failed in verify: d.loc[d.source_qc != 'USABLE', 'area_growth_pct_day'].isna().all()")
    measured = d.loc[d.area_growth_pct_day.notna()]
    require((measured.window_end_date == measured.date).all(), 'Validation condition failed in verify: (measured.window_end_date == measured.date).all()')
    require(measured.growth_window_days.between(MIN_SPAN_DAYS, WINDOW_DAYS - 1).all(), 'Validation condition failed in verify: measured.growth_window_days.between(MIN_SPAN_DAYS, WINDOW_DAYS - 1).all()')
    bridge = measured.window_start_date.lt(pd.Timestamp('2026-06-12')) & measured.date.ge(pd.Timestamp('2026-06-12'))
    require(not bridge.any(), 'Validation condition failed in verify: not bridge.any()')
    ids = d.groupby('lettuce_key').top_camera.nunique()
    require(ids.le(1).all(), 'Validation condition failed in verify: ids.le(1).all()')
    require(d.scale_verified.eq(False).all(), 'Validation condition failed in verify: d.scale_verified.eq(False).all()')
    require(d.loc[d.date.eq(pd.Timestamp('2026-06-12')), 'area_growth_pct_day'].isna().all(), "Validation condition failed in verify: d.loc[d.date.eq(pd.Timestamp('2026-06-12')), 'area_growth_pct_day'].isna().all()")
    expected = 100 * ((measured.window_end_area_cm2 / measured.window_start_area_cm2) ** (1 / measured.growth_window_days) - 1)
    require(np.allclose(expected, measured.area_growth_pct_day, atol=1e-09), 'Validation condition failed in verify: np.allclose(expected, measured.area_growth_pct_day, atol=1e-09)')
    chosen = d.loc[d.source_qc.eq('USABLE')].groupby('lettuce_key').size().sort_values(ascending=False).head(20).index
    prefix = d.loc[d.lettuce_key.isin(chosen) & d.date.le(pd.Timestamp('2026-06-16'))].copy()
    check = temporal(prefix, progress=False)
    require(np.allclose(prefix.area_growth_pct_day, check.area_growth_pct_day, equal_nan=True), 'Validation condition failed in verify: np.allclose(prefix.area_growth_pct_day, check.area_growth_pct_day, equal_nan=True)')
    require(prefix.growth_qc.tolist() == check.growth_qc.tolist(), 'Validation condition failed in verify: prefix.growth_qc.tolist() == check.growth_qc.tolist()')
    base = dict(lettuce_key='fixture', top_camera='cam1', scale_signature='DIRECT', source_qc='USABLE', representative_record_id='fixture')
    for (values, expected_state) in [([100] * 7, 'PLATEAU_CANDIDATE'), ([100 * 1.02 ** i for i in range(7)], 'GROWTH_OBSERVED'), ([100 * 0.98 ** i for i in range(7)], 'DECREASE_OBSERVED')]:
        x = pd.DataFrame([dict(base, date=pd.Timestamp('2026-05-01') + pd.Timedelta(days=i), area_cm2_top=v) for (i, v) in enumerate(values)])
        require(temporal(x, False).iloc[-1].growth_qc == expected_state, 'Validation condition failed in verify: temporal(x, False).iloc[-1].growth_qc == expected_state')
    values = [100 * 1.1 ** i for i in range(7)] + [100 * 1.1 ** 6 * 1.02 ** (i + 1) for i in range(7)]
    x = pd.DataFrame([dict(base, date=pd.Timestamp('2026-05-01') + pd.Timedelta(days=i), area_cm2_top=v) for (i, v) in enumerate(values)])
    require(temporal(x, False).iloc[-1].growth_qc == 'GROWTH_SLOWING', "Validation condition failed in verify: temporal(x, False).iloc[-1].growth_qc == 'GROWTH_SLOWING'")
    x.loc[7:, 'date'] += pd.Timedelta(days=15)
    require(pd.isna(temporal(x, False).iloc[7].area_growth_pct_day), 'Validation condition failed in verify: pd.isna(temporal(x, False).iloc[7].area_growth_pct_day)')
    return dict(checks='PASS', formula_recomputed_rows=len(measured), prefix_tracks=len(chosen), front_required=False)

def write_streaming_excel(d, root, columns):
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.formatting.rule import CellIsRule
    wb = Workbook(write_only=True)
    body_font = Font(name='Calibri', size=10)
    for (name, cols) in [('operational_output', OP_COLUMNS), ('all_required_columns', columns)]:
        ws = wb.create_sheet(name)
        ws.freeze_panes = 'C2'
        ws.sheet_view.showGridLines = False
        for (j, c) in enumerate(cols, 1):
            ws.column_dimensions[get_column_letter(j)].width = 65 if c == 'final_message' else 36 if c.endswith('_ids') else max(16, min(30, len(c) + 2))
        header = []
        for c in cols:
            cell = WriteOnlyCell(ws, c)
            cell.font = Font(name='Calibri', size=10, bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='215B69')
            header.append(cell)
        ws.append(header)
        ws.row_dimensions[1].height = 30
        for row in tqdm(d[cols].itertuples(index=False, name=None), total=len(d), desc='Write ' + name):
            cells = []
            for (c, v) in zip(cols, row):
                if v is None or (not isinstance(v, str) and pd.isna(v)):
                    v = None
                if isinstance(v, np.generic):
                    v = v.item()
                if isinstance(v, pd.Timestamp):
                    v = v.to_pydatetime()
                cell = WriteOnlyCell(ws, v)
                cell.font = body_font
                if isinstance(v, str):
                    require(len(v) <= 32767, 'Validation condition failed in write_streaming_excel: len(v) <= 32767')
                    cell.data_type = 's'
                elif isinstance(v, datetime):
                    cell.number_format = 'yyyy-mm-dd'
                elif isinstance(v, float):
                    cell.number_format = '0.00'
                cells.append(cell)
            ws.append(cells)
        ws.auto_filter.ref = f'A1:{get_column_letter(len(cols))}{len(d) + 1}'
        qc = get_column_letter(cols.index('growth_qc') + 1)
        for (state, color) in [('DECREASE_OBSERVED', 'FADBD8'), ('PLATEAU_CANDIDATE', 'E8DAEF'), ('GROWTH_SLOWING', 'FDEBD0'), ('SOURCE_REVIEW', 'FCF3CF')]:
            ws.conditional_formatting.add(f'{qc}2:{qc}{len(d) + 1}', CellIsRule(operator='equal', formula=['"' + state + '"'], fill=PatternFill('solid', fgColor=color)))
    wb.save(root / 'growth_qc.xlsx')

def export_excel(d, root):
    columns = OP_COLUMNS + [c for c in d if c not in OP_COLUMNS]
    payload = {'op_columns': OP_COLUMNS, 'columns': columns, 'rows': [[json_safe(x) for x in r] for r in d[columns].itertuples(index=False, name=None)]}
    work = root / '_근거'
    work.mkdir(parents=True, exist_ok=True)
    payload_path = work / 'workbook_payload.json'
    payload_path.write_text(json.dumps(payload, ensure_ascii=False, allow_nan=False), encoding='utf-8')
    write_streaming_excel(d, root, columns)
    import openpyxl
    wb = openpyxl.load_workbook(root / 'growth_qc.xlsx', read_only=True, data_only=True)
    require(wb.sheetnames == ['operational_output', 'all_required_columns'], "Validation condition failed in export_excel: wb.sheetnames == ['operational_output', 'all_required_columns']")
    evidence_rows = []
    for (name, cols) in [('operational_output', OP_COLUMNS), ('all_required_columns', columns)]:
        values = wb[name].iter_rows(values_only=True)
        require(list(next(values)) == cols, 'Validation condition failed in export_excel: list(next(values)) == cols')
        count = 0
        for (row, expected) in zip(values, payload['rows']):
            target = [expected[columns.index(c)] for c in cols]
            require(len(row) == len(cols), 'Validation condition failed in export_excel: len(row) == len(cols)')
            for (actual, want, col) in zip(row, target, cols):
                if col in ['date', 'window_start_date', 'window_end_date'] and want is not None:
                    actual = actual.strftime('%Y-%m-%d') if isinstance(actual, datetime) else actual
                if isinstance(want, float):
                    require(actual is not None and math.isclose(actual, want, rel_tol=1e-12, abs_tol=1e-09), 'Validation condition failed in export_excel: actual is not None and math.isclose(actual, want, rel_tol=1e-12, abs_tol=1e-09)')
                else:
                    require(actual == want or (want == '' and actual is None), "Validation condition failed in export_excel: actual == want or (want == '' and actual is None)")
            count += 1
            if name == 'all_required_columns':
                evidence_rows.append(row)
        require(count == len(d), 'Validation condition failed in export_excel: count == len(d)')
        require(next(values, None) is None, 'Validation condition failed in export_excel: next(values, None) is None')
    loaded = pd.DataFrame(evidence_rows, columns=columns)
    wb.close()
    for c in ['date', 'window_start_date', 'window_end_date']:
        loaded[c] = pd.to_datetime(loaded[c])
    payload_path.unlink()
    return loaded

def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--check', action='store_true')
    parser.add_argument('--output', type=Path, default=OUTPUT_ROOT)
    args = parser.parse_args()
    start = time.perf_counter()
    code_signature = sha(__file__)
    (raw, fc, protected, source_counts) = read_source()
    daily = build_daily(raw)
    del raw
    d = temporal(daily)
    checks = verify(d)
    report = dict(status='CHECK_PASS' if args.check else 'GROWTH_QC_COMPLETE_WITH_REVIEW', input_rows=fc['input_rows'], code_sha256=code_signature, visual_review_complete=False, source_feature_counts=source_counts, daily_rows=len(d), keys=d.lettuce_key.nunique(), top_source_rows=int(d.n_all_top_records.sum()), primary_camera_absent_daily_rows=int((~d.primary_camera_observed).sum()), growth_counts=d.growth_qc.value_counts().to_dict(), source_counts=d.source_qc.value_counts().to_dict(), growth_basis='TOP_AREA_ONLY', front_required=False, scale_verified=False, held_frames_preserved=134, workbook_writer='OPENPYXL_WRITE_ONLY', provisional_threshold=AREA_THRESHOLD_PCT_DAY, checks=checks, inference_calls=0, transplant_date_verified=False, disease_labels_created=False, observation_first=d.date.min(), observation_last=d.date.max())
    if not args.check:
        root = args.output
        root.mkdir(parents=True, exist_ok=True)
        d.to_csv(root / 'growth_qc.csv', index=False, encoding='utf-8-sig')
        d = export_excel(d, root)
        require(len(d) == report['daily_rows'], "Validation condition failed in main: len(d) == report['daily_rows']")
        files = [root / 'growth_qc.xlsx', root / 'growth_qc.csv']
        report['outputs'] = [dict(path=str(p), sha256=sha(p), bytes=p.stat().st_size) for p in files]
        report['axis'] = 'observed dates compressed; gap boundary marked; lines broken; rates use actual days'
    require(all((sha(p) == h for (p, h) in protected.items())), 'Validation condition failed in main: all((sha(p) == h for (p, h) in protected.items()))')
    report['protected_completion_unchanged'] = True
    report['elapsed_seconds'] = round(time.perf_counter() - start, 2)
    if not args.check:
        (args.output / 'growth_qc_completion.json').write_text(json.dumps(json_safe(report), ensure_ascii=False, indent=2, allow_nan=False), encoding='utf-8')
    print(json.dumps(json_safe(report), ensure_ascii=False, indent=2, allow_nan=False), flush=True)
if __name__ == '__main__':
    main()
