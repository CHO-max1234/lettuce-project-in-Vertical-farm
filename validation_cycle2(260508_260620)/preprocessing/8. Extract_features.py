"""Calculate image features from the preserved target masks, assess shape quality and pair views only at the same recorded capture."""
from __future__ import annotations
import sys as _sys
from pathlib import Path as _Path
_CYCLE_ROOT = _Path(__file__).resolve().parents[1]
if str(_CYCLE_ROOT) not in _sys.path:
    _sys.path.insert(0, str(_CYCLE_ROOT))
_sys.path.insert(0, str(_CYCLE_ROOT / 'pipeline'))
import settings as CYCLE
from pathlib import Path
import sys
CODE_ROOT = Path(__file__).resolve().parent
sys.path.insert(0, str(CODE_ROOT.parent))
from runtime import CROP_FOLDER, moved_path, moved_tree
DATA_ROOT = CYCLE.PREPROCESS_OUTPUT
CROP_ROOT = DATA_ROOT / CROP_FOLDER
SAMPLE_ROOT = CROP_ROOT / 'samples'
UPSTREAM_ROOT = DATA_ROOT / '4. 자리배치'
OLD_FEATURE = CYCLE.MAIN_PREPROCESS / '8. Extract_features.py'
WORKERS = 8
CHUNK_SIZE = 96
MIN_DATE_PEERS = 20
MIN_FRAME_PEERS = 4
FEATURE_VERSION = 'S3_CROP_FEATURE_PAIR_SHAPE_V1'
PRE_MOVE_SIGNATURES = {'sample': '8c2934048b1f7580af09baa954a18c47e3f54918d7ad31e1d5207cbabf735a1f', 'full': '1c18a26baf57b7258037f164fecb94d4e71c82034e0a2eaaa8cd6961deb6d983'}
import argparse
from collections import Counter, defaultdict
from concurrent.futures import ThreadPoolExecutor
import csv
import hashlib
import importlib.util
import json
import math
import os
import re
import sqlite3
from datetime import datetime
import cv2
import numpy as np
from openpyxl import Workbook, load_workbook
from tqdm import tqdm
cv2.setNumThreads(1)
csv.field_size_limit(64 * 1024 * 1024)
sys.dont_write_bytecode = True
sys.path.insert(0, str(OLD_FEATURE.parent))
spec = importlib.util.spec_from_file_location('s3_reference_feature', OLD_FEATURE)
old = importlib.util.module_from_spec(spec)
spec.loader.exec_module(old)
CASE_NAMES = ('bed29_20260620_014645_cam1_rf__t4.png', 'bed32_20260620_045637_cam1__b9.png')
PAIR_FRAME_IDS = ('0f92b01ca15379f6b3567e6b', 'cfed287b237ef0e86f14b7ba', '546a3a62e094b7f3c8d54e4b')
META = ['lettuce_key', 'capture_key', 'record_id', 'camera', 'bed_id', 'capture_date', 'capture_time', 'physical_slot', 'source_image_name', 'crop_path', 'crop_sha256', 'status', 'selection_status', 'ocr_status', 'scale_status', 'scale_source', 'scale_verified', 'mm_per_px', 'feature_status', 'feature_qc', 'feature_flags', 'shape_qc', 'shape_reasons', 'shape_reference', 'shape_peer_n', 'shape_area_ratio', 'straight_cut_fraction', 'green_fraction', 'pair_status', 'chosen_top_record_id', 'front_record_id', 'front_crop_path', 'paired_top_area_cm2', 'paired_front_height_cm', 'height_index', 'volume_proxy_top_front', 'volume_proxy', 'mask_pixels', 'mask_policy', 'error']
FEATURES = [x for x in old.FEATURES if x != 'volume_proxy']
ALL_FIELDS = list(dict.fromkeys(META + FEATURES + ['front_area_x_height_legacy', 'source_path', 'frame_id', 'season_id', 'upstream_qc', 'scale_qc', 'source_boundary_contact', 'is_crop_clipped', 'evidence_kind', 'target_available', 'target_mask_database', 'target_mask_record_id']))
MAIN_FIELDS = list(dict.fromkeys(META + FEATURES + ['front_area_x_height_legacy']))
PAIR_FIELDS = ['lettuce_key', 'capture_key', 'capture_date', 'capture_time', 'bed_id', 'physical_slot', 'top_camera', 'chosen_top_record_id', 'front_record_id', 'top_crop_path', 'front_crop_path', 'area_cm2_top', 'front_height_cm', 'height_index', 'volume_proxy_top_front', 'volume_proxy', 'pair_status', 'top_shape_qc', 'front_shape_qc', 'top_scale_status', 'front_scale_status']
UNAVAILABLE_FIELDS = ['lettuce_key', 'capture_key', 'record_id', 'camera', 'bed_id', 'capture_date', 'capture_time', 'physical_slot', 'selection_status', 'feature_status', 'shape_qc', 'shape_reasons']

def sha(path: Path) -> str:
    h = hashlib.sha256()
    with path.open('rb') as f:
        for block in iter(lambda : f.read(1024 * 1024), b''):
            h.update(block)
    return h.hexdigest()

def yes(value) -> bool:
    return str(value).lower() in ('1', 'true', 'yes')

def remap(path: str) -> Path:
    p = Path(moved_path(path))
    if not p.is_absolute():
        raise ValueError('Invalid input or state in remap (not p.is_absolute())')
    if p == CROP_ROOT or CROP_ROOT in p.parents:
        return p
    raise ValueError('Invalid input or state in remap')

def normalize(row: dict) -> dict:
    row = moved_tree(dict(row))
    for col in ('crop_path', 'mask_path'):
        if row.get(col):
            row[col] = str(remap(row[col]))
    required = ('record_id', 'season_id', 'camera', 'bed_id', 'capture_date', 'capture_time', 'physical_slot')
    if any((not row.get(x) for x in required)):
        raise ValueError('Invalid input or state in normalize (any((not row.get(x) for x in required)))')
    if row['camera'] not in ('cam0', 'cam1', 'cam2'):
        raise ValueError("Invalid input or state in normalize (row['camera'] not in ('cam0', 'cam1', 'cam2'))")
    if not re.fullmatch('[tb](?:[1-9]|1[0-2])', row['physical_slot']):
        raise ValueError("Invalid input or state in normalize (not re.fullmatch('[tb](?:[1-9]|1[0-2])', row['physical_slot']))")
    row['lettuce_key'] = '|'.join((row[x] for x in ('season_id', 'bed_id', 'physical_slot')))
    row['capture_key'] = '|'.join((row[x] for x in ('season_id', 'bed_id', 'capture_date', 'capture_time', 'physical_slot')))
    return row

def source_db() -> sqlite3.Connection:
    p = UPSTREAM_ROOT / '_whole_run' / 'state.sqlite3'
    db = sqlite3.connect(p.as_uri() + '?mode=ro', uri=True)
    db.execute('PRAGMA query_only=ON')
    return db

def full_rows():
    p = CROP_ROOT / 'crop_manifest.csv'
    with p.open(encoding='utf-8-sig', newline='') as f:
        for row in csv.DictReader(f):
            yield normalize(row)

def longest_run(vector: np.ndarray) -> int:
    binary = vector.astype(np.int16)
    change = np.diff(np.pad(binary, (1, 1)))
    starts = np.where(change == 1)[0]
    ends = np.where(change == -1)[0]
    return int(max(ends - starts, default=0))

def shape_evidence(crop: np.ndarray, mask: np.ndarray) -> tuple[float, float]:
    binary = mask > 0
    if not binary.any():
        return (0.0, 0.0)
    (ys, xs) = np.where(binary)
    (x1, x2, y1, y2) = (xs.min(), xs.max(), ys.min(), ys.max())
    height = max(1, y2 - y1 + 1)
    width = max(1, x2 - x1 + 1)
    cut = max(longest_run(binary[y1:y2 + 1, x1]) / height, longest_run(binary[y1:y2 + 1, x2]) / height, longest_run(binary[y1, x1:x2 + 1]) / width, longest_run(binary[y2, x1:x2 + 1]) / width)
    hsv = cv2.cvtColor(crop, cv2.COLOR_BGR2HSV)
    green = (hsv[:, :, 0] >= 25) & (hsv[:, :, 0] <= 95) & (hsv[:, :, 1] >= 40) & (hsv[:, :, 2] >= 30)
    return (round(float(cut), 4), round(float(green[binary].mean()), 4))

def compute(row: dict) -> dict:
    out = {key: row.get(key, '') for key in ALL_FIELDS}
    out.update(feature_status='NO_CROP', shape_qc='NOT_APPLICABLE', pair_status='UNPAIRED', mask_policy='', error='')
    if not yes(row.get('has_crop', bool(row.get('crop_path')))):
        return out
    try:
        crop_path = Path(row['crop_path'])
        data = np.fromfile(str(crop_path), dtype=np.uint8)
        if hashlib.sha256(data).hexdigest() != row['crop_sha256']:
            raise ValueError("Invalid input or state in compute (hashlib.sha256(data).hexdigest() != row['crop_sha256'])")
        crop = cv2.imdecode(data, cv2.IMREAD_COLOR)
        if crop is None:
            raise ValueError('Invalid input or state in compute (crop is None)')
        if row.get('mask_data'):
            mask = old.unpack_region(row)
        elif row.get('mask_path'):
            path = Path(row['mask_path'])
            raw = np.fromfile(str(path), dtype=np.uint8)
            if hashlib.sha256(raw).hexdigest() != row['mask_sha256']:
                raise ValueError("Invalid input or state in compute (hashlib.sha256(raw).hexdigest() != row['mask_sha256'])")
            mask = cv2.imdecode(raw, cv2.IMREAD_UNCHANGED)
        else:
            raise ValueError("Invalid input or state in compute (row.get('mask_path'))")
        if mask is None or mask.ndim != 2 or mask.shape != crop.shape[:2]:
            raise ValueError('Invalid input or state in compute (mask is None or mask.ndim != 2 or mask.shape != crop.shape[:2])')
        if not np.all(np.isin(np.unique(mask), [0, 1, 255])):
            raise ValueError('Invalid input or state in compute (not np.all(np.isin(np.unique(mask), [0, 1, 255])))')
        mask = (mask > 0).astype(np.uint8) * 255
        (mm, scale_status, _) = old.scale_for(row, {})
        (vals, qc) = old.FeatureExtractorAndQC().extract_30_features(crop, mask, mm or 1.0, row['camera'] == 'cam2')
        if qc['qc_status'] not in ('QC_PASS', 'QC_WARNING'):
            out.update(feature_status='INVALID_MASK', feature_qc=qc['qc_status'], error=qc.get('qc_details', ''))
            return out
        if not all((math.isfinite(float(x)) for x in vals.values())):
            raise ValueError('Invalid input or state in compute (not all((math.isfinite(float(x)) for x in vals.values())))')
        if row['camera'] == 'cam2':
            out['front_area_x_height_legacy'] = vals.pop('volume_proxy')
            vals['front_height_cm'] = math.sqrt(vals['front_height_cm2'])
        if mm is None:
            for key in ('area_cm2', 'front_height_cm', 'front_height_cm2', 'front_area_x_height_legacy'):
                vals[key] = None
        (cut, green) = shape_evidence(crop, mask)
        out.update(vals)
        out.update(feature_status='OK', feature_qc=qc['qc_status'], feature_flags=qc.get('qc_flags', 'NONE'), scale_status=scale_status, scale_verified=yes(row.get('scale_verified')), mm_per_px=mm, straight_cut_fraction=cut, green_fraction=green, mask_pixels=int((mask > 0).sum()), mask_policy='EXPLICIT_BINARY_MASK', shape_qc='PENDING_COMPARISON')
        return out
    except Exception as exc:
        out.update(feature_status='ERROR', shape_qc='NOT_EVALUABLE', error=repr(exc))
        return out

def initialize(db: sqlite3.Connection, signature: str, mode: str, out: Path, source_signature: str):
    db.execute('CREATE TABLE IF NOT EXISTS config (key TEXT PRIMARY KEY,value TEXT)')
    db.execute('CREATE TABLE IF NOT EXISTS features (record_id TEXT PRIMARY KEY,seq INTEGER,payload TEXT)')
    db.execute('CREATE INDEX IF NOT EXISTS features_order ON features(seq)')
    db.execute('CREATE TABLE IF NOT EXISTS pairs (capture_key TEXT PRIMARY KEY,payload TEXT)')
    prior = db.execute("SELECT value FROM config WHERE key='signature'").fetchone()
    if prior and prior[0] != signature:
        report_path = out / 'feature_completion.json'
        if prior[0] != PRE_MOVE_SIGNATURES[mode] or not report_path.is_file():
            raise ValueError('Invalid input or state in initialize (prior[0] != PRE_MOVE_SIGNATURES[mode] or not report_path.is_file())')
        report = json.loads(report_path.read_text(encoding='utf-8'))
        required_status = 'FULL_FEATURE_COMPLETE_WITH_REVIEW'
        if report.get('status') != required_status or report.get('signature') != prior[0] or report.get('source_crop_signature') != source_signature or (report.get('feature_errors') != 0):
            raise ValueError("Invalid input or state in initialize (report.get('status') != required_status or report.get('signature') != prior[0] or report.get('source_crop_signature') != source_signature or (report.get('feature_errors') != 0))")
        for (name, hash_key) in (('features_all.csv', 'csv_sha256'), ('crop_features.xlsx', 'xlsx_sha256')):
            path = out / name
            if not path.is_file() or sha(path) != report['verification'][hash_key]:
                raise ValueError("Invalid input or state in initialize (not path.is_file() or sha(path) != report['verification'][hash_key])")
        if db.execute('SELECT count(*) FROM features').fetchone()[0] != report['input_rows']:
            raise ValueError("Invalid input or state in initialize (db.execute('SELECT count(*) FROM features').fetchone()[0] != report['input_rows'])")
    db.execute('INSERT OR REPLACE INTO config VALUES(?,?)', ('signature', signature))
    db.commit()

def chunks(iterator, size):
    batch = []
    for item in iterator:
        batch.append(item)
        if len(batch) >= size:
            yield batch
            batch = []
    if batch:
        yield batch

def calculate(rows, db: sqlite3.Connection, workers: int):
    counts = Counter()
    with ThreadPoolExecutor(max_workers=workers) as pool:
        for batch in tqdm(chunks(enumerate(rows), CHUNK_SIZE), desc='crop → feature', unit='chunk'):
            pending = []
            for (seq, row) in batch:
                if db.execute('SELECT 1 FROM features WHERE record_id=?', (row['record_id'],)).fetchone():
                    counts['reused'] += 1
                else:
                    pending.append((seq, row))
            for ((seq, row), result) in zip(pending, pool.map(compute, (r for (_, r) in pending))):
                db.execute('INSERT INTO features VALUES(?,?,?)', (row['record_id'], seq, json.dumps(result, ensure_ascii=False, allow_nan=False)))
                counts[result['feature_status']] += 1
            db.commit()
            counts['input_rows'] += len(batch)
    return counts

def rows_db(db):
    for (rid, payload) in db.execute('SELECT record_id,payload FROM features ORDER BY seq'):
        yield (rid, moved_tree(json.loads(payload)))

def shape_reference(db):
    tracks = defaultdict(list)
    frames = defaultdict(list)
    for (rid, r) in rows_db(db):
        if r['feature_status'] != 'OK':
            continue
        area = r.get('area_cm2')
        if area is None or area <= 0:
            area = r.get('roi_area_px')
        if area is None or area <= 0:
            continue
        basis = 'cm2' if r.get('area_cm2') is not None else 'px'
        group = (r['capture_date'], r['camera'], r['physical_slot'][0], basis)
        tracks[group, r['lettuce_key']].append(float(area))
        frames[r.get('frame_id'), r['camera'], basis].append((rid, float(area)))
    by_group = defaultdict(list)
    for ((group, key), numbers) in tracks.items():
        by_group[group].append((key, float(np.median(numbers))))
    return (by_group, frames)

def apply_shape_qc(db):
    (groups, frames) = shape_reference(db)
    counts = Counter()
    for (rid, r) in tqdm(rows_db(db), desc='shape QC', unit='row'):
        if r['feature_status'] != 'OK':
            continue
        basis = 'cm2' if r.get('area_cm2') is not None else 'px'
        group = (r['capture_date'], r['camera'], r['physical_slot'][0], basis)
        area = r.get('area_cm2')
        if area is None or area <= 0:
            area = r.get('roi_area_px')
        comparable = [value for (key, value) in groups[group] if key != r['lettuce_key']]
        peer_count = len(comparable)
        reference = 'DATE_CAMERA_ROW' if peer_count >= MIN_DATE_PEERS else ''
        if not reference:
            frame_peers = [value for (other, value) in frames[r.get('frame_id'), r['camera'], basis] if other != rid]
            if len(frame_peers) >= MIN_FRAME_PEERS:
                comparable = frame_peers
                peer_count = len(frame_peers)
                reference = 'SAME_FRAME'
        ratio = float(area) / float(np.median(comparable)) if reference else None
        flags = []
        evidence = str(r.get('evidence_kind') or '').lower()
        if reference and ratio < 0.18 and (float(r['green_fraction']) < 0.45) and ('hole' in evidence):
            flags.append('HOLE_RIM_SMALL_NONGREEN')
        if reference and ratio < 0.55 and (float(r['straight_cut_fraction']) >= 0.8) and yes(r.get('is_crop_clipped')):
            flags.append('CUT_FRAGMENT_SMALL_STRAIGHT')
        r.update(shape_area_ratio=round(ratio, 5) if ratio is not None else None, shape_peer_n=peer_count, shape_reference=reference or 'INSUFFICIENT', shape_qc='SUSPECT' if flags else 'PASS' if reference else 'REVIEW_REFERENCE', shape_reasons=';'.join(flags))
        counts[r['shape_qc']] += 1
        db.execute('UPDATE features SET payload=? WHERE record_id=?', (json.dumps(r, ensure_ascii=False, allow_nan=False), rid))
    db.commit()
    return counts

def choose_top(rows: list[dict]) -> tuple[dict | None, str]:
    if not rows:
        return (None, 'NO_TOP')
    col = int(rows[0]['physical_slot'][1:])
    preferred = 'cam0' if col <= 4 else 'cam1'
    by_camera = defaultdict(list)
    for r in rows:
        by_camera[r['camera']].append(r)
    if len(by_camera[preferred]) == 1:
        return (by_camera[preferred][0], 'PREFERRED_TOP')
    if len(by_camera[preferred]) > 1:
        return (None, 'AMBIGUOUS_TOP')
    alternate = 'cam1' if preferred == 'cam0' else 'cam0'
    if len(by_camera[alternate]) == 1:
        return (by_camera[alternate][0], 'ALTERNATE_TOP_REVIEW')
    return (None, 'AMBIGUOUS_TOP' if by_camera[alternate] else 'NO_TOP')

def pair_features(db):
    grouped = defaultdict(list)
    for (rid, r) in rows_db(db):
        if r['feature_status'] == 'OK':
            grouped[r['capture_key']].append(r)
    counts = Counter()
    for (capture_key, items) in tqdm(grouped.items(), desc='top/front pair', unit='capture slot'):
        (top, policy) = choose_top([r for r in items if r['camera'] in ('cam0', 'cam1')])
        fronts = [r for r in items if r['camera'] == 'cam2']
        front = fronts[0] if len(fronts) == 1 else None
        if top is None or front is None:
            counts[policy if top is None else 'NO_FRONT' if not fronts else 'AMBIGUOUS_FRONT'] += 1
            continue
        if top.get('area_cm2') is None or front.get('front_height_cm') is None:
            counts['SCALE_UNAVAILABLE'] += 1
            continue
        area = float(top['area_cm2'])
        height = float(front['front_height_cm'])
        volume = area * height
        height_index = height / math.sqrt(area) if area > 0 else None
        review = policy != 'PREFERRED_TOP' or any((r['shape_qc'] != 'PASS' for r in (top, front)))
        status = 'PAIRED_REVIEW' if review else 'PAIRED_ESTIMATED'
        pair = dict(lettuce_key=top['lettuce_key'], capture_key=capture_key, capture_date=top['capture_date'], capture_time=top['capture_time'], bed_id=top['bed_id'], physical_slot=top['physical_slot'], top_camera=top['camera'], chosen_top_record_id=top['record_id'], front_record_id=front['record_id'], top_crop_path=top['crop_path'], front_crop_path=front['crop_path'], area_cm2_top=area, front_height_cm=height, height_index=height_index, volume_proxy_top_front=volume, volume_proxy=volume, pair_status=status, top_shape_qc=top['shape_qc'], front_shape_qc=front['shape_qc'], top_scale_status=top['scale_status'], front_scale_status=front['scale_status'])
        db.execute('INSERT OR REPLACE INTO pairs VALUES(?,?)', (capture_key, json.dumps(pair, ensure_ascii=False, allow_nan=False)))
        for r in (top, front):
            r.update(pair_status=status, chosen_top_record_id=top['record_id'], front_record_id=front['record_id'], front_crop_path=front['crop_path'], paired_top_area_cm2=area, paired_front_height_cm=height, height_index=height_index, volume_proxy_top_front=volume, volume_proxy=volume)
            db.execute('UPDATE features SET payload=? WHERE record_id=?', (json.dumps(r, ensure_ascii=False, allow_nan=False), r['record_id']))
        counts[status] += 1
    db.commit()
    return counts

def excel_value(value):
    if value is None or value == '':
        return None
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False)
    if isinstance(value, str):
        return value[:32760]
    return value

def export(db, out: Path, summary: dict):
    out.mkdir(parents=True, exist_ok=True)
    csv_target = out / 'features_all.csv'
    xlsx_target = out / 'crop_features.xlsx'
    csv_temp = out / 'features_all.partial.csv'
    xlsx_temp = out / 'crop_features.partial.xlsx'
    wb = Workbook(write_only=True)
    intro = wb.create_sheet('Summary')
    intro.append(['Item', 'Value'])
    for (key, value) in summary.items():
        intro.append([key, json.dumps(value, ensure_ascii=False) if isinstance(value, (dict, list)) else value])
    intro.append(['volume_proxy_top_front', '동일 베드·날짜·시각·물리자리의 윗면 cm² × 정면 높이 cm; 영상 상대 지표'])
    intro.append(['scale_verified', 'Preserve unverified upstream scale; ESTIMATED does not establish physical validation'])
    all_sheet = wb.create_sheet('전체_crop')
    all_sheet.append(MAIN_FIELDS)
    unavailable = wb.create_sheet('Unavailable')
    unavailable.append(UNAVAILABLE_FIELDS)
    paired = wb.create_sheet('Paired_views')
    paired.append(PAIR_FIELDS)
    with csv_temp.open('w', encoding='utf-8-sig', newline='') as f:
        writer = csv.DictWriter(f, fieldnames=ALL_FIELDS, extrasaction='ignore')
        writer.writeheader()
        order = "json_extract(payload,'$.lettuce_key'),json_extract(payload,'$.capture_date'),json_extract(payload,'$.capture_time'),json_extract(payload,'$.camera'),record_id"
        for (payload,) in tqdm(db.execute('SELECT payload FROM features ORDER BY ' + order), desc='CSV / XLSX', unit='row'):
            row = moved_tree(json.loads(payload))
            writer.writerow({k: row.get(k) for k in ALL_FIELDS})
            if row['feature_status'] == 'NO_CROP':
                unavailable.append([excel_value(row.get(k)) for k in UNAVAILABLE_FIELDS])
            else:
                all_sheet.append([excel_value(row.get(k)) for k in MAIN_FIELDS])
        for (payload,) in db.execute("SELECT payload FROM pairs ORDER BY json_extract(payload,'$.lettuce_key'),json_extract(payload,'$.capture_date'),json_extract(payload,'$.capture_time')"):
            row = moved_tree(json.loads(payload))
            paired.append([excel_value(row.get(k)) for k in PAIR_FIELDS])
    wb.save(xlsx_temp)
    os.replace(csv_temp, csv_target)
    os.replace(xlsx_temp, xlsx_target)
    return (csv_target, xlsx_target)

def verify_export(db, csv_path, xlsx_path, expected):
    count = 0
    with csv_path.open(encoding='utf-8-sig', newline='') as f:
        for row in csv.DictReader(f):
            count += 1
            if not row['lettuce_key'] or not row['record_id']:
                raise ValueError("Invalid input or state in verify_export (not row['lettuce_key'] or not row['record_id'])")
    if count != expected:
        raise ValueError(f'Invalid input or state in verify_export (count != expected): {count} {expected} ')
    wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    crop_count = db.execute("SELECT count(*) FROM features WHERE json_extract(payload,'$.feature_status')!='NO_CROP'").fetchone()[0]
    missing = expected - crop_count
    pairs = db.execute('SELECT count(*) FROM pairs').fetchone()[0]
    sheet_counts = tuple((sum((1 for _ in wb[name].iter_rows())) - 1 for name in ('전체_crop', 'Unavailable', 'Paired_views')))
    if sheet_counts != (crop_count, missing, pairs):
        raise ValueError('Invalid input or state in verify_export (sheet_counts != (crop_count, missing, pairs))')
    wb.close()
    return {'csv_rows': count, 'xlsx_crop_rows': crop_count, 'xlsx_unavailable_rows': missing, 'xlsx_pair_rows': pairs, 'csv_sha256': sha(csv_path), 'xlsx_sha256': sha(xlsx_path)}

def run(mode, workers):
    completion = json.loads((CROP_ROOT / 'completion.json').read_text(encoding='utf-8'))
    if completion['verification_status'] != 'PASS' or completion['crop_count'] != 244456:
        raise ValueError("Invalid input or state in run (completion['verification_status'] != 'PASS' or completion['crop_count'] != 244456)")
    signature = hashlib.sha256((sha(Path(__file__)) + sha(OLD_FEATURE) + completion['signature'] + mode).encode()).hexdigest()
    out = CROP_ROOT
    state = out / '_feature_run'
    state.mkdir(parents=True, exist_ok=True)
    lock_path = state / '.run.lock'
    import msvcrt
    with lock_path.open('a+b') as lock:
        lock.seek(0)
        if lock_path.stat().st_size == 0:
            lock.write(b'0')
            lock.flush()
            lock.seek(0)
        msvcrt.locking(lock.fileno(), msvcrt.LK_NBLCK, 1)
        db = sqlite3.connect(state / 'state.sqlite3')
        try:
            initialize(db, signature, mode, out, completion['signature'])
            rows = full_rows()
            calculated = calculate(rows, db, workers)
            count = db.execute('SELECT count(*) FROM features').fetchone()[0]
            expected = 717121
            if count != expected:
                raise ValueError(f'Invalid input or state in run (count != expected): {count} {expected} ')
            selected = db.execute("SELECT count(*) FROM features WHERE json_extract(payload,'$.feature_status')!='NO_CROP'").fetchone()[0]
            if selected != 244456:
                raise ValueError(f'Invalid input or state in run (selected != 244456): {selected} ')
            db.execute('DELETE FROM pairs')
            db.execute("UPDATE features SET payload=json_remove(payload,'$.pair_status','$.chosen_top_record_id','$.front_record_id','$.front_crop_path','$.paired_top_area_cm2','$.paired_front_height_cm','$.height_index','$.volume_proxy_top_front','$.volume_proxy')")
            db.commit()
            shape = apply_shape_qc(db)
            pair = pair_features(db)
            errors = db.execute("SELECT count(*) FROM features WHERE json_extract(payload,'$.feature_status')='ERROR'").fetchone()[0]
            summary = dict(status='FULL_FEATURE_COMPLETE_WITH_REVIEW', version=FEATURE_VERSION, signature=signature, finished_at=datetime.now().astimezone().isoformat(), mode=mode, source_crop_signature=completion['signature'], input_rows=expected, crops=244456, feature_errors=errors, calculation=dict(calculated), shape_qc=dict(shape), pairs=dict(pair), scale_verified=False, auto_qc_approval=False)
            (csv_path, xlsx_path) = export(db, out, summary)
            summary['verification'] = verify_export(db, csv_path, xlsx_path, expected)
            summary['csv_path'] = str(csv_path)
            summary['xlsx_path'] = str(xlsx_path)
            (out / 'feature_completion.json').write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding='utf-8')
            print(json.dumps(summary, ensure_ascii=False, indent=2))
            return 0 if errors == 0 else 1
        finally:
            db.close()
            msvcrt.locking(lock.fileno(), msvcrt.LK_UNLCK, 1)

def check_relocated():
    source = json.loads((CROP_ROOT / 'completion.json').read_text(encoding='utf-8'))
    report = json.loads((CROP_ROOT / 'feature_completion.json').read_text(encoding='utf-8'))
    if source.get('verification_status') != 'PASS' or source.get('crop_count') != 244456 or report.get('status') != 'FULL_FEATURE_COMPLETE_WITH_REVIEW' or (report.get('input_rows') != 717121) or (report.get('feature_errors') != 0) or (report.get('source_crop_signature') != source.get('signature')):
        raise ValueError("Invalid input or state in check_relocated (source.get('verification_status') != 'PASS' or source.get('crop_count') != 244456 or report.get('status') != 'FULL_FEATURE_COMPLETE_WITH_REVIEW' or (report.get('input_rows') != 717121) or (report.get('feature_errors') != 0) or (report.get('source_crop_signature') != source.get('signature')))")
    for name in ('crop_manifest.csv', 'features_all.csv', 'crop_features.xlsx'):
        if not (CROP_ROOT / name).is_file():
            raise FileNotFoundError('Invalid input or state in check_relocated (not (CROP_ROOT / name).is_file())')
    if not any((CROP_ROOT / 'crops').glob('cam*/*/*.png')):
        raise FileNotFoundError("Invalid input or state in check_relocated (not any((CROP_ROOT / 'crops').glob('cam*/*/*.png')))")
    db = sqlite3.connect((CROP_ROOT / '_feature_run/state.sqlite3').as_uri() + '?mode=ro&immutable=1', uri=True)
    try:
        if db.execute('SELECT count(*) FROM features').fetchone()[0] != 717121 or db.execute('SELECT count(*) FROM pairs').fetchone()[0] != 52790:
            raise ValueError("Invalid input or state in check_relocated (db.execute('SELECT count(*) FROM features').fetchone()[0] != 717121 or db.execute('SELECT count(*) FROM pairs').fetchone()[0] != 52790)")
        for name in CASE_NAMES:
            item = db.execute("SELECT payload FROM features WHERE json_extract(payload,'$.crop_path') LIKE ?", ('%' + name,)).fetchone()
            if not item:
                raise ValueError('Invalid input or state in check_relocated (not item)')
            row = moved_tree(json.loads(item[0]))
            if row.get('shape_qc') != 'SUSPECT' or not Path(row['crop_path']).is_file():
                raise ValueError("Invalid input or state in check_relocated (row.get('shape_qc') != 'SUSPECT' or not Path(row['crop_path']).is_file())")
    finally:
        db.close()
    return dict(status='EXISTING_FULL_FEATURE_RELOCATED_PASS', root=str(CROP_ROOT), rows=717121, crops=244456, pairs=52790, errors=0, example_crops_checked=len(CASE_NAMES))

def main():
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument('--run-full', action='store_true', help='main')
    p.add_argument('--check', action='store_true', help='main')
    p.add_argument('--workers', type=int, default=WORKERS)
    args = p.parse_args()
    if not 1 <= args.workers <= 20:
        raise ValueError('Invalid input or state in main (not 1 <= args.workers <= 20)')
    if args.check:
        print(json.dumps(check_relocated(), ensure_ascii=False, indent=2))
        return 0
    return run('full', args.workers)
if __name__ == '__main__':
    raise SystemExit(main())
