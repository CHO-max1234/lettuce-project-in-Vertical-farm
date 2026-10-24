"""Estimate local pixel scale from the known 50 mm reference. Preserve estimated/unverified scale and geometry-group provenance."""
import sys as _sys
from pathlib import Path as _Path
_CYCLE_ROOT = _Path(__file__).resolve().parents[1]
if str(_CYCLE_ROOT) not in _sys.path:
    _sys.path.insert(0, str(_CYCLE_ROOT))
_sys.path.insert(0, str(_CYCLE_ROOT / 'pipeline'))
import settings as CYCLE
from pathlib import Path
ROOT = CYCLE.BRIGHT_OUTPUT
INPUT_DIR = ROOT
OUTPUT_DIR = ROOT / '_실행기록' / 'scale'
VERSION = 'scale_final_v2_estimates_3crop'
REAL_DIAM_MM = 50.0
WORKERS = 4
CHUNK_SIZE = 8
SAVE_OVERLAYS = False
REFERENCE_IMAGES = 3
MIN_REFERENCE_IMAGES = 3
CAM1_FLAT_ANGLE_DEG = 1.5
CAM2_OTHER_WIDTH_PX = 2300
CAM01_HALF_WIDTH_MAX_PX = 1600
DIRECT_GROUP_DIAMETER_REVIEW_RATIO = 0.3
MAX_REFERENCE_ANGLE_DIFF_DEG = 2.5
MAX_REFERENCE_DIAMETER_RANGE_RATIO = 0.25
CAMERA = {'cam0': dict(x=(0.095, 0.385), top=(0.1, 0.32, 0.18), bottom=(0.65, 0.92, 0.82)), 'cam1': dict(x=(0.46, 0.965), top=(0.06, 0.38, 0.2), bottom=(0.48, 0.88, 0.73)), 'cam2': dict(x=(0.003, 0.34), top=(0.045, 0.38, 0.17), bottom=(0.66, 0.96, 0.85))}
import argparse, hashlib, json, re, time, shutil
from collections import Counter
from concurrent.futures import ThreadPoolExecutor
import cv2
import numpy as np
from tqdm import tqdm
cv2.setNumThreads(1)

def read_image(path):
    image = cv2.imdecode(np.fromfile(str(path), np.uint8), cv2.IMREAD_COLOR)
    if image is None:
        raise ValueError('Cannot read ' + str(path))
    return image

def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()

def save_png(path, image):
    (ok, buf) = cv2.imencode('.png', image)
    if not ok:
        raise ValueError('PNG encode failed')
    buf.tofile(str(path))

def sample(im, x, y):
    return cv2.remap(im, np.asarray(x, np.float32).reshape(1, -1), np.asarray(y, np.float32).reshape(1, -1), cv2.INTER_LINEAR, borderMode=cv2.BORDER_CONSTANT).ravel().astype(np.float32)

def boundary_line(gray, green, xrange, settings, kind, max_slope=0.28, *, edge_base=None, neutral_u8=None):
    (h, w) = gray.shape
    (x1, x2) = xrange
    (lo, hi, expected) = settings
    if edge_base is None:
        edge_base = cv2.Canny(cv2.GaussianBlur(gray, (5, 5), 0), 35, 90)
    edge = edge_base.copy()
    if neutral_u8 is None:
        neutral_u8 = (~green).astype(np.uint8) * 255
    edge[:int(lo * h)] = 0
    edge[int(hi * h):] = 0
    edge[:, :x1] = 0
    edge[:, x2:] = 0
    lines = cv2.HoughLinesP(edge, 1, np.pi / 720, threshold=35, minLineLength=max(45, int((x2 - x1) * 0.2)), maxLineGap=24)
    candidates = []
    if lines is not None:
        for (ax, ay, bx, by) in lines.reshape(-1, 4):
            if abs(bx - ax) < 20:
                continue
            a = float((by - ay) / (bx - ax))
            b = float(ay - a * ax)
            mid = a * ((x1 + x2) / 2) + b
            if abs(a) > max_slope or not lo * h < mid < hi * h:
                continue
            xs = np.linspace(max(x1, min(ax, bx)), min(x2, max(ax, bx)), 80)
            ys = a * xs + b
            above = sample(gray, xs, ys - 8)
            below = sample(gray, xs, ys + 8)
            neutral = sample(neutral_u8, xs, ys + (-12 if kind == 'bottom' else 12)) / 255
            contrast = float(np.median(above - below))
            if kind == 'top':
                contrast = abs(contrast)
            score = min(abs(bx - ax) / (x2 - x1), 1) * 35 + max(0, contrast) * 0.8 + neutral.mean() * 12 - abs(mid / h - expected) * 65
            if contrast > 6:
                candidates.append((float(score), a, b, float(contrast), int(abs(bx - ax))))
    if not candidates:
        return ((0.0, expected * h), {'method': 'CAMERA_WINDOW_REVIEW', 'score': 0.0})
    (score, a, b, contrast, length) = max(candidates)
    return ((a, b), {'method': 'PLATE_RAIL_EDGE', 'score': score, 'contrast': contrast, 'length_px': length})

def search_region(image, cam):
    (h, w) = image.shape[:2]
    cfg = CAMERA[cam]
    gray = cv2.cvtColor(image, cv2.COLOR_BGR2GRAY).astype(np.float32)
    hsv = cv2.cvtColor(image, cv2.COLOR_BGR2HSV)
    green = (hsv[:, :, 0] >= 22) & (hsv[:, :, 0] <= 105) & (hsv[:, :, 1] >= 42)
    gray_u8 = gray.astype(np.uint8)
    smooth = cv2.GaussianBlur(gray_u8, (5, 5), 0)
    boundary_inputs = dict(edge_base=cv2.Canny(smooth, 35, 90), neutral_u8=(~green).astype(np.uint8) * 255)
    (xl, xr) = (int(v * w) for v in cfg['x'])
    left_info = {'method': 'BOUNDED_CAMERA_WINDOW'}
    if cam == 'cam0':
        gx = np.abs(cv2.Sobel(gray, cv2.CV_32F, 1, 0, ksize=3))
        profile = gx[int(0.32 * h):int(0.7 * h)].mean(axis=0)
        profile = cv2.GaussianBlur(profile[None, :], (15, 1), 0).ravel()
        indices = np.arange(int(0.075 * w), int(0.15 * w))
        scores = profile[indices] / max(profile[indices].max(), 1) - 0.5 * np.abs(indices / w - 0.105) / 0.075
        xl = int(indices[np.argmax(scores)]) + 5
        left_info = {'method': 'CAM0_VERTICAL_PLATE_EDGE', 'x': xl, 'strength': float(profile[xl - 5])}
    elif cam == 'cam1':
        gx = np.abs(cv2.Sobel(gray, cv2.CV_32F, 1, 0, ksize=3))
        profile = gx[int(0.14 * h):int(0.78 * h)].mean(axis=0)
        profile = cv2.GaussianBlur(profile[None, :], (15, 1), 0).ravel()
        indices = np.arange(int(0.84 * w), int(0.98 * w))
        scores = profile[indices] / max(profile[indices].max(), 1) - 0.38 * np.abs(indices / w - 0.92) / 0.14
        xr = int(indices[np.argmax(scores)]) - 4
        xl = max(0, xr - int(0.5 * w))
        left_info = {'method': 'CAM1_RIGHT_STRUCTURAL_EDGE', 'right_x': xr, 'width_fraction': 0.5, 'strength': float(profile[xr + 4])}
    (bottom, bi) = boundary_line(gray_u8, green, (xl, xr), cfg['bottom'], 'bottom', **boundary_inputs)
    (top, ti) = boundary_line(gray_u8, green, (xl, xr), cfg['top'], 'top', 0.045 if cam == 'cam2' else 0.28, **boundary_inputs)
    if cam == 'cam2':
        (bottom, bi) = boundary_line(gray_u8, green, (xl, int(0.325 * w)), cfg['bottom'], 'bottom', 0.075, **boundary_inputs)
        if w >= 2300:
            top = (bottom[0], bottom[1] - 0.1 * w)
            ti = {'method': 'FRONT_RAIL_OFFSET_WINDOW', 'depth_px': 0.1 * w}
        else:
            (top, ti) = boundary_line(gray_u8, green, (int(0.13 * w), xr), (0.2, 0.38, 0.28), 'top', 0.075, **boundary_inputs)
            ti['framing_review'] = True
    if min((bottom[0] * x + bottom[1] - top[0] * x - top[1] for x in (xl, xr))) < 0.32 * h:
        top = (0.0, cfg['top'][2] * h)
        ti = {'method': 'CAMERA_WINDOW_REVIEW', 'score': 0.0}
    poly = np.array([[xl, top[0] * xl + top[1] + 3], [xr, top[0] * xr + top[1] + 3], [xr, bottom[0] * xr + bottom[1] - 4], [xl, bottom[0] * xl + bottom[1] - 4]], np.int32)
    poly[:, 1] = np.clip(poly[:, 1], 0, h - 1)
    if cam == 'cam2':
        edge = cv2.Canny(smooth, 38, 125)[:, :int(0.16 * w)]
        lines = cv2.HoughLinesP(edge, 1, np.pi / 360, threshold=18, minLineLength=max(24, int(0.18 * h)), maxLineGap=int(0.1 * h))
        rails = []
        front_y = float(np.mean(poly[2:, 1]))
        if lines is not None:
            for (ax, ay, bx, by) in lines.reshape(-1, 4):
                if abs(by - ay) < 20:
                    continue
                a = float((bx - ax) / (by - ay))
                b = float(ax - a * ay)
                if not -1.35 < a < -0.2:
                    continue
                cross = a * front_y + b
                if abs(cross) > 0.1 * w:
                    continue
                cost = 1.35 * abs(cross) / (0.035 * w) + abs(a + 0.695) / 0.24 - 0.18 * np.hypot(bx - ax, by - ay) / h
                rails.append((float(cost), a, b))
        if rails:
            (_, a, b) = min(rails)
            anchor_x = a * front_y + b
            poly[0, 0] = poly[3, 0] = max(xl, int(anchor_x - 0.06 * w))
            left_info = {'method': 'CAM2_CONTINUOUS_LEFT_RAIL', 'a': a, 'b': b, 'front_cross_x': anchor_x}
    mask = np.zeros((h, w), np.uint8)
    cv2.fillPoly(mask, [poly], 255)
    return (mask, poly, {'left': left_info, 'top': ti, 'bottom': bi}, green)

def fit_candidate(contour, gray, green, roi, source, *, green_u8=None):
    if len(contour) < 16:
        return None
    area = float(cv2.contourArea(contour))
    if area < 140:
        return None
    ((cx, cy), (ax, ay), angle) = cv2.fitEllipse(contour)
    (h, w) = gray.shape
    major = max(ax, ay)
    minor = min(ax, ay)
    if not 0.011 * w <= major <= 0.047 * w or minor < 12 or (not 0.25 <= minor / major <= 1):
        return None
    if not (0 <= cx < w and 0 <= cy < h) or roi[int(cy), int(cx)] == 0:
        return None
    rad = np.deg2rad(angle)
    u = np.array([np.cos(rad), np.sin(rad)])
    v = np.array([-np.sin(rad), np.cos(rad)])
    points = contour[:, 0, :].astype(float) - [cx, cy]
    norm = np.column_stack([points @ u / (ax / 2), points @ v / (ay / 2)])
    err = np.abs(np.linalg.norm(norm, axis=1) - 1)
    fill = area / max(np.pi * ax * ay / 4, 1)
    if fill < 0.7 or fill > 1.18:
        return None
    theta = np.linspace(0, 2 * np.pi, 96, endpoint=False)
    xx = np.cos(theta) * ax / 2 * u[0] + np.sin(theta) * ay / 2 * v[0]
    yy = np.cos(theta) * ax / 2 * u[1] + np.sin(theta) * ay / 2 * v[1]
    outside_x = cx + xx * 1.16
    outside_y = cy + yy * 1.16
    roi_cover = float(np.mean(sample(roi, outside_x, outside_y) > 250))
    if green_u8 is None:
        green_u8 = green.astype(np.uint8) * 255
    g_cover = float(np.mean(sample(green_u8, outside_x, outside_y) > 100))
    outside = sample(gray, outside_x, outside_y)
    inside = sample(gray, cx + xx * 0.78, cy + yy * 0.78)
    contrast = outside - inside
    support = float(np.mean(contrast > 9))
    median_contrast = float(np.median(contrast))
    bins = ((np.arctan2(norm[:, 1], norm[:, 0]) + 2 * np.pi) % (2 * np.pi) / (2 * np.pi) * 36).astype(int) % 36
    contour_coverage = len(np.unique(bins)) / 36
    flags = []
    if roi_cover < 0.98:
        flags.append('CUT_BY_ROI_OR_FRAME')
    if g_cover > 0.1:
        flags.append('LEAF_OR_TOWEL_AT_RIM')
    if np.quantile(err, 0.9) > 0.115 or fill < 0.84:
        flags.append('INCOMPLETE_OR_NONELLIPTIC')
    if contour_coverage < 0.94:
        flags.append('INCOMPLETE_CONTOUR')
    if support < 0.8 or median_contrast < 22:
        flags.append('RIM_CONTRAST_INCOMPLETE')
    patch = np.zeros_like(roi)
    cv2.ellipse(patch, ((cx, cy), (ax * 0.85, ay * 0.85), angle), 255, -1)
    inside_green = float(green[patch > 0].mean()) if np.any(patch) else 1
    if inside_green > 0.035:
        flags.append('OCCUPIED_OR_OCCLUDED')
    score = float(2 * support + contour_coverage + min(median_contrast / 100, 1) - 3 * np.quantile(err, 0.9) - g_cover)
    return {'cx': float(cx), 'cy': float(cy), 'axes_px': [float(ax), float(ay)], 'angle_deg': float(angle), 'major_px': float(major), 'minor_px': float(minor), 'axis_ratio': float(minor / major), 'source': source, 'rim_support': support, 'rim_green_fraction': g_cover, 'inside_green_fraction': inside_green, 'ellipse_error_p90': float(np.quantile(err, 0.9)), 'contour_coverage': float(contour_coverage), 'fill_ratio': fill, 'roi_cover': roi_cover, 'score': score, 'flags': flags, 'accepted': not flags}

def detect_holes(image, roi, green):
    gray = cv2.cvtColor(image, cv2.COLOR_BGR2GRAY)
    gray_float = gray.astype(np.float32)
    green_u8 = green.astype(np.uint8) * 255
    blur = cv2.GaussianBlur(gray, (3, 3), 0)
    inputs = []
    for threshold in (55, 75, 95, 115, 135):
        inputs.append((f'dark_{threshold}', cv2.inRange(blur, 0, threshold)))
    inputs.append(('edge', cv2.Canny(blur, 30, 80)))
    candidates = []
    for (source, binary) in inputs:
        (contours, _) = cv2.findContours(binary, cv2.RETR_LIST, cv2.CHAIN_APPROX_NONE)
        for c in contours:
            (x, y, bw, bh) = cv2.boundingRect(c)
            if not 0.01 * image.shape[1] < bw < 0.05 * image.shape[1] or bh < 10:
                continue
            if roi[min(y + bh // 2, roi.shape[0] - 1), min(x + bw // 2, roi.shape[1] - 1)] == 0:
                continue
            item = fit_candidate(c, gray_float, green, roi, source, green_u8=green_u8)
            if item:
                candidates.append(item)
    candidates.sort(key=lambda c: (c['accepted'], c['score']), reverse=True)
    kept = []
    for c in candidates:
        if any((np.hypot(c['cx'] - k['cx'], c['cy'] - k['cy']) < 0.55 * max(c['major_px'], k['major_px']) for k in kept)):
            continue
        kept.append(c)
    for (i, c) in enumerate(sorted(kept, key=lambda c: (c['cy'], c['cx'])), 1):
        c['id'] = i
    return kept

def overlay(image, poly, candidates, status):
    canvas = image.copy()
    cv2.polylines(canvas, [poly], True, (255, 230, 0), 3, cv2.LINE_AA)
    good = [c for c in candidates if c['accepted']]
    selected = max(good, key=lambda c: c['score']) if good else None
    for c in candidates:
        if not c['accepted'] and (c['inside_green_fraction'] > 0.1 or c['ellipse_error_p90'] > 0.2 or c['rim_support'] < 0.4):
            continue
        p = (int(c['cx']), int(c['cy']))
        ell = ((c['cx'], c['cy']), tuple(c['axes_px']), c['angle_deg'])
        color = (20, 40, 245) if c is selected else (40, 210, 40) if c['accepted'] else (0, 150, 230)
        cv2.ellipse(canvas, ell, color, 3 if c['accepted'] else 1, cv2.LINE_AA)
        text = f"H{c['id']} {c['major_px']:.1f}x{c['minor_px']:.1f}" if c['accepted'] else f"X{c['id']}"
        cv2.putText(canvas, text, (p[0] + 8, p[1] - 12), cv2.FONT_HERSHEY_SIMPLEX, 0.55, color, 2, cv2.LINE_AA)
    header = np.full((78, canvas.shape[1], 3), 26, np.uint8)
    title = f'scale ver1 | {status} | complete={len(good)} | candidate 50mm holes; scale NOT verified'
    cv2.putText(header, title, (12, 29), cv2.FONT_HERSHEY_SIMPLEX, 0.66, (240, 240, 240), 1, cv2.LINE_AA)
    cv2.putText(header, 'CYAN: search region   GREEN: complete rim   RED: best candidate   ORANGE X: rejected', (12, 59), cv2.FONT_HERSHEY_SIMPLEX, 0.61, (240, 240, 240), 1, cv2.LINE_AA)
    return np.vstack([header, canvas])
import csv
import math
import sys
import statistics
from datetime import datetime
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
SCALE_FIELDS = ['image_key', 'image_name', 'camera', 'detected', 'n_candidates', 'selected', 'selected_overlay_id', 'selected_x_px', 'selected_y_px', 'selected_diameter_px', 'selected_minor_px', 'selected_mm_per_px', 'bottom_angle_deg', 'geometry_group', 'fallback_median_diameter_px', 'fallback_median_mm_per_px', 'diameter_px', 'mm_per_px', 'scale_source', 'scale_ready', 'scale_verified', 'scale_version', 'scale_qc', 'reference_count', 'reference_1', 'reference_2', 'reference_3', 'reference_max_angle_diff_deg', 'reference_diameter_range_ratio', 'qc', 'bottom_left_x', 'bottom_left_y', 'bottom_right_x', 'bottom_right_y', 'image_width', 'image_height', 'source_path', 'source_sha256', 'overlay_path']
CANDIDATE_FIELDS = ['image_key', 'image_name', 'camera', 'candidate', 'overlay_id', 'selected', 'center_x_px', 'center_y_px', 'diameter_px', 'minor_px', 'mm_per_px', 'ellipse_angle_deg', 'selection_score']

def identify_camera(path):
    matches = set(re.findall('(?<![a-z0-9])cam[012](?![0-9])', path.stem.lower()))
    if len(matches) != 1:
        raise ValueError('Invalid input or state in identify_camera (len(matches) != 1)')
    return matches.pop()

def geometry_group(camera, angle, width):
    if camera in ('cam0', 'cam1'):
        resolution = 'half' if width < CAM01_HALF_WIDTH_MAX_PX else 'full'
        if camera == 'cam0':
            return f'cam0_{resolution}'
        pose = 'flat' if abs(angle) < CAM1_FLAT_ANGLE_DEG else 'sloped'
        return f'cam1_{resolution}_{pose}'
    if camera == 'cam2':
        return 'cam2_other' if width < CAM2_OTHER_WIDTH_PX else 'cam2_standard'
    raise ValueError(f'Invalid input or state in geometry_group: {camera} ')

def process_final(path, input_root, output, image_key=None, image_name=None):
    key = image_key if image_key is not None else path.relative_to(input_root).as_posix()
    row = dict.fromkeys(SCALE_FIELDS)
    row.update(image_key=key, image_name=image_name or path.name, source_path=str(path.resolve()), detected=False, n_candidates=0, scale_ready=False, scale_verified=False, scale_version=VERSION, reference_count=0, scale_source='UNRESOLVED')
    candidates_out = []
    try:
        row['source_sha256'] = sha(path)
        camera = identify_camera(path)
        row['camera'] = camera
        im = read_image(path)
        (h, w) = im.shape[:2]
        row.update(image_width=w, image_height=h)
        (roi, poly, geometry, green) = search_region(im, camera)
        candidates = detect_holes(im, roi, green)
        accepted = sorted((c for c in candidates if c['accepted']), key=lambda c: c['id'])
        selected = max(accepted, key=lambda c: c['score']) if accepted else None
        (xl, yl) = map(int, poly[3])
        (xr, yr) = map(int, poly[2])
        if xr <= xl:
            raise ValueError('Invalid input or state in process_final (xr <= xl)')
        angle = math.degrees(math.atan2(yr - yl, xr - xl))
        row.update(bottom_left_x=xl, bottom_left_y=yl, bottom_right_x=xr, bottom_right_y=yr, bottom_angle_deg=angle, geometry_group=geometry_group(camera, angle, w), detected=bool(accepted), n_candidates=len(accepted))
        for (no, c) in enumerate(accepted, 1):
            cand = dict(zip(CANDIDATE_FIELDS, [key, row['image_name'], camera, f'candidate_{no}', f"H{c['id']}", c is selected, c['cx'], c['cy'], c['major_px'], c['minor_px'], REAL_DIAM_MM / c['major_px'], c['angle_deg'], c['score']]))
            candidates_out.append(cand)
            if c is selected:
                row.update(selected=cand['candidate'], selected_overlay_id=cand['overlay_id'], selected_x_px=c['cx'], selected_y_px=c['cy'], selected_diameter_px=c['major_px'], selected_minor_px=c['minor_px'], selected_mm_per_px=REAL_DIAM_MM / c['major_px'])
        flags = []
        if any(('REVIEW' in v.get('method', '') for v in geometry.values())):
            flags.append('STRUCTURE_WINDOW_REVIEW')
        if geometry['top'].get('framing_review'):
            flags.append('OTHER_FRAMING')
        row['qc'] = ';'.join(flags)
        if SAVE_OVERLAYS:
            target = output / 'overlays' / Path(key + '.png')
            target.parent.mkdir(parents=True, exist_ok=True)
            status = 'COMPLETE_HOLE_CANDIDATE' if accepted else 'NO_COMPLETE_HOLE'
            save_png(target, overlay(im, poly, candidates, status))
            row['overlay_path'] = target.relative_to(output).as_posix()
        if sha(path) != row['source_sha256']:
            raise ValueError("Invalid input or state in process_final (sha(path) != row['source_sha256'])")
        audit = dict(image_key=key, roi_polygon=poly.tolist(), structure=geometry, candidates=candidates, selected_id=selected['id'] if selected else None)
        return (row, candidates_out, audit)
    except Exception as exc:
        import traceback
        row.update(detected=False, n_candidates=0, selected=None, selected_overlay_id=None, selected_x_px=None, selected_y_px=None, selected_diameter_px=None, selected_minor_px=None, selected_mm_per_px=None, scale_source='ERROR', qc=f'{type(exc).__name__}: {exc}')
        return (row, [], {'image_key': key, 'error': row['qc'], 'traceback': traceback.format_exc()})

def reference_index(pool):
    groups = {}
    for row in pool:
        groups.setdefault(row['bottom_angle_deg'], []).append(row)
    angles = sorted(groups)
    return (angles, {angle: sorted(group, key=lambda row: row['image_key'])[:REFERENCE_IMAGES] for (angle, group) in groups.items()})

def nearest_references(index, angle):
    from bisect import bisect_left
    (angles, groups) = index
    position = bisect_left(angles, angle)
    nearby = angles[max(0, position - REFERENCE_IMAGES):position + REFERENCE_IMAGES]
    candidates = [row for candidate_angle in nearby for row in groups[candidate_angle]]
    return sorted(candidates, key=lambda row: (abs(row['bottom_angle_deg'] - angle), row['image_key']))[:REFERENCE_IMAGES]

def complete_scale_rows(rows):
    pools = {}

    def valid(value):
        try:
            return math.isfinite(float(value)) and float(value) > 0
        except (TypeError, ValueError):
            return False
    for row in rows:
        if row['detected'] and row['scale_source'] != 'ERROR' and valid(row.get('selected_diameter_px')):
            pools.setdefault((row['camera'], row['geometry_group']), []).append(row)
    indices = {key: reference_index(pool) for (key, pool) in pools.items()}
    for row in rows:
        if row['scale_source'] == 'ERROR':
            continue
        flags = list(filter(None, [row.get('qc', '')]))
        if row['detected'] and valid(row.get('selected_diameter_px')):
            row.update(diameter_px=row['selected_diameter_px'], mm_per_px=REAL_DIAM_MM / float(row['selected_diameter_px']), scale_source='DIRECT', scale_ready=True)
            continue
        pool = pools.get((row['camera'], row['geometry_group']), [])
        index = indices.get((row['camera'], row['geometry_group']))
        source = 'MEDIAN_SAME_GROUP'
        if not pool:
            row.update(scale_source='UNRESOLVED', scale_ready=False, qc=';'.join(flags + ['NO_SAME_GROUP_REFERENCES']))
            continue
        refs = nearest_references(index, row['bottom_angle_deg'])
        diameters = [float(r['selected_diameter_px']) for r in refs]
        diam = statistics.median(diameters)
        max_diff = max((abs(r['bottom_angle_deg'] - row['bottom_angle_deg']) for r in refs))
        spread = (max(diameters) - min(diameters)) / diam
        if len(refs) < MIN_REFERENCE_IMAGES:
            flags.append('FEW_REFERENCES')
        if max_diff > MAX_REFERENCE_ANGLE_DIFF_DEG:
            flags.append('REFERENCE_ANGLE_DIFFERENCE_HIGH')
        if spread > MAX_REFERENCE_DIAMETER_RANGE_RATIO:
            flags.append('REFERENCE_DIAMETER_SPREAD_HIGH')
        if not index[0][0] <= row['bottom_angle_deg'] <= index[0][-1]:
            flags.append('OUTSIDE_REFERENCE_ANGLE_RANGE')
        row.update(reference_count=len(refs), reference_max_angle_diff_deg=max_diff, reference_diameter_range_ratio=spread, fallback_median_diameter_px=diam, fallback_median_mm_per_px=REAL_DIAM_MM / diam, diameter_px=diam, mm_per_px=REAL_DIAM_MM / diam, scale_source=source, scale_ready=True, qc=';'.join(flags))
        for (i, ref) in enumerate(refs, 1):
            row[f'reference_{i}'] = ref['image_key']
    for row in rows:
        row['scale_qc'] = row.get('qc', '')
    return rows

def flag_direct_group_outliers(rows):
    groups = {}
    for row in rows:
        if row['scale_source'] == 'DIRECT':
            groups.setdefault(row['geometry_group'], []).append(float(row['diameter_px']))
    medians = {group: statistics.median(values) for (group, values) in groups.items()}
    marker = 'DIRECT_DIAMETER_OUTSIDE_GROUP_30PCT'
    flagged = 0
    for row in rows:
        flags = [flag for flag in (row.get('qc') or '').split(';') if flag and flag != marker]
        if row['scale_source'] == 'DIRECT':
            ratio = float(row['diameter_px']) / medians[row['geometry_group']]
            if abs(ratio - 1) > DIRECT_GROUP_DIAMETER_REVIEW_RATIO:
                flags.append(marker)
                flagged += 1
        row['qc'] = ';'.join(flags)
        row['scale_qc'] = row['qc']
    return flagged

def save_scale_csv(rows, output):
    with (output / 'scale_map.csv').open('w', encoding='utf-8-sig', newline='') as f:
        writer = csv.DictWriter(f, fieldnames=SCALE_FIELDS)
        writer.writeheader()
        writer.writerows(rows)

def export_workbook(rows, output, candidate_data=None, max_sheet_rows=1048576, workbook_path=None):
    if not 2 <= max_sheet_rows <= 1048576:
        raise ValueError('Invalid input or state in export_workbook (not 2 <= max_sheet_rows <= 1048576)')
    book = Workbook(write_only=True)
    header_fill = PatternFill('solid', fgColor='24445C')
    amber = PatternFill('solid', fgColor='FFF0D5')
    red = PatternFill('solid', fgColor='FCE5E2')
    header_font = Font(name='Calibri', size=10, bold=True, color='FFFFFF')
    body_font = Font(name='Calibri', size=10)
    link_font = Font(name='Calibri', size=10, color='176DB0', underline='single')
    header_alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    body_alignment = Alignment(vertical='center', wrap_text=False)
    notes_alignment = Alignment(vertical='center', wrap_text=True)

    def new_sheet(sheet_name, headers):
        ws = book.create_sheet(sheet_name)
        ws.freeze_panes = 'C2'
        ws.sheet_view.showGridLines = False
        for (i, heading) in enumerate(headers, 1):
            width = 19
            if heading in ('image_key', 'image_name', 'source_path', 'overlay_path') or (heading.startswith('reference_') and heading[-1:].isdigit()):
                width = 48
            if heading == 'qc':
                width = 52
            if heading == 'camera':
                width = 10
            if sheet_name.startswith('Calculation'):
                width = [29, 92, 84][i - 1]
            ws.column_dimensions[get_column_letter(i)].width = width
        ws.row_dimensions[1].height = 48
        cells = []
        for name in headers:
            c = WriteOnlyCell(ws, name)
            c.font = header_font
            c.fill = header_fill
            c.alignment = header_alignment
            cells.append(c)
        ws.append(cells)
        return ws

    def sheet(sheet_name, headers, records):
        ws = new_sheet(sheet_name, headers)
        count = 1
        page = 1
        for record in records:
            if count >= max_sheet_rows:
                ws.auto_filter.ref = f'A1:{get_column_letter(len(headers))}{count}'
                page += 1
                ws = new_sheet(f'{sheet_name}_{page}', headers)
                count = 1
            count += 1
            if sheet_name == 'Calculation':
                ws.row_dimensions[count].height = 45
            cells = []
            for name in headers:
                value = record.get(name)
                c = WriteOnlyCell(ws, value)
                if isinstance(value, str):
                    c.data_type = 's'
                c.font = body_font
                c.alignment = notes_alignment if sheet_name == 'Calculation' else body_alignment
                if isinstance(value, float):
                    c.number_format = '0.0000' if 'mm_per_px' in name else '0.00'
                    if name == 'reference_diameter_range_ratio':
                        c.number_format = '0.0%'
                if name == 'scale_source' and value != 'DIRECT' or (name == 'detected' and value is False):
                    c.fill = amber
                if name == 'selected' and value is True:
                    c.fill = red
                if name == 'overlay_path' and value:
                    c.hyperlink = value
                    c.font = link_font
                cells.append(c)
            ws.append(cells)
        ws.auto_filter.ref = f'A1:{get_column_letter(len(headers))}{count}'
        return ws

    def candidate_rows():
        with (output / 'candidates.csv').open(encoding='utf-8-sig', newline='') as f:
            for row in csv.DictReader(f):
                for k in CANDIDATE_FIELDS[6:]:
                    row[k] = float(row[k])
                row['selected'] = row['selected'] == 'True'
                yield row
    sheet('scale_map', SCALE_FIELDS, rows)
    sheet('candidate', CANDIDATE_FIELDS, candidate_rows() if candidate_data is None else candidate_data)
    notes = [('reference_diameter_mm', REAL_DIAM_MM, 'Known physical reference diameter'), ('diameter_px', 'Ellipse major axis', 'Minor axis retained separately'), ('mm_per_px', 'reference_diameter_mm / diameter_px', 'Local scale; recalculate after resizing'), ('fallback', 'Median of up to three direct references', 'Same camera, resolution and geometry group only; no cross-group substitution'), ('unresolved', 'No direct reference in the same group', 'Missing scale remains unresolved'), ('scale_verified', False, 'Estimated scale is not a physical measurement validation'), ('direct_group_review', DIRECT_GROUP_DIAMETER_REVIEW_RATIO, 'Flag direct diameters differing by more than this ratio; retain the value'), ('input', 'Brightness-normalized PNG', 'Original source identity and hash are retained')]
    sheet('Calculation', ['Item', 'Value', 'Description'], (dict(zip(['Item', 'Value', 'Description'], n)) for n in notes))
    notes_ws = book.worksheets[-1]
    notes_ws.column_dimensions['A'].width = 29
    notes_ws.column_dimensions['B'].width = 92
    notes_ws.column_dimensions['C'].width = 84
    book.save(workbook_path or output / 'scale_map.xlsx')

def annotate_final(rows, output):
    if not SAVE_OVERLAYS:
        return
    for row in rows:
        if not row['overlay_path']:
            continue
        p = output / row['overlay_path']
        im = read_image(p)
        band = np.full((42, im.shape[1], 3), 26, np.uint8)
        if row['scale_ready']:
            text = f"{row['scale_source']} | diameter={row['diameter_px']:.2f}px | 50mm / diameter = {row['mm_per_px']:.4f} mm/px | bottom={row['bottom_angle_deg']:.2f}deg"
        else:
            text = row['scale_source'] + ' | ' + (row['qc'] or '')
        cv2.putText(band, text, (12, 28), cv2.FONT_HERSHEY_SIMPLEX, 0.65, (100, 225, 245) if not row['detected'] else (220, 240, 220), 1, cv2.LINE_AA)
        save_png(p, np.vstack([im, band]))

def available_output(path):
    if not path.exists() or not any(path.iterdir()):
        return path
    i = 2
    while True:
        p = path.with_name(path.name + f'_{i:02d}')
        if not p.exists():
            return p
        i += 1

def run(input_dir=INPUT_DIR, output_dir=OUTPUT_DIR):
    source = Path(input_dir).expanduser().resolve()
    requested = Path(output_dir).expanduser().resolve()
    if not source.is_dir():
        raise FileNotFoundError('Invalid input or state in run (not source.is_dir())')
    camera_dirs = [source / cam for cam in CAMERA]
    if any((not folder.is_dir() for folder in camera_dirs)):
        raise FileNotFoundError('Invalid input or state in run (any((not folder.is_dir() for folder in camera_dirs)))')
    paths = sorted((p for folder in camera_dirs for p in folder.rglob('*') if p.is_file() and p.suffix.lower() in {'.png', '.jpg', '.jpeg', '.bmp', '.tif', '.tiff'}), key=lambda p: p.relative_to(source).as_posix())
    if not paths:
        raise ValueError('Invalid input or state in run (not paths)')
    if not 0 < REAL_DIAM_MM or not math.isfinite(REAL_DIAM_MM):
        raise ValueError('Invalid input or state in run (not 0 < REAL_DIAM_MM or not math.isfinite(REAL_DIAM_MM))')
    if not 1 <= MIN_REFERENCE_IMAGES <= REFERENCE_IMAGES <= 3:
        raise ValueError('Invalid input or state in run (not 1 <= MIN_REFERENCE_IMAGES <= REFERENCE_IMAGES <= 3)')
    if WORKERS < 1 or CHUNK_SIZE < 1:
        raise ValueError('Invalid input or state in run (WORKERS < 1 or CHUNK_SIZE < 1)')
    output = available_output(requested)
    output.mkdir(parents=True, exist_ok=True)
    suffix = output.name.removeprefix(requested.name)
    workbook_path = source / f'3작기_scale_map{suffix}.xlsx'
    if workbook_path.exists():
        raise FileExistsError('Invalid input or state in run (workbook_path.exists())')
    started = time.perf_counter()
    rows = []
    with (output / 'candidates.csv').open('w', encoding='utf-8-sig', newline='') as cf, (output / 'detection_audit.jsonl').open('w', encoding='utf-8') as jf, ThreadPoolExecutor(max_workers=WORKERS) as pool, tqdm(total=len(paths), desc='scale final', unit='image') as bar:
        writer = csv.DictWriter(cf, fieldnames=CANDIDATE_FIELDS)
        writer.writeheader()
        for start in range(0, len(paths), CHUNK_SIZE):
            for (row, candidates, audit) in pool.map(lambda p: process_final(p, source, output), paths[start:start + CHUNK_SIZE]):
                rows.append(row)
                writer.writerows(candidates)
                jf.write(json.dumps(audit, ensure_ascii=False, allow_nan=False) + '\n')
                bar.update(1)
    complete_scale_rows(rows)
    direct_outliers = flag_direct_group_outliers(rows)
    save_scale_csv(rows, output)
    export_workbook(rows, output)
    annotate_final(rows, output)
    shutil.copy2(output / 'scale_map.xlsx', workbook_path)
    summary = dict(version=VERSION, input_dir=str(source), output_dir=str(output), workbook_path=str(workbook_path), input_images=len(paths), camera_counts=dict(Counter((r['camera'] or 'UNKNOWN' for r in rows))), status_counts=dict(Counter((r['scale_source'] for r in rows))), direct_group_diameter_review=direct_outliers, candidates=sum((r['n_candidates'] for r in rows)), real_diameter_mm=REAL_DIAM_MM, diameter_definition='ellipse_major_axis', detection_version='scale_ver1_frozen', fallback=dict(reference_images=REFERENCE_IMAGES, minimum=MIN_REFERENCE_IMAGES, cam1_flat_angle_deg=CAM1_FLAT_ANGLE_DEG, cam01_half_width_max_px=CAM01_HALF_WIDTH_MAX_PX, cam2_other_width_px=CAM2_OTHER_WIDTH_PX, max_angle_diff_deg=MAX_REFERENCE_ANGLE_DIFF_DEG, max_diameter_range_ratio=MAX_REFERENCE_DIAMETER_RANGE_RATIO), seconds=time.perf_counter() - started, created_at=datetime.now().isoformat(timespec='seconds'))
    (output / 'summary.json').write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    print('scale workbook:', workbook_path)
    return (rows, output)

def main(argv=None):
    if argv is None and ('ipykernel' in sys.modules or 'IPython' in sys.modules):
        argv = []
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--input', type=Path, default=INPUT_DIR)
    parser.add_argument('--output', type=Path, default=OUTPUT_DIR)
    args = parser.parse_args(argv)
    return run(args.input, args.output)
if __name__ == '__main__':
    main()
