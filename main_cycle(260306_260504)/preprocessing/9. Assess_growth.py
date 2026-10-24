"""Compute paired-view daily growth QC and write its numerical workbook. The volume proxy is projected area times front height. These provisional display rules do not diagnose biological growth failure or disease. Figure reproduction is provided separately."""
import sys as _sys
from pathlib import Path as _Path
_CYCLE_ROOT = _Path(__file__).resolve().parents[1]
if str(_CYCLE_ROOT) not in _sys.path:
    _sys.path.insert(0, str(_CYCLE_ROOT))
import settings as CYCLE
from pathlib import Path
import argparse, datetime, hashlib, io, json, math, os, re, sqlite3, tempfile
import numpy as np
import pandas as pd
from openpyxl import load_workbook
try:
    from tqdm import tqdm
except ImportError:

    def tqdm(it, desc='', **kwargs):
        return it
HERE = Path(__file__).resolve().parent
DEFAULT_OUTPUT_ROOT = CYCLE.PREPROCESS_OUTPUT
INPUT_XLSX = None
OUTPUT_DIR = None
PREDICTIONS_FILE = None
CONFIRMED_COHORTS = [{'cycle_id': '2작기', 'cohort_id': '2nd_bed44_45_49_same_cultivar_transplant', 'beds': ['bed44', 'bed45', 'bed49'], 'transplant_date': None}]
DEFAULT_CONFIG = {'transplant_date': None, 'chunk_rows': 5000, 'window_days': 7, 'min_valid_days': 3, 'min_span_days': 2, 'good_valid_days': 5, 'max_gap_days': 3, 'min_peer_n': 5, 'min_peer_beds': 2, 'size_growth_threshold_pct_day': 1, 'height_growth_threshold_cm_day': 0.1, 'slowing_ratio': 0.5, 'peer_band_sd': 0.5, 'relation_gap_z_threshold': 1, 'relation_gap_std_threshold': 1.25, 'structure_delta_threshold': 0.02, 'persistence_days': 3, 'protocol_version': 'growth_qc_daily_v1_20260913', 'threshold_status': 'PROVISIONAL_DISPLAY_RULES_NOT_BIOLOGICALLY_VALIDATED'}
for _key in ['input_csv', 'output_dir', 'output_name', 'cohort_id', 'cohort_beds', 'cohort_confirmed_by_user']:
    DEFAULT_CONFIG.pop(_key, None)
DEFAULT_CONFIG['protocol_version'] = 'integrated_growth_qc_v2_20260913'

def default_paths(discover=True):
    root = Path(os.environ.get('LETTUCE_OUTPUT_ROOT', DEFAULT_OUTPUT_ROOT))
    feature = root / '8. feature' / 'crop_features.xlsx'
    try:
        if (HERE / 'pipeline_paths.py').is_file():
            import pipeline_paths as paths
            root = paths.OUTPUT_ROOT
            feature = paths.FEATURE_ROOT / 'crop_features.xlsx'
    except ImportError:
        pass
    if INPUT_XLSX:
        feature = Path(INPUT_XLSX)
    elif discover and (not feature.is_file()):
        candidates = sorted(feature.parent.glob('crop_features*.xlsx'))
        candidates = [p for p in candidates if p.is_file() and '.writing.' not in p.name]
        if len(candidates) > 1:
            raise ValueError('Invalid input or state in default_paths (rejected condition: len(candidates) > 1)')
        if candidates:
            feature = candidates[0]
    return (feature, Path(OUTPUT_DIR) if OUTPUT_DIR else root / '9. 생육QC')
OP_COLUMNS = ['cycle_id', 'date', 'bed_name', 'slot_name', 'plant_age_days', 'source_pair_id', 'area_cm2_top', 'front_height_cm', 'volume_proxy_top_front', 'circularity_top', 'circularity_front', 'area_growth_pct_day', 'height_growth_cm_day', 'volume_growth_pct_day', 'growth_window_days', 'valid_obs_n', 'peer_group_id', 'peer_n', 'volume_vs_peer_pct', 'growth_vs_peer_pp_day', 'relative_growth_status', 'growth_qc', 'relation_qc', 'slot_qc', 'timeseries_qc', 'scale_qc', 'final_message', 'predicted_weight_g', 'predicted_stage', 'harvest_status']
SOURCE_COLUMNS = ['record_id', 'camera', 'bed_id', 'capture_date', 'capture_time', 'physical_slot', 'crop_path', 'status', 'feature_status', 'feature_qc', 'feature_flags', 'scale_status', 'mm_per_px', 'scale_version', 'scale_source', 'scale_qc', 'mask_policy', 'upstream_qc', 'metadata_warnings', 'reasons', 'track_id', 'pair_id', 'season_id', 'slot_namespace_version', 'observation_view', 'selection_status', 'selected_camera', 'selected_record_id', 'selected_hard_valid', 'ocr_action', 'area_cm2', 'front_height_cm', 'circularity', 'solidity']
GROWTH_TEXT = {'GROWTH_OBSERVED': 'Image-derived growth observed', 'GROWTH_SLOWING': 'Image-derived growth slowed', 'PLATEAU_CANDIDATE': 'Possible plateau', 'DECREASE_OBSERVED': 'Image-derived size decrease observed', 'STRUCTURAL_TRANSITION_SIGNAL': 'Possible structural transition', 'MIXED_CHANGE': 'Mixed changes', 'INSUFFICIENT_DATA': 'Insufficient observations'}
RELATION_TEXT = {'COHERENT': 'Paired views are consistent', 'DIRECTION_DIFFERENCE': 'Paired views change in different directions', 'FRONT_DEFICIT': 'Front-view size is relatively small', 'FRONT_EXCESS': 'Front-view size is relatively large', 'UNSTABLE_RELATION': 'Paired-view relationship is unstable', 'NOT_EVALUABLE': 'Paired-view relationship is unavailable'}

def finite(x):
    try:
        return bool(np.isfinite(float(x)))
    except (ValueError, TypeError):
        return False

def unique_text(values):
    return ';'.join(sorted({str(v) for v in values if pd.notna(v) and str(v).strip()}))

def json_list(values):
    return json.dumps(sorted({str(v) for v in values if pd.notna(v) and str(v).strip()}), ensure_ascii=False)

def sha256(path):
    h = hashlib.sha256()
    with open(path, 'rb') as f:
        for b in iter(lambda : f.read(1024 * 1024), b''):
            h.update(b)
    return h.hexdigest()

def normalize_features(d):
    d = d.fillna('').copy()
    for c in SOURCE_COLUMNS:
        if c not in d:
            raise ValueError('Missing required columns: ' + c)
    for c in SOURCE_COLUMNS:
        d[c] = d[c].map(lambda v: v.strftime('%Y%m%d') if isinstance(v, (pd.Timestamp, datetime.datetime, datetime.date)) else str(v).strip())
    for c in ['area_cm2', 'front_height_cm', 'circularity', 'solidity', 'mm_per_px', 'selected_hard_valid']:
        d[c] = pd.to_numeric(d[c].replace({'True': '1', 'False': '0', 'true': '1', 'false': '0'}), errors='coerce')
    d['day'] = pd.to_datetime(d.capture_date, format='%Y%m%d', errors='raise')
    for c in ['record_id', 'pair_id', 'track_id', 'season_id', 'bed_id', 'physical_slot']:
        if d[c].eq('').any():
            raise ValueError("Invalid input or state in normalize_features (rejected condition: d[c].eq('').any())")
    if d.record_id.duplicated().any() or d.duplicated(['pair_id', 'observation_view']).any():
        raise ValueError("Invalid input or state in normalize_features (rejected condition: d.record_id.duplicated().any() or d.duplicated(['pair_id', 'observation_view']).any())")
    if not d.physical_slot.str.fullmatch('[tb](?:[1-9]|1[0-2])').all():
        raise ValueError("Invalid input or state in normalize_features (rejected condition: not d.physical_slot.str.fullmatch('[tb](?:[1-9]|1[0-2])').all())")
    if not d.bed_id.str.fullmatch('bed\\d+').all():
        raise ValueError('Invalid bed ID')
    if not d.slot_namespace_version.eq('physical_12col_v1').all():
        raise ValueError("Invalid input or state in normalize_features (rejected condition: not d.slot_namespace_version.eq('physical_12col_v1').all())")
    if not d.track_id.eq(d.season_id + '|' + d.bed_id + '|' + d.physical_slot).all():
        raise ValueError("Invalid input or state in normalize_features (rejected condition: not d.track_id.eq(d.season_id + '|' + d.bed_id + '|' + d.physical_slot).all())")
    if not (d.observation_view.eq('top') & d.camera.isin(['cam0', 'cam1']) | d.observation_view.eq('front') & d.camera.eq('cam2')).all():
        raise ValueError('Invalid camera/view')
    selected = d.selected_record_id.ne('')
    if not d.loc[selected, 'selected_record_id'].eq(d.loc[selected, 'record_id']).all():
        raise ValueError("Invalid input or state in normalize_features (rejected condition: not d.loc[selected, 'selected_record_id'].eq(d.loc[selected, 'record_id']).all())")
    keys = ['season_id', 'bed_id', 'physical_slot', 'capture_date', 'capture_time']
    if d.groupby('pair_id')[keys].nunique().gt(1).any().any():
        raise ValueError("Invalid input or state in normalize_features (rejected condition: d.groupby('pair_id')[keys].nunique().gt(1).any().any())")
    return assess_source(d)

def assess_source(d):
    d = d.copy()
    d['numeric_ok'] = d.feature_status.eq('OK') & d.area_cm2.gt(0) & np.isfinite(d.area_cm2) & d.mm_per_px.gt(0) & np.isfinite(d.mm_per_px) & d.mask_policy.eq('EXPLICIT_BINARY_MASK')
    d.loc[d.observation_view.eq('front') & (~d.front_height_cm.gt(0) | ~np.isfinite(d.front_height_cm)), 'numeric_ok'] = False
    accepted = d.status.isin(['OK', 'ACCEPTED', 'REVIEW_METADATA_OR_UPSTREAM', 'REVIEW_SHAPE_CLEANUP'])
    upstream_bad = d.upstream_qc.str.contains('REVIEW|INFERRED|CONFLICT|DISAGREEMENT', case=False, regex=True)
    metadata_ok = d.metadata_warnings.isin(['', '[]', '["NEIGHBORS_NOT_CERTIFIED_COMPLETE"]'])
    d['analysis_ok'] = d.numeric_ok & d.feature_qc.eq('QC_PASS') & accepted & ~upstream_bad & metadata_ok & ~d.ocr_action.eq('REVIEW')
    d.loc[d.observation_view.eq('top') & ~d.selected_hard_valid.eq(1), 'analysis_ok'] = False

    def reasons(r):
        result = []
        if not r.numeric_ok:
            result.append('Missing or invalid measurements')
        if r.numeric_ok and r.feature_qc != 'QC_PASS':
            result.append('Mask feature quality warning')
        if r.ocr_action == 'REVIEW' or 'OCR_IDENTITY_REVIEW' in r.upstream_qc:
            result.append('Review OCR identity')
        if 'NEIGHBORS_NOT_CERTIFIED_COMPLETE' in r.metadata_warnings:
            result.append('Neighbor-boundary validation incomplete')
        if 'INFERRED' in r.upstream_qc or 'SLOT_COUNT_REVIEW' in r.upstream_qc:
            result.append('Review slot location or count')
        if not accepted.loc[r.name]:
            result.append('Review crop ownership, shape or recovery')
        if r.status == 'REVIEW_SHAPE_CLEANUP':
            result.append('Review cleaned mask shape')
        if r.scale_qc:
            result.append('Review scale-reference geometry')
        return ';'.join(result)
    d['analysis_reason'] = d.apply(reasons, axis=1)
    return d

def build_pairs(d):
    keys = ['pair_id', 'season_id', 'day', 'bed_id', 'physical_slot', 'capture_time', 'track_id']
    top = d.loc[d.observation_view.eq('top')].copy()
    front = d.loc[d.observation_view.eq('front')].copy()
    top = top.rename(columns={c: c + '_top' for c in top if c not in keys})
    front = front.rename(columns={c: c + '_front' for c in front if c not in keys})
    p = top.merge(front, on=keys, how='outer', validate='one_to_one')
    for c in ['numeric_ok_top', 'numeric_ok_front', 'analysis_ok_top', 'analysis_ok_front']:
        p[c] = p[c].eq(True)
    p['paired_numeric'] = p.numeric_ok_top & p.numeric_ok_front
    p['paired_usable'] = p.analysis_ok_top & p.analysis_ok_front
    return p

def median(sub, col):
    return float(sub[col].median()) if not sub.empty and sub[col].notna().any() else np.nan

def build_daily(d, p, config):
    rows = []
    keys = ['season_id', 'day', 'bed_id', 'physical_slot', 'track_id']
    groups = p.groupby(keys, sort=True)
    for ((season, day, bed, slot, track), g) in tqdm(groups, desc='Daily paired observations'):
        used = g.loc[g.paired_usable]
        numeric = g.loc[g.paired_numeric]
        chosen = used if not used.empty else numeric
        paired = not chosen.empty
        t = chosen if paired else g.loc[g.numeric_ok_top]
        f = chosen if paired else g.loc[g.numeric_ok_front]
        a = median(t, 'area_cm2_top')
        h = median(f, 'front_height_cm_front')
        cams = unique_text(t.camera_top)
        version = unique_text(t.scale_version_top) + ' / ' + unique_text(f.scale_version_front)
        scales = unique_text(pd.concat([t.scale_status_top, f.scale_status_front]))
        scalar = 'MISSING' if not (finite(a) or finite(h)) else 'ESTIMATED_REVIEW' if 'ESTIMATED' in scales else 'VERIFIED'
        eligible = not used.empty and ';' not in cams and (t.scale_version_top.nunique() == 1) and (f.scale_version_front.nunique() == 1)
        rep = (chosen if paired else g).sort_values(['capture_time', 'pair_id']).iloc[-1]
        reason = unique_text(pd.concat([g.analysis_reason_top, g.analysis_reason_front]))
        if ';' in cams:
            reason += ';Mixed top-view cameras on the same day'
        if not paired:
            reason += ';Insufficient valid top/front pairs at the same capture time'
        if paired and used.empty:
            reason += ';Observed values fail time-series or peer quality requirements'
        raw_n = int(g.record_id_top.notna().sum() + g.record_id_front.notna().sum())
        slotqc = 'WEAK' if eligible else 'MISSING' if not paired else 'UNUSABLE'
        row = {'cycle_id': season, 'date': day, 'bed_name': bed, 'slot_name': slot, 'plant_age_days': (day - pd.Timestamp(config['transplant_date'])).days if config.get('transplant_date') else np.nan, 'source_pair_id': f'{season}|{day:%Y%m%d}|{bed}|{slot}|daily', 'area_cm2_top': a, 'front_height_cm': h, 'volume_proxy_top_front': a * h if paired else np.nan, 'circularity_top': median(t, 'circularity_top'), 'circularity_front': median(f, 'circularity_front'), 'area_cm2_front': median(f, 'area_cm2_front'), 'solidity_top': median(t, 'solidity_top'), 'solidity_front': median(f, 'solidity_front'), 'track_id': track, 'top_camera': cams, 'scale_signature': version + ' / ' + scales, 'measurement_basis': 'PAIRED_DAILY_MEDIAN' if not used.empty else 'REVIEW_PAIRED_DAILY_MEDIAN' if paired else 'UNPAIRED_DAILY_VALUES', 'analysis_eligible': eligible, 'slot_qc': slotqc, 'scale_qc': scalar, 'n_source_records': raw_n, 'n_capture_pairs': len(g), 'n_numeric_pairs': len(numeric), 'n_used_pairs': len(used), 'all_pair_ids': json_list(g.pair_id), 'used_pair_ids': json_list(used.pair_id), 'top_record_ids': json_list(g.record_id_top), 'front_record_ids': json_list(g.record_id_front), 'representative_pair_id': rep.pair_id, 'representative_top_crop': rep.crop_path_top, 'representative_front_crop': rep.crop_path_front, 'source_status_top': unique_text(g.status_top), 'source_status_front': unique_text(g.status_front), 'source_scale_qc': unique_text(pd.concat([g.scale_qc_top, g.scale_qc_front])), 'scale_source_top': unique_text(t.scale_source_top), 'scale_source_front': unique_text(f.scale_source_front), 'mm_per_px_top_median': median(t, 'mm_per_px_top'), 'mm_per_px_front_median': median(f, 'mm_per_px_front'), 'qc_reason': reason.strip(';'), 'predicted_weight_g': np.nan, 'predicted_stage': None, 'harvest_status': None}
        rows.append(row)
    result = pd.DataFrame(rows)
    result['_slot_row'] = result.slot_name.str[0]
    result['_slot_number'] = result.slot_name.str[1:].astype(int)
    return result.sort_values(['date', 'bed_name', '_slot_row', '_slot_number']).drop(columns=['_slot_row', '_slot_number']).reset_index(drop=True)

def window_metrics(sub, config):
    n = len(sub)
    if n < config['min_valid_days']:
        return None
    days = int((sub.date.iloc[-1] - sub.date.iloc[0]).days)
    if days < config['min_span_days']:
        return None
    result = {'valid_obs_n': n, 'growth_window_days': days, 'window_start_date': sub.date.iloc[0]}
    for (col, prefix) in [('area_cm2_top', 'area'), ('volume_proxy_top_front', 'volume'), ('area_cm2_front', 'front_area')]:
        (first, last) = (float(sub[col].iloc[0]), float(sub[col].iloc[-1]))
        if first <= 0 or last <= 0 or (not (finite(first) and finite(last))):
            return None
        result[prefix + '_growth_pct_day'] = 100 * np.expm1(np.log(last / first) / days)
        result[prefix + '_change_pct'] = 100 * (last / first - 1)
        result[prefix + '_window_start'] = first
    result['height_growth_cm_day'] = (sub.front_height_cm.iloc[-1] - sub.front_height_cm.iloc[0]) / days
    result['height_window_start'] = sub.front_height_cm.iloc[0]
    for c in ['circularity_top', 'circularity_front', 'solidity_front']:
        result[c + '_delta'] = sub[c].iloc[-1] - sub[c].iloc[0]
    return result

def add_temporal(d, config):
    out = d.copy()
    out['valid_obs_n'] = 0
    out['growth_window_days'] = np.nan
    out['timeseries_qc'] = 'INSUFFICIENT'
    out['temporal_segment_id'] = ''
    for c in ['area_growth_pct_day', 'height_growth_cm_day', 'volume_growth_pct_day', 'front_area_growth_pct_day', 'area_change_pct', 'volume_change_pct', 'front_area_change_pct', 'area_window_start', 'volume_window_start', 'front_area_window_start', 'height_window_start', 'circularity_top_delta', 'circularity_front_delta', 'solidity_front_delta', 'previous_volume_growth_pct_day']:
        out[c] = np.nan
    out['window_start_date'] = pd.NaT
    for (track, sub) in tqdm(out.groupby('track_id', sort=False), desc='Past-only growth windows'):
        seg = 0
        history = []
        last_signature = None
        last_date = None
        for (idx, r) in sub.sort_values('date').iterrows():
            out.at[idx, 'temporal_segment_id'] = f'{track}|{seg}'
            if not r.analysis_eligible:
                continue
            signature = (r.top_camera, r.scale_signature)
            if last_signature is not None and signature != last_signature:
                seg += 1
                history = []
            if last_date is not None and (r.date - last_date).days > config['max_gap_days']:
                seg += 1
                history = []
            out.at[idx, 'temporal_segment_id'] = f'{track}|{seg}'
            history.append(idx)
            last_signature = signature
            last_date = r.date
            cutoff = r.date - pd.Timedelta(days=2 * config['window_days'] - 1)
            history = [i for i in history if out.at[i, 'date'] >= cutoff]
            h = out.loc[history]
            w = h.loc[h.date.ge(r.date - pd.Timedelta(days=config['window_days'] - 1))]
            out.at[idx, 'valid_obs_n'] = len(w)
            result = window_metrics(w, config)
            if result:
                for (c, v) in result.items():
                    out.at[idx, c] = v
                out.at[idx, 'timeseries_qc'] = 'GOOD' if len(w) >= config['good_valid_days'] else 'USABLE'
                prev = h.loc[(h.date >= r.date - pd.Timedelta(days=2 * config['window_days'] - 1)) & (h.date < r.date - pd.Timedelta(days=config['window_days'] - 1))]
                pr = window_metrics(prev, config)
                if pr:
                    out.at[idx, 'previous_volume_growth_pct_day'] = pr['volume_growth_pct_day']
    return out

def add_peers(d, config):
    out = d.copy()
    out['peer_group_definition'] = [f'{r.cycle_id}|{r.cohort_id}|{r.date:%Y%m%d}|{r.slot_name[0]}|{r.top_camera}|{r.scale_signature}' for r in out.itertuples()]
    out['peer_group_id'] = out.peer_group_definition.map(lambda s: 'PG_' + hashlib.sha256(s.encode()).hexdigest()[:14])
    for c in ['peer_n', 'peer_bed_n', 'growth_peer_n', 'growth_peer_bed_n']:
        out[c] = 0
    for c in ['peer_volume_mean', 'peer_volume_sd', 'peer_growth_mean', 'peer_growth_sd', 'volume_vs_peer_pct', 'growth_vs_peer_pp_day', 'top_area_z', 'front_area_z', 'front_top_gap_z']:
        out[c] = np.nan
    out['relative_growth_status'] = 'NOT_EVALUABLE'
    for (_, g) in tqdm(out.groupby('peer_group_id', sort=False), desc='Equal-plant reference groups'):
        eligible = g.loc[g.analysis_eligible & g.cohort_confirmed & g.volume_proxy_top_front.gt(0)]
        if eligible.empty:
            continue
        for (idx, r) in g.iterrows():
            peers = eligible.loc[eligible.track_id.ne(r.track_id)]
            bed_n = peers.bed_name.nunique()
            out.at[idx, 'peer_n'] = len(peers)
            out.at[idx, 'peer_bed_n'] = bed_n
            ready = len(peers) >= config['min_peer_n'] and bed_n >= config['min_peer_beds']
            if not ready or not r.analysis_eligible:
                continue
            mu = peers.volume_proxy_top_front.mean()
            sd = peers.volume_proxy_top_front.std()
            out.at[idx, 'peer_volume_mean'] = mu
            out.at[idx, 'peer_volume_sd'] = sd
            out.at[idx, 'volume_vs_peer_pct'] = 100 * (r.volume_proxy_top_front / mu - 1)
            for (col, dest) in [('area_cm2_top', 'top_area_z'), ('area_cm2_front', 'front_area_z')]:
                mean = peers[col].mean()
                std = peers[col].std()
                if finite(std) and std > 1e-12:
                    out.at[idx, dest] = (r[col] - mean) / std
            out.at[idx, 'front_top_gap_z'] = out.at[idx, 'front_area_z'] - out.at[idx, 'top_area_z']
            gp = peers.loc[peers.volume_growth_pct_day.notna() & peers.window_start_date.eq(r.window_start_date)]
            growth_bed_n = gp.bed_name.nunique()
            out.at[idx, 'growth_peer_n'] = len(gp)
            out.at[idx, 'growth_peer_bed_n'] = growth_bed_n
            if not finite(r.volume_growth_pct_day) or len(gp) < config['min_peer_n'] or growth_bed_n < config['min_peer_beds']:
                continue
            gm = gp.volume_growth_pct_day.mean()
            gs = gp.volume_growth_pct_day.std()
            out.at[idx, 'peer_growth_mean'] = gm
            out.at[idx, 'peer_growth_sd'] = gs
            delta = r.volume_growth_pct_day - gm
            out.at[idx, 'growth_vs_peer_pp_day'] = delta
            if finite(gs) and gs > 1e-12:
                band = config['peer_band_sd'] * gs
                out.at[idx, 'relative_growth_status'] = 'FAST' if delta > band else 'SLOW' if delta < -band else 'TYPICAL'
    return out

def direction(v, threshold):
    if not finite(v):
        return 'UNKNOWN'
    return 'UP' if v > threshold else 'DOWN' if v < -threshold else 'FLAT'

def add_labels(d, config):
    out = d.copy()
    out['growth_qc'] = 'INSUFFICIENT_DATA'
    out['relation_qc'] = 'NOT_EVALUABLE'
    out['gap_history_std'] = np.nan
    out['gap_persistent_count'] = 0
    th = config['size_growth_threshold_pct_day']
    ht = config['height_growth_threshold_cm_day']
    gapth = config['relation_gap_z_threshold']
    for (idx, r) in out.iterrows():
        if not r.analysis_eligible or not finite(r.volume_growth_pct_day):
            continue
        a = direction(r.area_growth_pct_day, th)
        h = direction(r.height_growth_cm_day, ht)
        v = direction(r.volume_growth_pct_day, th)
        f = direction(r.front_area_growth_pct_day, th)
        structure = finite(r.circularity_front_delta) and abs(r.circularity_front_delta) >= config['structure_delta_threshold'] or (finite(r.solidity_front_delta) and abs(r.solidity_front_delta) >= config['structure_delta_threshold'])
        if a == 'FLAT' and (h != 'FLAT' or structure):
            label = 'STRUCTURAL_TRANSITION_SIGNAL'
        elif v == 'DOWN':
            label = 'DECREASE_OBSERVED'
        elif v == 'FLAT' and a == 'FLAT' and (h == 'FLAT') and (r.valid_obs_n >= config['good_valid_days']):
            label = 'PLATEAU_CANDIDATE'
        elif v == 'UP' and a != 'DOWN' and (h != 'DOWN'):
            prev = r.previous_volume_growth_pct_day
            label = 'GROWTH_SLOWING' if finite(prev) and prev > th and (r.volume_growth_pct_day < prev * config['slowing_ratio']) else 'GROWTH_OBSERVED'
        else:
            label = 'MIXED_CHANGE'
        out.at[idx, 'growth_qc'] = label
        if a != f:
            relation = 'DIRECTION_DIFFERENCE'
        elif finite(r.front_top_gap_z):
            relation = 'FRONT_DEFICIT' if r.front_top_gap_z <= -gapth else 'FRONT_EXCESS' if r.front_top_gap_z >= gapth else 'COHERENT'
        else:
            relation = 'NOT_EVALUABLE'
        out.at[idx, 'relation_qc'] = relation
    for (_, g) in out.groupby('temporal_segment_id', sort=False):
        g = g.sort_values('date')
        for (idx, r) in g.iterrows():
            if not r.analysis_eligible or not finite(r.volume_growth_pct_day):
                continue
            recent = g.loc[(g.date <= r.date) & (g.date >= r.date - pd.Timedelta(days=config['window_days'] - 1))]
            valid = recent.loc[recent.front_top_gap_z.notna() & recent.analysis_eligible]
            if len(valid) >= config['min_valid_days']:
                std = valid.front_top_gap_z.std()
                out.at[idx, 'gap_history_std'] = std
                if std >= config['relation_gap_std_threshold']:
                    out.at[idx, 'relation_qc'] = 'UNSTABLE_RELATION'
            direction_now = -1 if finite(r.front_top_gap_z) and r.front_top_gap_z <= -gapth else 1 if finite(r.front_top_gap_z) and r.front_top_gap_z >= gapth else 0
            count = 0
            last = r.date
            if direction_now:
                for (_, rr) in recent.iloc[::-1].iterrows():
                    if (last - rr.date).days > 1 or not rr.analysis_eligible or (not finite(rr.front_top_gap_z)):
                        break
                    if direction_now * rr.front_top_gap_z < gapth:
                        break
                    count += 1
                    last = rr.date
            out.at[idx, 'gap_persistent_count'] = count
    return out

def make_message(r):
    """Summarize measured image proxies without diagnosing biological status."""
    parts = []
    if finite(r.predicted_weight_g):
        parts.append(f'Predicted weight: {r.predicted_weight_g:.0f} g.')
    if pd.notna(r.predicted_stage):
        parts.append(f'Predicted stage: {r.predicted_stage}.')
    if finite(r.volume_proxy_top_front):
        parts.append(f'Top area: {r.area_cm2_top:.1f} cm2; front height: {r.front_height_cm:.1f} cm; volume proxy: {r.volume_proxy_top_front:.0f} cm3.')
    else:
        parts.append('No valid paired-view volume proxy.')
    if finite(r.volume_vs_peer_pct):
        parts.append(f'Volume relative to {r.peer_n} matched peers: {r.volume_vs_peer_pct:+.1f}%.')
    else:
        parts.append('Matched-peer size comparison unavailable.')
    if finite(r.volume_growth_pct_day):
        parts.append(f'Volume change: {r.volume_growth_pct_day:+.2f}% per day over {int(r.growth_window_days)} days ({r.valid_obs_n} valid days).')
        parts.append(GROWTH_TEXT[r.growth_qc] + '. ' + RELATION_TEXT[r.relation_qc] + '.')
    else:
        parts.append('Growth trend unavailable: insufficient valid consecutive observations.')
    if finite(r.growth_vs_peer_pp_day):
        parts.append(f'Growth-rate difference from peers: {r.growth_vs_peer_pp_day:+.2f} percentage points per day.')
    if str(r.qc_reason).strip():
        parts.append('Review: ' + str(r.qc_reason) + '.')
    if r.scale_qc == 'ESTIMATED_REVIEW':
        parts.append('Size estimates use unverified scale.')
    parts.append('QC uses provisional image-based rules.')
    return '\n'.join(parts)

def input_sheets(wb):
    names = [s for s in wb.sheetnames if re.fullmatch('전체(?:_\\d+)?', s)]
    if not names:
        raise ValueError('Invalid input or state in input_sheets (rejected condition: not names)')
    return sorted(names, key=lambda s: 1 if s == '전체' else int(s.split('_')[-1]))

def iter_feature_chunks(path, chunk_rows=5000):
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        for name in input_sheets(wb):
            rows = wb[name].iter_rows(values_only=True)
            header = next(rows, None)
            if not header or len(set(header)) != len(header):
                raise ValueError('Invalid input or state in iter_feature_chunks (rejected condition: not header or len(set(header)) != len(header))')
            missing = set(SOURCE_COLUMNS) - set(header)
            if missing:
                raise ValueError('Invalid input or state in iter_feature_chunks (rejected condition: missing)')
            positions = [header.index(c) for c in SOURCE_COLUMNS]
            batch = []
            for row in rows:
                if all((v is None for v in row)):
                    continue
                batch.append([row[i] if i < len(row) else None for i in positions])
                if len(batch) >= chunk_rows:
                    yield pd.DataFrame(batch, columns=SOURCE_COLUMNS)
                    batch = []
            if batch:
                yield pd.DataFrame(batch, columns=SOURCE_COLUMNS)
    finally:
        wb.close()

def attach_cohorts(d, cohorts):
    out = d.copy()
    mapping = {}
    for group in cohorts:
        for bed in group['beds']:
            key = (group['cycle_id'], bed)
            if key in mapping:
                raise ValueError('Invalid input or state in attach_cohorts (rejected condition: key in mapping)')
            mapping[key] = group
    out['cohort_id'] = ''
    out['cohort_confirmed'] = False
    for (idx, r) in out.iterrows():
        match = mapping.get((r.cycle_id, r.bed_name))
        out.at[idx, 'cohort_id'] = match['cohort_id'] if match else 'UNCONFIRMED:' + r.bed_name
        out.at[idx, 'cohort_confirmed'] = match is not None
        if match and match.get('transplant_date'):
            age = (r.date - pd.Timestamp(match['transplant_date'])).days
            if age < 0:
                raise ValueError('Invalid input or state in attach_cohorts (rejected condition: age < 0)')
            out.at[idx, 'plant_age_days'] = age
    return out

def compute_qc(path, config, db_path=':memory:', cohorts=None):
    conn = sqlite3.connect(str(db_path))
    columns = ','.join(('"' + c + '" TEXT' for c in SOURCE_COLUMNS))
    conn.execute('CREATE TABLE raw (' + columns + ', UNIQUE(record_id), UNIQUE(pair_id,observation_view))')
    conn.execute('CREATE INDEX raw_bed ON raw(season_id,bed_id)')
    count = 0
    try:
        for batch in tqdm(iter_feature_chunks(path, config['chunk_rows']), desc='feature XLSX chunks'):
            values = []
            for row in batch.itertuples(index=False, name=None):
                values.append(tuple(('' if v is None or (isinstance(v, float) and math.isnan(v)) else v.strftime('%Y%m%d') if isinstance(v, (datetime.datetime, datetime.date, pd.Timestamp)) else str(v) for v in row)))
            try:
                conn.executemany('INSERT INTO raw VALUES (' + ','.join(('?' for _ in SOURCE_COLUMNS)) + ')', values)
            except sqlite3.IntegrityError as exc:
                raise ValueError('Invalid input or state in compute_qc') from exc
            count += len(values)
        conn.commit()
        if count == 0:
            raise ValueError('Invalid input or state in compute_qc (rejected condition: count == 0)')
        conflicts = ' OR '.join(('COUNT(DISTINCT "' + key + '")>1' for key in ['season_id', 'bed_id', 'physical_slot', 'capture_date', 'capture_time']))
        clash = conn.execute('SELECT pair_id FROM raw GROUP BY pair_id HAVING ' + conflicts + ' LIMIT 1').fetchone()
        if clash:
            raise ValueError('Invalid input or state in compute_qc (rejected condition: clash)')
        parts = []
        pair_n = 0
        groups = conn.execute('SELECT DISTINCT season_id,bed_id FROM raw ORDER BY season_id,bed_id').fetchall()
        for (season, bed) in tqdm(groups, desc='bed → daily QC'):
            records = conn.execute('SELECT * FROM raw INDEXED BY raw_bed WHERE season_id=? AND bed_id=?', (season, bed)).fetchall()
            source = normalize_features(pd.DataFrame.from_records(records, columns=SOURCE_COLUMNS))
            pairs = build_pairs(source)
            daily = build_daily(source, pairs, config)
            pair_n += len(pairs)
            parts.append(daily)
        d = pd.concat(parts, ignore_index=True)
    finally:
        conn.close()
    d = attach_cohorts(d, CONFIRMED_COHORTS if cohorts is None else cohorts)
    d = add_temporal(d, config)
    d = add_peers(d, config)
    d = add_labels(d, config)
    d['final_message'] = d.apply(make_message, axis=1)
    assert int(d.n_source_records.sum()) == count and int(d.n_capture_pairs.sum()) == pair_n
    assert d.source_pair_id.is_unique and (not d.duplicated(['cycle_id', 'date', 'bed_name', 'slot_name']).any())
    assert not np.isinf(d.select_dtypes(include='number').to_numpy()).any()
    assert d.loc[~d.analysis_eligible, ['volume_growth_pct_day', 'volume_vs_peer_pct']].isna().all().all()
    d['_row'] = d.slot_name.str[0]
    d['_n'] = d.slot_name.str[1:].astype(int)
    d = d.sort_values(['cycle_id', 'date', 'bed_name', '_row', '_n']).drop(columns=['_row', '_n']).reset_index(drop=True)
    return (d, {'input_rows': count, 'pair_rows': pair_n, 'daily_rows': len(d), 'growth_rows': int(d.volume_growth_pct_day.notna().sum()), 'unconfirmed_cohort_rows': int((~d.cohort_confirmed).sum())})

def merge_predictions(d, path=None):
    out = d.copy()
    for col in ['prediction_source', 'prediction_model_version']:
        out[col] = ''
    if not path:
        return (out, {'provided': False, 'matched_rows': 0})
    path = Path(path)
    if not path.is_file():
        raise FileNotFoundError('Invalid input or state in merge_predictions (rejected condition: not path.is_file())')
    pred = pd.read_excel(path, dtype=object) if path.suffix.lower() == '.xlsx' else pd.read_csv(path, dtype=object)
    value_cols = [c for c in ['predicted_weight_g', 'predicted_stage', 'harvest_status'] if c in pred]
    if not value_cols:
        raise ValueError('Invalid input or state in merge_predictions (rejected condition: not value_cols)')
    if 'source_pair_id' not in pred:
        keys = ['cycle_id', 'date', 'bed_name', 'slot_name']
        if not set(keys) <= set(pred):
            raise ValueError('Invalid input or state in merge_predictions (rejected condition: not set(keys) <= set(pred))')
        if pred[keys].isna().any().any():
            raise ValueError('Invalid input or state in merge_predictions (rejected condition: pred[keys].isna().any().any())')
        dt = pd.to_datetime(pred['date'].astype(str).str.replace('\\.0$', '', regex=True), errors='raise')
        pred['source_pair_id'] = pred.cycle_id.astype(str) + '|' + dt.dt.strftime('%Y%m%d') + '|' + pred.bed_name.astype(str) + '|' + pred.slot_name.astype(str) + '|daily'
    if pred.source_pair_id.isna().any() or pred.source_pair_id.astype(str).str.strip().eq('').any() or pred.source_pair_id.duplicated().any():
        raise ValueError("Invalid input or state in merge_predictions (rejected condition: pred.source_pair_id.isna().any() or pred.source_pair_id.astype(str).str.strip().eq('').any() or pred.source_pair_id.duplicated().any())")
    if 'predicted_weight_g' in pred:
        raw = pred.predicted_weight_g
        v = pd.to_numeric(raw, errors='coerce')
        if (raw.notna() & (~np.isfinite(v) | v.lt(0))).any():
            raise ValueError('Invalid input or state in merge_predictions (rejected condition: (raw.notna() & (~np.isfinite(v) | v.lt(0))).any())')
        pred['predicted_weight_g'] = v
    if 'harvest_status' in pred:
        if not pred.harvest_status.dropna().isin(['READY', 'NOT_READY', 'UNKNOWN']).all():
            raise ValueError("Invalid input or state in merge_predictions (rejected condition: not pred.harvest_status.dropna().isin(['READY', 'NOT_READY', 'UNKNOWN']).all())")
    if 'predicted_stage' in pred:
        pred['predicted_stage'] = pred.predicted_stage.map(lambda v: str(v) if pd.notna(v) else None)
    lookup = pred.set_index('source_pair_id')
    for col in value_cols:
        out[col] = out.source_pair_id.map(lookup[col])
    matched = out.source_pair_id.isin(lookup.index)
    out.loc[matched, 'prediction_source'] = str(path.resolve())
    if 'model_version' in pred:
        out['prediction_model_version'] = out.source_pair_id.map(lookup.model_version).fillna('')
    out['final_message'] = out.apply(make_message, axis=1)
    words = {'READY': 'eligible', 'NOT_READY': 'not eligible', 'UNKNOWN': 'unavailable'}
    for (idx, r) in out.loc[out.harvest_status.notna()].iterrows():
        out.at[idx, 'final_message'] = 'Model harvest prediction: ' + words[r.harvest_status] + '.\n' + r.final_message
    return (out, {'provided': True, 'matched_rows': int(matched.sum()), 'unmatched_prediction_rows': int((~pred.source_pair_id.isin(out.source_pair_id)).sum())})


def write_qc_excel(d, target):
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    wb = Workbook(write_only=True)
    ordered = OP_COLUMNS + [c for c in d if c not in OP_COLUMNS]
    for (name, cols) in [('operational_output', OP_COLUMNS), ('all_required_columns', ordered)]:
        if len(d) > 1048575:
            raise ValueError('Invalid input or state in _write_qc_excel (rejected condition: len(d) > 1048575)')
        ws = wb.create_sheet(name)
        ws.freeze_panes = 'G2'
        for (i, col) in enumerate(cols, 1):
            ws.column_dimensions[get_column_letter(i)].width = 80 if col == 'final_message' else 30 if col.endswith('_id') else max(15, min(28, len(col) + 2))
        header = []
        for col in cols:
            cell = WriteOnlyCell(ws, col)
            cell.font = Font(name='Calibri', color='FFFFFF', bold=True)
            cell.fill = PatternFill('solid', fgColor='215B69')
            header.append(cell)
        ws.append(header)
        ws.row_dimensions[1].height = 30
        for row in d[cols].itertuples(index=False, name=None):
            cells = []
            for val in row:
                if val is None or (not isinstance(val, (list, dict)) and pd.isna(val)):
                    val = None
                if isinstance(val, np.generic):
                    val = val.item()
                if isinstance(val, pd.Timestamp):
                    val = val.to_pydatetime()
                cell = WriteOnlyCell(ws, val)
                if isinstance(val, str):
                    if len(val) > 32767:
                        raise ValueError('Excel string length limit exceeded')
                    cell.data_type = 's'
                elif isinstance(val, datetime.datetime):
                    cell.number_format = 'yyyy-mm-dd'
                elif isinstance(val, float):
                    cell.number_format = '0.00'
                cells.append(cell)
            ws.append(cells)
        ws.auto_filter.ref = f'A1:{get_column_letter(len(cols))}{len(d) + 1}'
    wb.save(target)

def check_excel(target, d, full=True):
    from openpyxl import load_workbook
    counts = None
    if not full:
        from output_inventory import _xlsx_rows
        counts = _xlsx_rows(Path(target))
    if hasattr(target, 'seek'):
        target.seek(0)
    wb = load_workbook(target, read_only=True, data_only=True)
    try:
        if wb.sheetnames != ['operational_output', 'all_required_columns']:
            raise AssertionError("Invalid input or state in check_excel (rejected condition: wb.sheetnames != ['operational_output', 'all_required_columns'])")
        for name in wb.sheetnames:
            rows = wb[name].iter_rows(values_only=True)
            header = next(rows)
            if counts is not None:
                ordered = OP_COLUMNS + [c for c in d if c not in OP_COLUMNS]
                assert list(header) == (OP_COLUMNS if name == 'operational_output' else ordered), 'Excel header mismatch'
                assert counts[name] == len(d), 'Excel row count mismatch'
                continue
            expected = d[list(header)]
            n = 0
            for wanted in expected.itertuples(index=False, name=None):
                actual = next(rows, None)
                assert actual is not None, 'Missing Excel rows'
                actual = actual + (None,) * (len(wanted) - len(actual))
                assert len(actual) == len(wanted), 'Excel column count mismatch'
                for (a, b) in zip(actual, wanted):
                    if b is None or pd.isna(b):
                        assert a is None
                    elif isinstance(b, (float, np.floating)):
                        assert a is not None and math.isclose(a, float(b), rel_tol=1e-10, abs_tol=1e-10)
                    elif b == '':
                        assert a in ('', None)
                    else:
                        assert a == b, (a, b)
                n += 1
            assert n == len(d) and next(rows, None) is None
    finally:
        wb.close()

def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--input', type=Path)
    parser.add_argument('--output', type=Path)
    parser.add_argument('--predictions', type=Path, help='Optional saved predictions; QC does not determine harvest eligibility')
    parser.add_argument('--check', action='store_true')
    args = parser.parse_args(argv)
    (source, default_output) = default_paths(discover=args.input is None)
    source = args.input or source
    target = args.output or default_output
    if not source.is_file():
        raise FileNotFoundError(source)
    before = sha256(source)
    if args.check:
        workbook = load_workbook(source, read_only=True, data_only=True)
        try:
            for name in input_sheets(workbook):
                header = next(workbook[name].iter_rows(values_only=True))
                missing = set(SOURCE_COLUMNS) - set(header)
                if missing:
                    raise ValueError(f'Missing feature columns: {sorted(missing)}')
        finally:
            workbook.close()
        print(json.dumps({'status': 'READY', 'input': str(source), 'writes': 0}))
        return 0
    target.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix='.qc_work_', dir=target) as temp:
        work = Path(temp)
        (data, report) = compute_qc(source, dict(DEFAULT_CONFIG), work / 'source.sqlite3')
        (data, prediction_report) = merge_predictions(data, args.predictions)
        output = work / 'growth_qc.xlsx'
        write_qc_excel(data, output)
        check_excel(output, data)
        if sha256(source) != before:
            raise RuntimeError('Source changed during QC calculation')
        output.replace(target / 'growth_qc.xlsx')
    report.update(source_sha256=before, source_unchanged=True, predictions=prediction_report)
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0
if __name__ == '__main__':
    raise SystemExit(main())
