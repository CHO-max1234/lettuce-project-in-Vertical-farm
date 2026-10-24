"""Processing utilities for final xlsx."""
from __future__ import annotations
import sys as _sys
from pathlib import Path as _Path
_CYCLE_ROOT = _Path(__file__).resolve().parents[1]
if str(_CYCLE_ROOT) not in _sys.path:
    _sys.path.insert(0, str(_CYCLE_ROOT))
import csv
import json
import math
import os
from collections import Counter
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from stage_inputs import image_key, norm, yes
COLUMNS = ['lettuce_key', 'capture_key', 'date', 'bed_name', 'slot_name', 'growth_qc', 'relation_qc', 'timeseries_qc', 'area_cm2', 'height_cm', 'volume_proxy', 'circularity', 'harvest_decision', 'pred_weight', 'cnn_pred_stage', 'cnn_p0', 'cnn_p1', 'cnn_p2', 'cnn_p3', 'cnn_p4', 'harvest_ornot', 'fresh_weight', 'message']

def number(value: object) -> float | None:
    try:
        result = float(value)
        return result if math.isfinite(result) else None
    except (TypeError, ValueError):
        return None

def daily_qc(path: Path, targets: set[str]) -> dict[str, dict]:
    book = load_workbook(path, read_only=True, data_only=True)
    found = {}
    try:
        rows = book.worksheets[0].iter_rows(values_only=True)
        header = [norm(x) for x in next(rows)]
        fields = ('source_pair_id', 'growth_qc', 'relation_qc', 'timeseries_qc')
        indexes = {field: header.index(field) for field in fields}
        for values in rows:
            parts = norm(values[indexes['source_pair_id']]).split('|')
            if len(parts) != 5:
                continue
            try:
                key = image_key(parts[1], parts[2], parts[3])
            except ValueError:
                continue
            if key not in targets:
                continue
            if key in found:
                raise ValueError(f'Duplicate QC key: {key}')
            found[key] = {field: norm(values[indexes[field]]) for field in fields[1:]}
    finally:
        book.close()
    if set(found) != targets:
        raise ValueError(f'Invalid input or state in daily_qc (rejected condition: set(found) != targets): {len(targets - set(found))} ')
    return found

def exact_front(front_csv: Path, gate_csv: Path) -> dict[str, float]:
    passed = set()
    with gate_csv.open(encoding='utf-8-sig', newline='') as handle:
        for row in csv.DictReader(handle):
            if yes(row['model_eligible']):
                passed.add(row['record_id'])
    pairs: dict[str, float] = {}
    ambiguous = set()
    with front_csv.open(encoding='utf-8-sig', newline='') as handle:
        for row in csv.DictReader(handle):
            if row['record_id'] not in passed:
                continue
            pair = row['pair_id']
            height = number(row['front_height_cm'])
            if not pair or height is None or height <= 0:
                continue
            if pair in pairs:
                ambiguous.add(pair)
            else:
                pairs[pair] = height
    for pair in ambiguous:
        del pairs[pair]
    return pairs

def label_weights(path: Path) -> dict[str, float]:
    values = {}
    with path.open(encoding='utf-8-sig', newline='') as handle:
        for row in csv.DictReader(handle):
            if row['status'] == 'OK':
                values[row['lettuce_key']] = float(row['weight_g'])
    return values

def representative_predictions(path: Path, targets: set[str], fronts: dict[str, float]) -> tuple[dict[str, dict], dict]:
    """Select representative images by upstream quality and capture time, without using prediction error."""
    chosen: dict[str, tuple[tuple, dict]] = {}
    counts = Counter()
    with path.open(encoding='utf-8-sig', newline='') as handle:
        reader = csv.DictReader(handle)
        required = {'record_id', 'capture_date', 'capture_time', 'bed_id', 'physical_slot', 'pair_id'}
        if required - set(reader.fieldnames or []):
            raise ValueError(f'Missing A-step output columns: {sorted(required - set(reader.fieldnames or []))}')
        for row in reader:
            counts['a_step_top_rows'] += 1
            try:
                key = image_key(row['capture_date'], row['bed_id'], row['physical_slot'])
            except ValueError:
                counts['bad_image_key'] += 1
                continue
            if key not in targets:
                continue
            counts['measured_prediction_rows'] += 1
            quality = number(row.get('selected_quality_score'))
            time_text = norm(row['capture_time']).removesuffix('.0').zfill(6)
            rank = (int(yes(row.get('분석가능'))), int(row.get('pair_id', '') in fronts), quality if quality is not None else -1000000000000.0, time_text, row['record_id'])
            if key not in chosen or rank > chosen[key][0]:
                chosen[key] = (rank, row)
    return ({key: value[1] for (key, value) in chosen.items()}, dict(counts))

def message(record: dict, paired: bool) -> str:
    weight = 'Unavailable' if record['pred_weight'] is None else f"{record['pred_weight']:.0f} g"
    stage = 'Unavailable' if record['cnn_pred_stage'] is None else str(record['cnn_pred_stage'])
    area = 'Unavailable' if record['area_cm2'] is None else f"{record['area_cm2']:.1f} cm2"
    volume = 'Unavailable' if record['volume_proxy'] is None else f"{record['volume_proxy']:.1f} cm3"
    return f"Predicted weight: {weight}\nPredicted stage: {stage}\nHarvest decision: {record['harvest_decision']}\nTop area: {area}; volume proxy: {volume}; paired views: {paired}. Growth QC: {record['growth_qc']}; relation QC: {record['relation_qc']}; time-series QC: {record['timeseries_qc']}."

def write_final(*, labels_csv: Path, qc_xlsx: Path, front_csv: Path, gate_csv: Path, a_csv: Path, d_csv: Path, targets: set[str], output: Path, report_path: Path) -> dict:
    weights = label_weights(labels_csv)
    if not targets <= set(weights):
        raise ValueError('Invalid input or state in write_final (rejected condition: not targets <= set(weights))')
    qc = daily_qc(qc_xlsx, targets)
    fronts = exact_front(front_csv, gate_csv)
    (representatives, prediction_counts) = representative_predictions(a_csv, targets, fronts)
    chosen_ids = {row['record_id'] for row in representatives.values()}
    d_rows = {}
    with d_csv.open(encoding='utf-8-sig', newline='') as handle:
        for row in csv.DictReader(handle):
            rid = row['record_id']
            if rid in chosen_ids:
                if rid in d_rows:
                    raise ValueError(f'Duplicate representative D-step record_id: {rid}')
                d_rows[rid] = row
    book = Workbook()
    sheet = book.active
    sheet.title = 'Predictions'
    sheet.append(COLUMNS)
    sheet.freeze_panes = 'A2'
    errors = []
    counts = Counter()
    for key in sorted(targets):
        (day, bed, slot) = key.split('_')
        representative = representatives.get(key)
        source = d_rows.get(representative['record_id'], {}) if representative else {}
        pair = source.get('pair_id', '')
        height = fronts.get(pair) if source else None
        area = number(source.get('area_cm2'))
        circularity = number(source.get('circularity'))
        pred = number(source.get('d_step_pred_weight_g')) if source.get('d_step_status') == 'OK' else None
        stage_raw = number(source.get('d_step_pred_stage'))
        stage = int(stage_raw) if stage_raw is not None else None
        harvest = norm(source.get('d_step_harvest_ornot'))
        harvest_num = int(float(harvest)) if harvest in {'0', '1', '0.0', '1.0'} else None
        decision = 'Eligible' if harvest_num == 1 else 'Not eligible' if harvest_num == 0 else 'Unavailable'
        time_text = norm(source.get('capture_time')).removesuffix('.0').zfill(6) if source else ''
        record = {'lettuce_key': key, 'capture_key': f'{day}_{time_text}_{bed}_{slot}' if source else '', 'date': f'20{day[:2]}-{day[2:4]}-{day[4:6]}', 'bed_name': bed, 'slot_name': slot, **qc[key], 'area_cm2': area, 'height_cm': height, 'volume_proxy': area * height if area is not None and height is not None else None, 'circularity': circularity, 'harvest_decision': decision, 'pred_weight': pred, 'cnn_pred_stage': stage, **{f'cnn_p{i}': number(source.get(f'd_step_p{i}')) for i in range(5)}, 'harvest_ornot': harvest_num, 'fresh_weight': weights[key]}
        record['message'] = message(record, paired=record['volume_proxy'] is not None)
        sheet.append([record.get(column) for column in COLUMNS])
        counts['final_rows'] += 1
        counts['predicted_rows'] += pred is not None
        counts['exact_front_pairs'] += record['volume_proxy'] is not None
        counts['unavailable_rows'] += pred is None
        if pred is not None:
            errors.append(pred - weights[key])
    sheet.auto_filter.ref = f'A1:{get_column_letter(len(COLUMNS))}{sheet.max_row}'
    header_fill = PatternFill('solid', fgColor='284761')
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(color='FFFFFF', bold=True)
    for (column, name) in enumerate(COLUMNS, 1):
        width = 75 if name == 'message' else 22 if name in {'lettuce_key', 'capture_key'} else 18
        sheet.column_dimensions[get_column_letter(column)].width = width
    msg_col = COLUMNS.index('message') + 1
    for cells in sheet.iter_rows(min_row=2, min_col=msg_col, max_col=msg_col):
        cells[0].alignment = Alignment(wrap_text=True, vertical='top')
    metrics = {'comparison_keys': len(errors), 'MAE_g': None, 'RMSE_g': None, 'bias_g': None, 'R2': None}
    if errors:
        metrics['MAE_g'] = sum((abs(error) for error in errors)) / len(errors)
        metrics['RMSE_g'] = math.sqrt(sum((error * error for error in errors)) / len(errors))
        metrics['bias_g'] = sum(errors) / len(errors)
        measured = [sheet.cell(row, COLUMNS.index('fresh_weight') + 1).value for row in range(2, sheet.max_row + 1) if sheet.cell(row, COLUMNS.index('pred_weight') + 1).value is not None]
        mean = sum(measured) / len(measured)
        denominator = sum(((actual - mean) ** 2 for actual in measured))
        metrics['R2'] = 1 - sum((error * error for error in errors)) / denominator if denominator > 0 else None
    summary = book.create_sheet('Summary')
    summary.append(['Item', 'Value'])
    for (name, value) in {**dict(counts), **metrics}.items():
        summary.append([name, value])
    summary.append(['Scope', 'Image keys linked to measured weights; frozen-model inference'])
    summary.append(['Stage and harvest accuracy', 'Not calculated without measured stage labels'])
    summary.append(['Representative capture', '같은 촬영의 정면 쌍, 기존 카메라 품질점수, 시각 순으로 선정; 실측Value 미사용'])
    summary.column_dimensions['A'].width = 32
    summary.column_dimensions['B'].width = 100
    output.parent.mkdir(parents=True, exist_ok=True)
    temporary = output.with_name(output.stem + '.writing.xlsx')
    try:
        book.save(temporary)
        os.replace(temporary, output)
    finally:
        if temporary.exists():
            temporary.unlink()
    report = {**dict(counts), **metrics, **prediction_counts, 'output': str(output), 'model_probability_source': 'D-step A-fold mean probabilities', 'qc_source': str(qc_xlsx), 'exact_front_source': str(front_csv)}
    report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
    return report
