"""Processing utilities for growth labels."""
from __future__ import annotations
import sys as _sys
from pathlib import Path as _Path
_CYCLE_ROOT = _Path(__file__).resolve().parents[1]
if str(_CYCLE_ROOT) not in _sys.path:
    _sys.path.insert(0, str(_CYCLE_ROOT))
import csv
import json
import math
import os
import re
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path
from openpyxl import load_workbook
from tqdm import tqdm
ALIASES = {'date': ('조사일', 'survey_date', 'date'), 'bed': ('베드번호', 'bed_id', 'bed'), 'individual': ('개체번호', 'individual_number', 'plant_number'), 'weight': ('무게(g)', '무게', 'fresh weight', 'fresh_weight(g)', 'fresh_weight_g', 'freshweight')}
FIELDS = ['lettuce_key', 'weight_g', 'status', 'source_excel_rows', 'heading']

def _normal(value: object) -> str:
    return re.sub('[\\s_()\\-]', '', str(value).strip().lower())

def inspect(path: Path) -> dict:
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = book.worksheets[0]
        header = next(sheet.iter_rows(values_only=True), ())
        lookup = defaultdict(list)
        for (i, value) in enumerate(header):
            if value is not None:
                lookup[_normal(value)].append(i)
        columns = {}
        warnings = []
        for (kind, names) in ALIASES.items():
            found = list(dict.fromkeys((i for name in names for i in lookup[_normal(name)])))
            if not found:
                raise ValueError(f'Invalid input or state in inspect (rejected condition: not found): {kind} {names} ')
            if len(found) > 1:
                if kind != 'weight':
                    raise ValueError(f"Invalid input or state in inspect (rejected condition: kind != 'weight'): {kind} {names} ")
                warnings.append(f'Weight columns: {len(found)} candidates; selected priority {header[found[0]]!s}')
            columns[kind] = found[0]
        columns['heading'] = next(iter(lookup[_normal('결구')] or lookup[_normal('heading')]), None)
        return {'sheet': sheet.title, 'columns': columns, 'warnings': warnings}
    finally:
        book.close()

def lettuce_key(survey: object, bed: object, individual: object) -> str:
    if isinstance(survey, datetime):
        day = survey.date()
    elif isinstance(survey, date):
        day = survey
    else:
        day = datetime.fromisoformat(str(survey)[:10]).date()
    match = re.fullmatch('(?:bed)?0*(\\d+)(?:\\.0)?', str(bed).strip(), re.I)
    if not match:
        raise ValueError(f'Invalid input or state in lettuce_key (rejected condition: not match): {bed} ')
    plant = int(float(individual))
    if not 1 <= plant <= 24 or float(individual) != plant:
        raise ValueError(f'Invalid input or state in lettuce_key (rejected condition: not 1 <= plant <= 24 or float(individual) != plant): {individual} ')
    slot = f't{plant}' if plant <= 12 else f'b{plant - 12}'
    return f'{day:%y%m%d}_bed{int(match.group(1)):02d}_{slot}'

def image_key(capture_date: object, bed_id: object, physical_slot: object) -> str:
    day = str(capture_date).strip().replace('-', '').removesuffix('.0')
    if not re.fullmatch('20\\d{6}', day):
        raise ValueError(f"Invalid input or state in image_key (rejected condition: not re.fullmatch('20\\\\d{{6}}', day)): {capture_date} ")
    bed = re.fullmatch('bed0*(\\d+)', str(bed_id).strip(), re.I)
    slot = re.fullmatch('([tb])([1-9]|1[012])', str(physical_slot).strip(), re.I)
    if not bed or not slot:
        raise ValueError(f'Invalid input or state in image_key (rejected condition: not bed or not slot): {bed_id} {physical_slot} ')
    return f'{day[2:]}_bed{int(bed.group(1)):02d}_{slot.group(1).lower()}{int(slot.group(2))}'

def build(source: Path, output: Path) -> dict:
    meta = inspect(source)
    ix = meta['columns']
    grouped = defaultdict(list)
    invalid = []
    book = load_workbook(source, read_only=True, data_only=True)
    try:
        rows = book.worksheets[0].iter_rows(values_only=True)
        next(rows)
        for (number, row) in enumerate(tqdm(rows, desc='Growth labels', unit='row'), start=2):
            if not any((value is not None for value in row)):
                continue
            try:
                key = lettuce_key(row[ix['date']], row[ix['bed']], row[ix['individual']])
                weight = float(row[ix['weight']])
                if not math.isfinite(weight) or weight < 0:
                    raise ValueError('Invalid input or state in build (rejected condition: not math.isfinite(weight) or weight < 0)')
                heading = row[ix['heading']] if ix['heading'] is not None else None
                grouped[key].append((number, weight, heading))
            except (ValueError, TypeError, IndexError) as exc:
                invalid.append({'excel_row': number, 'reason': str(exc)})
    finally:
        book.close()
    output.parent.mkdir(parents=True, exist_ok=True)
    temp = output.with_suffix(output.suffix + '.writing')
    conflicts = 0
    with temp.open('w', encoding='utf-8-sig', newline='') as handle:
        writer = csv.DictWriter(handle, FIELDS)
        writer.writeheader()
        for (key, records) in sorted(grouped.items()):
            weights = {weight for (_, weight, _) in records}
            status = 'OK' if len(weights) == 1 else 'CONFLICT'
            conflicts += status == 'CONFLICT'
            writer.writerow({'lettuce_key': key, 'weight_g': records[0][1] if status == 'OK' else '', 'status': status, 'source_excel_rows': ';'.join((str(x[0]) for x in records)), 'heading': records[0][2] if status == 'OK' and records[0][2] is not None else ''})
    os.replace(temp, output)
    report = {'status': 'READY', 'source': str(source), 'sheet': meta['sheet'], 'output': str(output), 'keys': len(grouped), 'conflicts': conflicts, 'invalid_rows': len(invalid), 'invalid_examples': invalid[:20], 'warnings': meta['warnings'], 'rule': '1..12=t1..t12; 13..24=b1..b12; exact survey date + bed + slot', 'role': 'post-inference evaluation only; unlabeled images remain available for history; weight never enters model X'}
    output.with_suffix('.summary.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
    return report
