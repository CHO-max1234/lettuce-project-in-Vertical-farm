"""Processing utilities for stage inputs."""
from __future__ import annotations
import sys as _sys
from pathlib import Path as _Path
_CYCLE_ROOT = _Path(__file__).resolve().parents[1]
if str(_CYCLE_ROOT) not in _sys.path:
    _sys.path.insert(0, str(_CYCLE_ROOT))
import csv
import hashlib
import json
import math
import os
import re
import zipfile
from collections import Counter, defaultdict
from pathlib import Path
from xml.etree import ElementTree as ET
from openpyxl import Workbook, load_workbook
from tqdm import tqdm
NS = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
REL = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}'
csv.field_size_limit(100000000)

def norm(value: object) -> str:
    if value is None:
        return ''
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()

def yes(value: object) -> bool:
    return norm(value).lower() in {'true', '1', '1.0', 'yes'}

def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open('rb') as handle:
        for block in iter(lambda : handle.read(1024 * 1024), b''):
            digest.update(block)
    return digest.hexdigest()

def image_key(day: object, bed: object, slot: object) -> str:
    date = norm(day).removesuffix('.0').replace('-', '')
    match = re.fullmatch('bed0*(\\d+)', norm(bed), re.I)
    place = re.fullmatch('([tb])([1-9]|1[012])', norm(slot), re.I)
    if not re.fullmatch('20\\d{6}', date) or not match or (not place):
        raise ValueError(f"Invalid input or state in image_key (rejected condition: not re.fullmatch('20\\\\d{{6}}', date) or not match or (not place)): {day} {bed} {slot} ")
    return f'{date[2:]}_bed{int(match.group(1)):02d}_{place.group(1).lower()}{int(place.group(2))}'

def labels_from_raw(path: Path) -> tuple[set[str], dict]:
    from growth_labels import inspect, lettuce_key
    info = inspect(path)
    ix = info['columns']
    values: dict[str, set[float]] = defaultdict(set)
    invalid = 0
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = book.worksheets[0].iter_rows(values_only=True)
        next(rows)
        for row in rows:
            if not any((v is not None for v in row)):
                continue
            try:
                key = lettuce_key(row[ix['date']], row[ix['bed']], row[ix['individual']])
                weight = float(row[ix['weight']])
                if not math.isfinite(weight) or weight < 0:
                    raise ValueError('Invalid input or state in labels_from_raw (rejected condition: not math.isfinite(weight) or weight < 0)')
                values[key].add(weight)
            except (ValueError, TypeError, IndexError):
                invalid += 1
    finally:
        book.close()
    valid = {key for (key, weights) in values.items() if len(weights) == 1}
    return (valid, {'valid_weight_keys': len(valid), 'conflict_keys': sum((len(v) > 1 for v in values.values())), 'invalid_rows': invalid, 'growth_sheet': info['sheet']})

def labels_from_csv(path: Path) -> set[str]:
    with path.open(encoding='utf-8-sig', newline='') as handle:
        return {row['lettuce_key'] for row in csv.DictReader(handle) if row['status'] == 'OK'}

def select_records(crop_list: Path, labels: set[str]) -> tuple[set[str], dict[str, tuple[str, str]], dict]:
    top_keys: set[str] = set()
    total = 0
    with crop_list.open(encoding='utf-8-sig', newline='') as handle:
        rows = csv.DictReader(handle)
        required = {'record_id', 'camera', 'bed_id', 'capture_date', 'physical_slot', 'crop_path', 'crop_sha256'}
        if required - set(rows.fieldnames or []):
            raise ValueError(f'Invalid input or state in select_records (rejected condition: required - set(rows.fieldnames or [])): {sorted(required - set(rows.fieldnames or []))} ')
        for row in rows:
            total += 1
            if row['camera'].lower() not in {'cam0', 'cam1'}:
                continue
            key = image_key(row['capture_date'], row['bed_id'], row['physical_slot'])
            if key in labels:
                top_keys.add(key)
    if not top_keys:
        raise ValueError('Invalid input or state in select_records (rejected condition: not top_keys)')
    latest: dict[tuple[str, str], str] = {}
    for key in top_keys:
        (day, bed, slot) = key.split('_')
        track = (bed, slot)
        latest[track] = max(latest.get(track, ''), '20' + day)
    selected: dict[str, tuple[str, str]] = {}
    count = Counter()
    with crop_list.open(encoding='utf-8-sig', newline='') as handle:
        for row in csv.DictReader(handle):
            (bed, slot, day) = (norm(row['bed_id']), norm(row['physical_slot']), norm(row['capture_date']))
            if day > latest.get((bed, slot), ''):
                continue
            rid = row['record_id']
            if rid in selected:
                raise ValueError(f'Invalid input or state in select_records (rejected condition: rid in selected): {rid} ')
            selected[rid] = (row['crop_path'], row['crop_sha256'])
            count[row['camera'].lower()] += 1
    if not selected:
        raise ValueError('Invalid input or state in select_records (rejected condition: not selected)')
    return (top_keys, selected, {'usable_crop_rows': total, 'measured_top_keys': len(top_keys), 'history_crop_rows': len(selected), 'history_cameras': dict(count)})

def sheet_rows(path: Path, sheet_name: str):
    with zipfile.ZipFile(path) as archive:
        if 'xl/sharedStrings.xml' in archive.namelist():
            raise ValueError("Invalid input or state in sheet_rows (rejected condition: 'xl/sharedStrings.xml' in archive.namelist())")
        workbook = ET.fromstring(archive.read('xl/workbook.xml'))
        relationships = {item.attrib['Id']: item.attrib['Target'] for item in ET.fromstring(archive.read('xl/_rels/workbook.xml.rels'))}
        matches = [item for item in workbook.iter(NS + 'sheet') if item.attrib.get('name') == sheet_name]
        if len(matches) != 1:
            raise ValueError(f'Invalid input or state in sheet_rows (rejected condition: len(matches) != 1): {sheet_name!r} ')
        target = relationships[matches[0].attrib[REL + 'id']]
        member = target.lstrip('/') if target.startswith('/') else 'xl/' + target
        with archive.open(member) as handle:
            parent = None
            for (event, element) in ET.iterparse(handle, events=('start', 'end')):
                if event == 'start' and element.tag == NS + 'sheetData':
                    parent = element
                if event != 'end' or element.tag != NS + 'row':
                    continue
                values = []
                for cell in element:
                    if cell.tag != NS + 'c':
                        continue
                    if cell.find(NS + 'f') is not None:
                        raise ValueError("Invalid input or state in sheet_rows (rejected condition: cell.find(NS + 'f') is not None)")
                    letters = re.match('[A-Z]+', cell.attrib['r']).group()
                    column = 0
                    for letter in letters:
                        column = column * 26 + ord(letter) - 64
                    values.extend([None] * max(0, column - len(values)))
                    kind = cell.attrib.get('t', 'n')
                    value = cell.find(NS + 'v')
                    if kind == 'inlineStr':
                        result = ''.join((node.text or '' for node in cell.iter(NS + 't')))
                    elif value is None or value.text is None:
                        result = None
                    elif kind == 'b':
                        result = value.text == '1'
                    elif kind == 'n':
                        result = float(value.text)
                    else:
                        result = value.text
                    values[column - 1] = result
                number = int(element.attrib['r'])
                element.clear()
                if parent is not None:
                    parent.remove(element)
                yield (number, values)

def feature_header(path: Path) -> list[str]:
    rows = sheet_rows(path, '전체')
    (number, values) = next(rows)
    if number != 1:
        raise ValueError('Invalid input or state in feature_header (rejected condition: number != 1)')
    header = [norm(value) for value in values]
    required = {'record_id', 'camera', 'bed_id', 'capture_date', 'capture_time', 'physical_slot', 'crop_path', 'crop_sha256', 'pair_id', 'selected_record_id', 'selected_camera', '분석가능', 'feature_status', 'front_height_cm', 'area_cm2', 'circularity'}
    if required - set(header) or len(header) != len(set(header)):
        raise ValueError(f'Invalid input or state in feature_header (rejected condition: required - set(header) or len(header) != len(set(header))): {sorted(required - set(header))} ')
    return header

def qc_check(path: Path, target_keys: set[str]) -> dict:
    columns = {'cycle_id', 'date', 'bed_name', 'slot_name', 'source_pair_id', 'growth_qc', 'relation_qc', 'timeseries_qc'}
    found: set[str] = set()
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = book.worksheets[0]
        rows = sheet.iter_rows(values_only=True)
        header = [norm(x) for x in next(rows)]
        if columns - set(header):
            raise ValueError(f'Missing QC worksheet columns: {sorted(columns - set(header))}')
        idx = header.index('source_pair_id')
        for row in tqdm(rows, desc='qc check: QC preflight', unit='row', mininterval=5):
            parts = norm(row[idx]).split('|')
            if len(parts) != 5 or not re.fullmatch('20\\d{6}', parts[1]):
                continue
            key = image_key(parts[1], parts[2], parts[3])
            if key in target_keys:
                if key in found:
                    raise ValueError(f'Invalid input or state in qc_check (rejected condition: key in found): {key} ')
                found.add(key)
    finally:
        book.close()
    return {'qc_sheet': sheet.title, 'matched_target_keys': len(found), 'target_without_qc': len(target_keys - found)}

def preflight(paths: dict[str, Path], *, check_hashes: bool=True) -> tuple[set[str], dict]:
    required = ('crop_list', 'crop_manifest', 'crop_root', 'features', 'qc', 'environment', 'growth', 'stage_model', 'a_model', 'c_models', 'd_rule')
    for name in required:
        if not paths[name].exists():
            raise FileNotFoundError(f'{name}: {paths[name]}')
    from environment_features import inspect as inspect_environment
    from model_contract import component_contracts
    env = inspect_environment(paths['environment'])
    (labels, growth) = labels_from_raw(paths['growth'])
    (targets, selected, selection) = select_records(paths['crop_list'], labels)
    header = feature_header(paths['features'])
    qc = qc_check(paths['qc'], targets)
    if qc['target_without_qc']:
        raise ValueError(f"Invalid input or state in preflight (rejected condition: qc['target_without_qc']): {qc['target_without_qc']} ")
    bundle = paths['a_model'] / 'selected_model_bundle.json'
    if not bundle.is_file():
        raise FileNotFoundError(bundle)
    for fold in json.loads(bundle.read_text(encoding='utf-8'))['folds']:
        for item in fold['files']:
            if not (paths['a_model'] / item['path']).is_file():
                raise FileNotFoundError(paths['a_model'] / item['path'])
    for (mode, (folder, metadata)) in component_contracts(paths['c_models']).items():
        for item in metadata['seeds']:
            model = folder / f"seed_{item['seed']}.cbm"
            if not model.is_file():
                raise FileNotFoundError(model)
    if check_hashes:
        expected = 'dc2b180b471ba47f67e617d4d6aa3f43c08016460dee6cf9cfb8ead226587eb2'
        if sha256(paths['stage_model']) != expected:
            raise ValueError("Invalid input or state in preflight (rejected condition: sha256(paths['stage_model']) != expected)")
    filter_summary = paths['crop_manifest'].parent / '_2차필터_20260915' / 'summary.json'
    if filter_summary.is_file():
        old = json.loads(filter_summary.read_text(encoding='utf-8'))
        if old.get('true_images') != selection['usable_crop_rows']:
            raise ValueError("Invalid input or state in preflight (rejected condition: old.get('true_images') != selection['usable_crop_rows'])")
    report = {**growth, **selection, **qc, 'feature_columns': len(header), 'environment_sheet': env['sheet'], 'history_policy': 'same bed/slot through last measured day', 'target_policy': 'measured date with usable top image; no model fitting', 'environment_schema': 'manuscript12'}
    return (targets, report)

def build_subset(paths: dict[str, Path], labels: set[str], output_dir: Path) -> dict:
    (targets, selected, stats) = select_records(paths['crop_list'], labels)
    source_stats = {name: (paths[name].stat().st_size, paths[name].stat().st_mtime_ns) for name in ('features', 'crop_manifest', 'crop_list')}
    output_dir.mkdir(parents=True, exist_ok=True)
    feature_out = output_dir / 'model_feature_subset.xlsx'
    front_out = output_dir / 'same_capture_front.csv'
    feature_tmp = output_dir / 'model_feature_subset.writing.xlsx'
    front_tmp = output_dir / 'same_capture_front.writing.csv'
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet('전체')
    sheet.freeze_panes = 'A2'
    found: set[str] = set()
    source_rows = 0
    rows = sheet_rows(paths['features'], '전체')
    (_, raw_header) = next(rows)
    header = [norm(x) for x in raw_header]
    position = {name: header.index(name) for name in ('record_id', 'camera', 'crop_path', 'crop_sha256', '분석가능', 'pair_id', 'selected_record_id', 'feature_status', 'front_height_cm')}
    sheet.append(header)
    try:
        with front_tmp.open('w', encoding='utf-8-sig', newline='') as handle:
            writer = csv.DictWriter(handle, fieldnames=['record_id', 'pair_id', 'front_height_cm'])
            writer.writeheader()
            for (excel_row, values) in tqdm(rows, desc='build subset: feature', unit='row', mininterval=5):
                source_rows += 1
                rid = norm(values[position['record_id']])
                if rid not in selected:
                    continue
                if rid in found:
                    raise ValueError(f'Invalid input or state in build_subset (rejected condition: rid in found): {rid} ')
                found.add(rid)
                (expected_path, expected_hash) = selected[rid]
                if norm(values[position['crop_path']]).casefold() != expected_path.casefold():
                    raise ValueError(f"Invalid input or state in build_subset (rejected condition: norm(values[position['crop_path']]).casefold() != expected_path.casefold()): {rid} ")
                if norm(values[position['crop_sha256']]) != expected_hash:
                    raise ValueError(f"Invalid input or state in build_subset (rejected condition: norm(values[position['crop_sha256']]) != expected_hash): {rid} ")
                if not yes(values[position['분석가능']]):
                    raise ValueError(f"Invalid input or state in build_subset (rejected condition: not yes(values[position['분석가능']])): {rid} ")
                values.extend([None] * (len(header) - len(values)))
                sheet.append(values)
                if norm(values[position['camera']]).lower() == 'cam2' and norm(values[position['selected_record_id']]) == rid and (norm(values[position['feature_status']]) == 'OK'):
                    height = values[position['front_height_cm']]
                    try:
                        height = float(height)
                    except (ValueError, TypeError):
                        continue
                    if math.isfinite(height) and height > 0:
                        writer.writerow({'record_id': rid, 'pair_id': norm(values[position['pair_id']]), 'front_height_cm': height})
        if source_rows <= 0 or found != set(selected):
            raise ValueError(f'Invalid input or state in build_subset (rejected condition: source_rows <= 0 or found != set(selected)): {source_rows} {len(selected)} {len(found)} ')
        workbook.save(feature_tmp)
        os.replace(feature_tmp, feature_out)
        os.replace(front_tmp, front_out)
    finally:
        workbook.close()
        for temporary in (feature_tmp, front_tmp):
            if temporary.exists():
                temporary.unlink()
    manifest_out = output_dir / 'model_crop_manifest.csv'
    manifest_tmp = output_dir / 'model_crop_manifest.writing.csv'
    manifest_found: set[str] = set()
    manifest_rows = 0
    try:
        with paths['crop_manifest'].open(encoding='utf-8-sig', newline='') as src, manifest_tmp.open('w', encoding='utf-8-sig', newline='') as dst:
            reader = csv.DictReader(src)
            writer = csv.DictWriter(dst, fieldnames=reader.fieldnames)
            writer.writeheader()
            for row in tqdm(reader, desc='build subset: mask manifest', unit='row', mininterval=5):
                manifest_rows += 1
                rid = row['record_id']
                if rid not in selected:
                    continue
                if rid in manifest_found:
                    raise ValueError(f'Invalid input or state in build_subset (rejected condition: rid in manifest_found): {rid} ')
                if row['crop_path'].casefold() != selected[rid][0].casefold():
                    raise ValueError(f"Invalid input or state in build_subset (rejected condition: row['crop_path'].casefold() != selected[rid][0].casefold()): {rid} ")
                manifest_found.add(rid)
                writer.writerow(row)
        if manifest_rows != source_rows or manifest_found != found:
            raise ValueError(f'Invalid input or state in build_subset (rejected condition: manifest_rows != source_rows or manifest_found != found): {manifest_rows} {len(manifest_found)} ')
        for (name, expected) in source_stats.items():
            path = paths[name]
            if (path.stat().st_size, path.stat().st_mtime_ns) != expected:
                raise RuntimeError(f'Invalid input or state in build_subset (rejected condition: (path.stat().st_size, path.stat().st_mtime_ns) != expected): {path} ')
        os.replace(manifest_tmp, manifest_out)
    finally:
        if manifest_tmp.exists():
            manifest_tmp.unlink()
    return {**stats, 'feature_source_rows': source_rows, 'manifest_source_rows': manifest_rows, 'subset_rows': len(found), 'targets': len(targets), 'feature_file': str(feature_out), 'manifest_file': str(manifest_out), 'front_file': str(front_out)}
