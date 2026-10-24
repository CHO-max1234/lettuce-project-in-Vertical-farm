"""Processing utilities for environment features."""
from __future__ import annotations
import sys as _sys
from pathlib import Path as _Path
_CYCLE_ROOT = _Path(__file__).resolve().parents[1]
if str(_CYCLE_ROOT) not in _sys.path:
    _sys.path.insert(0, str(_CYCLE_ROOT))
import csv
import json
import math
import os
from collections import defaultdict
from datetime import date, datetime, time
from pathlib import Path
from openpyxl import load_workbook
from tqdm import tqdm
BASE_COLUMNS = ['date', 'temp_mean', 'heat_hours_28_ratio', 'heat_degree_hours_28', 'vpd_hours_ratio_1_8', 'vpd_degree_hours_1_8', 'hot_dry_load_28_1_8', 'water_air_gap_day', 'water_air_gap_day_abs', 'ec_mean', 'minute_count', 'valid_stress_hour_count']
MIN_STRESS_MINUTES = 48
ALIASES = {'date': ('날짜', 'date', '측정일'), 'time': ('시간', 'time', '측정시간'), 'temp': ('temperature_avg', 'temperature', 'temp_avg', '온도'), 'rh': ('humidity_avg', 'humidity', 'rh_avg', '습도'), 'ec': ('ec_avg', 'ec', 'EC'), 'water': ('water_temp_avg', 'water_temp', '수온'), 'daynight': ('주야간', 'daynight', 'day_night')}

def _column(header: list[str], names: tuple[str, ...], *, required: bool=True) -> int | None:
    normalized = {str(v).strip().lower(): i for (i, v) in enumerate(header)}
    for name in names:
        if name.lower() in normalized:
            return normalized[name.lower()]
    if required:
        raise ValueError(f'Invalid input or state in _column (rejected condition: required): {names} ')
    return None

def inspect(path: Path) -> dict:
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = book.worksheets[0]
        header = next(sheet.iter_rows(values_only=True), ())
        index = {key: _column(list(header), names, required=key in {'date', 'time', 'temp', 'rh'}) for (key, names) in ALIASES.items()}
        return {'sheet': sheet.title, 'columns': index}
    finally:
        book.close()

def _number(value: object) -> float | None:
    try:
        result = float(value)
        return result if math.isfinite(result) else None
    except (ValueError, TypeError):
        return None

def _day(value: object) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return datetime.fromisoformat(str(value)[:10]).date()

def _hour(value: object) -> int:
    if isinstance(value, (datetime, time)):
        return value.hour
    return int(str(value).split(':', 1)[0])

def _vpd(temp_c: float, rh: float) -> float:
    saturation = 0.6108 * math.exp(17.27 * temp_c / (temp_c + 237.3))
    return saturation * (1 - rh / 100)

def build(raw: Path, output: Path) -> dict:
    metadata = inspect(raw)
    ix = metadata['columns']
    book = load_workbook(raw, read_only=True, data_only=True)
    hourly = defaultdict(lambda : {'n': 0, 'temp': 0.0, 'rh': 0.0, 'vpd': 0.0, 'ec': 0.0, 'ec_n': 0, 'water': 0.0, 'water_n': 0, 'day_temp': 0.0, 'day_rh': 0.0, 'day_n': 0, 'day_water': 0.0, 'day_water_n': 0})
    skipped = 0
    try:
        rows = book.worksheets[0].iter_rows(values_only=True)
        next(rows)
        for row in tqdm(rows, desc='Minute environment', unit='row', mininterval=5):
            try:
                (day, hour) = (_day(row[ix['date']]), _hour(row[ix['time']]))
                (temp, rh) = (_number(row[ix['temp']]), _number(row[ix['rh']]))
                if temp is None or rh is None or (not -20 <= temp <= 60) or (not 0 <= rh <= 100) or (not 0 <= hour <= 23):
                    skipped += 1
                    continue
            except (ValueError, TypeError, IndexError):
                skipped += 1
                continue
            bucket = hourly[day, hour]
            bucket['n'] += 1
            bucket['temp'] += temp
            bucket['rh'] += rh
            bucket['vpd'] += _vpd(temp, rh)
            ec = _number(row[ix['ec']]) if ix['ec'] is not None else None
            if ec is not None and ec >= 0:
                bucket['ec'] += ec / 1000 if ec > 20 else ec
                bucket['ec_n'] += 1
            water = _number(row[ix['water']]) if ix['water'] is not None else None
            if water is not None and -20 <= water <= 60:
                bucket['water'] += water
                bucket['water_n'] += 1
            label = str(row[ix['daynight']]).strip() if ix['daynight'] is not None else ''
            daylight = label in {'주간', 'day', 'DAY'} if label else False
            if daylight:
                bucket['day_temp'] += temp
                bucket['day_rh'] += rh
                bucket['day_n'] += 1
                if water is not None and -20 <= water <= 60:
                    bucket['day_water'] += water
                    bucket['day_water_n'] += 1
    finally:
        book.close()
    by_day = defaultdict(list)
    for ((day, hour), values) in sorted(hourly.items()):
        by_day[day].append(values)
    output.parent.mkdir(parents=True, exist_ok=True)
    temp_output = output.with_suffix(output.suffix + '.writing')
    with temp_output.open('w', encoding='utf-8-sig', newline='') as handle:
        writer = csv.DictWriter(handle, BASE_COLUMNS)
        writer.writeheader()
        for (day, hours) in sorted(by_day.items()):
            minutes = sum((h['n'] for h in hours))
            valid = [h for h in hours if h['n'] >= MIN_STRESS_MINUTES]
            stresses = []
            for h in valid:
                stresses.append((h['temp'] / h['n'], h['vpd'] / h['n']))
            n = len(stresses)
            (day_t, day_tn) = (sum((h['day_temp'] for h in hours)), sum((h['day_n'] for h in hours)))
            (day_w, day_wn) = (sum((h['day_water'] for h in hours)), sum((h['day_water_n'] for h in hours)))
            gap = day_w / day_wn - day_t / day_tn if day_wn and day_tn else None
            ec_n = sum((h['ec_n'] for h in hours))
            row = {'date': day.isoformat(), 'temp_mean': sum((h['temp'] for h in hours)) / minutes, 'minute_count': minutes, 'valid_stress_hour_count': n, 'heat_hours_28_ratio': sum((t > 28 for (t, _) in stresses)) / n if n else None, 'heat_degree_hours_28': sum((max(t - 28, 0) for (t, _) in stresses)) if n else None, 'vpd_hours_ratio_1_8': sum((v >= 1.8 for (_, v) in stresses)) / n if n else None, 'vpd_degree_hours_1_8': sum((max(v - 1.8, 0) for (_, v) in stresses)) if n else None, 'hot_dry_load_28_1_8': sum((max(t - 28, 0) * max(v - 1.8, 0) for (t, v) in stresses)) if n else None, 'water_air_gap_day': gap, 'water_air_gap_day_abs': abs(gap) if gap is not None else None, 'ec_mean': sum((h['ec'] for h in hours)) / ec_n if ec_n else None}
            writer.writerow(row)
    os.replace(temp_output, output)
    report = {'status': 'READY', 'source': str(raw), 'sheet': metadata['sheet'], 'daily_output': str(output), 'days': len(by_day), 'minute_rows_skipped': skipped, 'method': 'minute -> hourly means (>=48 minutes for stress) -> 12 daily base columns; C-step derives the 12 manuscript environmental features', 'scope': 'raw reconstruction; parity with frozen 2nd daily input must be checked before accuracy comparison'}
    output.with_suffix('.summary.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
    return report
